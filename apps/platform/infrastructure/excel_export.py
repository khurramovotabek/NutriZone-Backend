"""
"Excel'ga yuklash" for any Dashboard list: one mixin, the same professional workbook everywhere.

    class OrderAdmin(ExcelExportMixin, admin.ModelAdmin):
        excel_title = "Buyurtmalar"
        excel_columns = [Col("Raqam", "order_number", 18), Col("Summa", lambda o: o.x_total, 14, money=True), ...]

        def excel_prepare(self, objs):  # optional: bulk-load what the getters need (no N+1)
            ...

What it adds:
  * <app>/<model>/export.xlsx   the WHOLE list exactly as filtered / searched on the page (not one page), up to
                                `excel_limit` rows, for anyone who may view the list
  * an admin action             "Tanlanganlarni Excel'ga yuklash" for ticked rows
  * `nz_excel_url` in the page  rendered by admin/includes/excel_button.html (classic lists get it in the object
                                tools, premium / custom pages place the include in their hero actions)

The workbook: branded title row with what was exported and when, bold header, frozen header + autofilter, column
widths, money as "# ##0", dates in Tashkent time, a totals row for money columns.
"""
from __future__ import annotations

import re

import io
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.contrib import admin
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.urls import path, reverse
from django.utils import timezone

TZ = ZoneInfo("Asia/Tashkent")
BRAND = "1E3A8A"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class Col:
    header: str
    value: str | Callable
    width: int = 16
    money: bool = False  # "# ##0" format + a total at the bottom
    number: bool = False  # integer format, totalled as well

    def get(self, obj):
        if callable(self.value):
            return self.value(obj)
        cur = obj
        for part in self.value.split("."):
            cur = getattr(cur, part, None)
            if cur is None:
                return None
            if callable(cur):
                cur = cur()
        return cur


def cell(value):
    """A value openpyxl stores natively: aware datetimes in Tashkent time, enums / UUIDs / lists as text."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Ha" if value else "Yo'q"
    if isinstance(value, datetime):
        return timezone.localtime(value, TZ).replace(tzinfo=None) if timezone.is_aware(value) else value
    if isinstance(value, int | float | Decimal | date):
        return value
    if isinstance(value, list | tuple | set):
        return ", ".join(str(x) for x in value)
    return str(value)


def workbook(title: str, columns: list[Col], rows, subtitle: str = "") -> bytes:
    """rows: iterable of objects (read through `columns`)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:31] or "Sheet"  # Excel forbids []:*?/\ in sheet names
    n = len(columns)
    stamp = timezone.localtime(timezone.now(), TZ).strftime("%d.%m.%Y %H:%M")
    ws.append([f"NutriZone · {title}"] + [None] * (n - 1))
    ws.append([f"{subtitle + ' · ' if subtitle else ''}Yuklangan: {stamp}"] + [None] * (n - 1))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(n, 1))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(n, 1))
    ws["A1"].font = Font(bold=True, size=14, color=BRAND)
    ws["A2"].font = Font(italic=True, size=10, color="64748B")
    ws.append([c.header for c in columns])
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=BRAND)
    thin = Side(style="thin", color="CBD5E1")
    for c in ws[3]:
        c.font, c.fill = head, fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    ws.row_dimensions[3].height = 30
    count = 0
    for obj in rows:
        ws.append([cell(c.get(obj)) for c in columns])
        count += 1
    first, last = 4, 3 + count
    for k, col in enumerate(columns, 1):
        letter = get_column_letter(k)
        ws.column_dimensions[letter].width = col.width
        fmt = "#,##0" if (col.money or col.number) else None
        for (value_cell,) in ws.iter_rows(min_row=first, max_row=last, min_col=k, max_col=k):
            if isinstance(value_cell.value, datetime):
                value_cell.number_format = "dd.mm.yyyy hh:mm"
            elif isinstance(value_cell.value, date):
                value_cell.number_format = "dd.mm.yyyy"
            elif fmt and isinstance(value_cell.value, int | float | Decimal):
                value_cell.number_format = fmt
    if count and any(c.money or c.number for c in columns):
        total_row = last + 1
        ws.cell(row=total_row, column=1, value="JAMI").font = Font(bold=True)
        for k, col in enumerate(columns, 1):
            if (col.money or col.number) and k > 1:
                letter = get_column_letter(k)
                c = ws.cell(row=total_row, column=k, value=f"=SUM({letter}{first}:{letter}{last})")
                c.font, c.number_format = Font(bold=True), "#,##0"
                c.border = Border(top=Side(style="medium", color=BRAND))
    ws.freeze_panes = "A4"
    if count:
        ws.auto_filter.ref = f"A3:{get_column_letter(n)}{last}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def response(content: bytes, name: str) -> HttpResponse:
    stamp = timezone.localtime(timezone.now(), TZ).strftime("%Y-%m-%d")
    resp = HttpResponse(content, content_type=XLSX)
    resp["Content-Disposition"] = f'attachment; filename="nutrizone_{name}_{stamp}.xlsx"'
    return resp


class ExcelExportMixin:
    excel_title: str = ""
    excel_columns: list[Col] = []
    excel_limit = 50000
    excel_select_related: tuple = ()
    excel_prefetch: tuple = ()

    # ---- hooks
    def excel_prepare(self, objs: list) -> None:
        """Bulk-load what the column getters need (set attributes on the objects)."""

    def get_excel_columns(self, request) -> list[Col]:
        return list(self.excel_columns)

    def excel_name(self) -> str:
        return self.model._meta.model_name

    # ---- plumbing
    def _excel_url_name(self) -> str:
        return f"{self.model._meta.app_label}_{self.model._meta.model_name}_excel"

    def get_urls(self):
        return [path("export.xlsx", self.admin_site.admin_view(self.excel_export_view), name=self._excel_url_name()), *super().get_urls()]

    def excel_url(self, request) -> str:
        q = request.GET.copy()
        for key in ("p", "o", "_changelist_filters", "e"):
            q.pop(key, None)
        url = reverse(f"admin:{self._excel_url_name()}")
        return f"{url}?{q.urlencode()}" if q else url

    def changelist_view(self, request, extra_context=None):
        extra_context = {**(extra_context or {}), "nz_excel_url": self.excel_url(request) if self.has_view_permission(request) else None}
        return super().changelist_view(request, extra_context)

    def _excel_build(self, request, queryset, subtitle: str) -> HttpResponse:
        qs = queryset
        if self.excel_select_related:
            qs = qs.select_related(*self.excel_select_related)
        if self.excel_prefetch:
            qs = qs.prefetch_related(*self.excel_prefetch)
        objs = list(qs[: self.excel_limit])
        self.excel_prepare(objs)
        title = self.excel_title or str(self.model._meta.verbose_name_plural).capitalize()
        return response(workbook(title, self.get_excel_columns(request), objs, subtitle), self.excel_name())

    def excel_export_view(self, request):
        """The whole list as filtered on the page (same filters / search / tabs), not just the visible page."""
        if not self.has_view_permission(request):
            raise PermissionDenied
        cl = self.get_changelist_instance(request)
        qs = cl.get_queryset(request)
        n = qs.count()
        note = f"{n} ta yozuv" + (f" (birinchi {self.excel_limit} tasi)" if n > self.excel_limit else "")
        if request.GET:
            note += " · filtr qo'llangan"
        return self._excel_build(request, qs, note)

    @admin.action(description="📥 Tanlanganlarni Excel'ga yuklash")
    def excel_export_selected(self, request, queryset):
        return self._excel_build(request, queryset, f"{queryset.count()} ta tanlangan yozuv")

    def get_actions(self, request):
        actions = super().get_actions(request)
        if actions and self.has_view_permission(request):  # only where the list already has bulk actions (checkboxes)
            func = type(self).excel_export_selected
            actions["excel_export_selected"] = (func, "excel_export_selected", func.short_description)
        return actions


# ---------------------------------------------------------------- bulk lookups for excel_prepare


def people(ids) -> dict:
    """{user id: (full name, phone)} for customer / cashier UUID columns."""
    from apps.identity.infrastructure.models import User

    ids = {i for i in ids if i}
    return {u.id: (u.full_name or "", u.phone_number or "") for u in User.objects.filter(id__in=ids).only("full_name", "phone_number")}


def stores(ids) -> dict:
    from apps.stores.infrastructure.models import Store

    ids = {i for i in ids if i}
    return dict(Store.objects.filter(id__in=ids).values_list("id", "name"))
