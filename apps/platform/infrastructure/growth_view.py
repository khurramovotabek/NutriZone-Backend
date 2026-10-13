"""
Dashboard growth pages (same audience as Analitika: order + till-sale viewers):

  /django-admin/sales-plan/        Sotuv rejasi -- monthly targets per branch / company, progress, forecast, sellers
  /django-admin/segments/          Mijozlar segmentlari -- RFM segments, members, export, "send a message" shortcut

Targets are edited by superusers (or staff with platform.change_salestarget).
"""
from __future__ import annotations

import io
from datetime import timedelta
from decimal import Decimal, InvalidOperation

from django.contrib import admin, messages
from django.contrib.admin.views.decorators import staff_member_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST

from apps.platform.infrastructure.analytics_view import can_view


def _can_edit_plan(user) -> bool:
    return user.is_superuser or user.has_perm("platform.change_salestarget")


@staff_member_required
def sales_plan_view(request):
    from apps.platform.infrastructure import sales_plan as sp

    if not can_view(request.user):
        raise PermissionDenied
    month = sp.parse_month(request.GET.get("m"))
    data = sp.progress(month)
    prev = (month - timedelta(days=1)).replace(day=1)
    nxt = (month.replace(day=28) + timedelta(days=5)).replace(day=1)
    rows = ([data["company"]] if data["company"] else []) + data["branches"]
    return render(request, "admin/platform/sales_plan.html", {
        **admin.site.each_context(request), "title": "Sotuv rejasi", "d": data, "rows": rows, "month": month,
        "prev": prev.strftime("%Y-%m"), "next": nxt.strftime("%Y-%m"), "can_edit": _can_edit_plan(request.user),
        "is_current": month == sp.month_start(),
    })


@staff_member_required
@require_POST
def sales_plan_save_view(request):
    from apps.platform.infrastructure import sales_plan as sp

    if not _can_edit_plan(request.user):
        raise PermissionDenied
    month = sp.parse_month(request.POST.get("month"))
    values = {k[7:]: v for k, v in request.POST.items() if k.startswith("target_")}
    try:
        n = sp.save_targets(month, values, request.user.get_full_name() or request.user.get_username())
    except (ValueError, InvalidOperation):
        messages.error(request, "Reja summasi noto'g'ri — faqat raqam kiriting (masalan 120 000 000).")
    else:
        messages.success(request, f"{month:%m.%Y} uchun reja saqlandi ({n} ta).")
    return redirect(f"{reverse('admin:nz_sales_plan')}?m={month:%Y-%m}")


def _segment_rows(code: str, data: dict, limit: int | None = 200) -> list[dict]:
    from apps.identity.infrastructure.models import User

    ids = data["members"].get(code, [])
    profiles = data["profiles"]
    users = {str(u.id): u for u in User.objects.filter(id__in=ids[: limit or len(ids)]).only("id", "full_name", "phone_number", "created_at", "telegram_user_id")}
    rows = []
    for cid in ids:
        u = users.get(cid)
        if u is None:
            continue
        p = profiles.get(cid) or {}
        rows.append({"id": cid, "name": u.full_name or "—", "phone": u.phone_number, "telegram": bool(u.telegram_user_id), "created": u.created_at,
                     "purchases": p.get("purchases", 0), "revenue": Decimal(p.get("revenue") or 0), "last": p.get("last", ""), "channels": p.get("channels", [])})
    rows.sort(key=lambda r: (-r["revenue"], r["name"]))
    return rows


@staff_member_required
def segments_view(request):
    from apps.platform.infrastructure import segments as sg

    if not can_view(request.user):
        raise PermissionDenied
    data = sg.snapshot(force=request.GET.get("refresh") == "1")
    code = request.GET.get("s") if request.GET.get("s") in sg.SEGMENTS else "at_risk"
    current = data["segments"][code]
    return render(request, "admin/platform/segments.html", {
        **admin.site.each_context(request), "title": "Mijozlar segmentlari", "data": data, "cards": list(data["segments"].values()),
        "code": code, "current": current, "rows": _segment_rows(code, data), "total": data["total_customers"],
        "broadcast_url": f"{reverse('admin:nz_notify_send')}?audience=segment&segment={code}" if _has_compose() else "",
        "export_url": f"{reverse('admin:nz_segments_export')}?s={code}",
    })


def _has_compose() -> bool:
    try:
        reverse("admin:nz_notify_send")
        return True
    except Exception:  # noqa: BLE001
        return False


@staff_member_required
def segments_export_view(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from apps.platform.infrastructure import segments as sg

    if not can_view(request.user):
        raise PermissionDenied
    code = request.GET.get("s") if request.GET.get("s") in sg.SEGMENTS else "at_risk"
    data = sg.snapshot()
    wb = Workbook()
    ws = wb.active
    ws.title = sg.SEGMENTS[code]["label"][:31]
    ws.append(["Mijoz", "Telefon", "Telegram", "Xaridlar", "Jami xarid (so'm)", "Oxirgi xarid", "Kanal"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E293B")
    for r in _segment_rows(code, data, limit=None):
        ws.append([r["name"], r["phone"], "ha" if r["telegram"] else "", r["purchases"], float(r["revenue"]), r["last"][:10], ", ".join(r["channels"])])
    for col, w in zip("ABCDEFG", (28, 16, 10, 10, 18, 14, 14), strict=True):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="segment-{code}.xlsx"'
    return response
