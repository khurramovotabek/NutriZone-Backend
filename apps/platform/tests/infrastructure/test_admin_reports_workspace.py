"""
Phase 08.8: NutriZone Admin Reports & Business Intelligence tests.

Real database fixtures throughout -- every claim about a metric
changing is verified against real persisted data, per this phase's
own explicit "Real Data Verification" requirement.
"""
from __future__ import annotations

import uuid
from datetime import datetime
from decimal import Decimal

import pytest
from django.contrib.auth import get_user_model
from django.test import Client
from django.test.utils import CaptureQueriesContext
from django.db import connection

from apps.ordering.domain.aggregates import Order
from apps.ordering.domain.enums import OrderChannel, OrderStatus
from apps.ordering.domain.value_objects import OrderNumber
from apps.ordering.infrastructure import composition as ordering_composition
from apps.ordering.tests.factories import _next_order_number, make_address, make_line_item
from apps.platform.infrastructure.reports import ReportPeriod, ReportsAggregator
from common.utils.money import Money

pytestmark = [pytest.mark.django_db, pytest.mark.integration]


@pytest.fixture
def staff_client():
    User = get_user_model()
    user = User.objects.create(username=f"nz_staff_{uuid.uuid4().hex[:8]}", is_staff=True, is_superuser=True)
    user.set_password("pw")
    user.save()
    client = Client()
    client.force_login(user)
    return client


@pytest.fixture
def non_staff_client():
    User = get_user_model()
    user = User.objects.create(username=f"nz_customer_{uuid.uuid4().hex[:8]}", is_staff=False)
    user.set_password("pw")
    user.save()
    client = Client()
    client.force_login(user)
    return client


def _make_completed_order(amount=Decimal("100000.00")):
    line = make_line_item(unit_price=Money(amount=amount, currency="UZS"), quantity=1)
    order = Order(
        id=uuid.uuid4(), order_number=OrderNumber(_next_order_number()), customer_id=uuid.uuid4(), cart_id=uuid.uuid4(),
        shipping_address=make_address(), billing_address=make_address(), currency="UZS", channel=OrderChannel.WEB,
        status=OrderStatus.COMPLETED, lines={line.id: line}, placed_at=datetime.utcnow(),
    )
    return ordering_composition.build_order_repository().save(order)


class TestReportsAccessControl:
    def test_anonymous_denied(self):
        response = Client().get("/django-admin/reports/")
        assert response.status_code in (302, 403)

    def test_non_staff_denied(self, non_staff_client):
        response = non_staff_client.get("/django-admin/reports/")
        assert response.status_code in (302, 403)

    def test_staff_allowed(self, staff_client):
        response = staff_client.get("/django-admin/reports/")
        assert response.status_code == 200
        assert "Reports & Business Intelligence" in response.content.decode()


class TestSalesReportRealData:
    def test_empty_database_shows_zero_not_fake_data(self):
        period = ReportPeriod.last_n_days(30, "Last 30 Days")
        result = ReportsAggregator(period).sales_report()
        assert result["gross_sales"] == Decimal("0")
        assert result["orders"] == 0
        assert result["revenue_change"]["value"] is None  # no fabricated percentage

    def test_real_completed_order_changes_gross_sales(self):
        """Proves the report reads real data, not a static shape -- per this phase's own explicit real-data-verification requirement."""
        _make_completed_order(amount=Decimal("150000.00"))

        period = ReportPeriod.last_n_days(30, "Last 30 Days")
        result = ReportsAggregator(period).sales_report()
        assert result["gross_sales"] == Decimal("150000.00")
        assert result["orders"] == 1
        assert result["average_order_value"] == Decimal("150000.00")

    def test_rendered_report_shows_real_revenue_number(self, staff_client):
        """End-to-end: a real order's revenue actually appears in the rendered page."""
        _make_completed_order(amount=Decimal("77000.00"))
        response = staff_client.get("/django-admin/reports/")
        assert response.status_code == 200
        assert "77000" in response.content.decode()


class TestReportsPerformance:
    def test_bounded_query_count_regardless_of_order_volume(self, staff_client):
        """Reports must aggregate in the database, not per-row -- proves query count does not scale with order count."""
        for i in range(5):
            _make_completed_order(amount=Decimal("10000.00"))

        with CaptureQueriesContext(connection) as ctx:
            response = staff_client.get("/django-admin/reports/")
        assert response.status_code == 200
        count_5 = len(ctx.captured_queries)

        for i in range(5, 15):
            _make_completed_order(amount=Decimal("10000.00"))

        with CaptureQueriesContext(connection) as ctx:
            response = staff_client.get("/django-admin/reports/")
        assert response.status_code == 200
        count_15 = len(ctx.captured_queries)

        assert count_15 < count_5 + 5, f"Query count grew from {count_5} to {count_15} for 10 more orders -- Reports is not aggregating in the database"


class TestReportsFinancialIntegrityIsolation:
    """This phase's own explicit rule: Reports must be read-only and must never mutate business state."""

    def test_viewing_reports_does_not_change_order_state(self, staff_client):
        order = _make_completed_order(amount=Decimal("50000.00"))
        staff_client.get("/django-admin/reports/")
        staff_client.get("/django-admin/reports/")

        from apps.ordering.infrastructure import models as orm

        reloaded = orm.OrderModel.objects.get(id=order.id)
        assert reloaded.status == "completed"  # unchanged by merely viewing the report


class TestNoFakeSectionsForUnsupportedConcepts:
    """This phase's own explicit rule: do not fabricate analytics for domain concepts confirmed absent in prior phases."""

    def test_reports_page_does_not_mention_fake_banner_analytics(self, staff_client):
        """Updated: a real Banner model now exists (admin-registered, appears in the sidebar naturally and correctly) -- this no longer asserts "Banner" is absent entirely, only that no fabricated "Banner Analytics"-style reporting section was invented for it."""
        response = staff_client.get("/django-admin/reports/")
        assert "Banner Analytics" not in response.content.decode()

    def test_reports_page_does_not_mention_customer_segments(self, staff_client):
        response = staff_client.get("/django-admin/reports/")
        content = response.content.decode()
        assert "VIP" not in content
        assert "High-value" not in content


class TestCrossWorkspaceRegression:
    def test_dashboard_home_still_renders(self, staff_client):
        response = staff_client.get("/django-admin/")
        assert response.status_code == 200
        assert "nz-dashboard" in response.content.decode()

    def test_commerce_still_renders(self, staff_client):
        response = staff_client.get("/django-admin/products/product/")
        assert response.status_code == 200

    def test_marketing_still_renders(self, staff_client):
        response = staff_client.get("/django-admin/pricing/couponmodel/")
        assert response.status_code == 200


class TestSalesReportExcelExport:
    """Phase 16: the one genuinely missing capability found during reconnaissance -- no Excel export existed anywhere in this project before this phase."""

    def test_non_staff_denied(self, non_staff_client):
        response = non_staff_client.get("/django-admin/reports/export/sales.xlsx")
        assert response.status_code in (302, 403)

    def test_staff_receives_real_xlsx_file(self, staff_client):
        import io

        from openpyxl import load_workbook

        response = staff_client.get("/django-admin/reports/export/sales.xlsx")
        assert response.status_code == 200
        assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "attachment" in response["Content-Disposition"]
        assert ".xlsx" in response["Content-Disposition"]

        # The response body must be a genuinely valid, openable workbook -- not merely a 200 status.
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        assert sheet.title == "Sales Report"
        assert sheet["A1"].value == "NutriZone Sales Report"

    def test_export_reflects_real_persisted_order_data(self, staff_client):
        """Proves the export is not a static template -- it reflects real, current database state, using the exact same ReportsAggregator the HTML report itself uses."""
        import io

        from openpyxl import load_workbook

        order_id = uuid.uuid4()
        order = Order.create_from_cart_data(
            order_id=order_id, order_number=OrderNumber(_next_order_number()), cart_id=None, customer_id=uuid.uuid4(),
            shipping_address=make_address(), billing_address=make_address(), currency="UZS",
            lines=[make_line_item(unit_price=Money(Decimal("77000.00"), "UZS"), quantity=1)],
            coupon_ids=set(), channel=OrderChannel.WEB, at=datetime.now(),
        )[0]
        order.status = OrderStatus.COMPLETED
        ordering_composition.build_order_repository().save(order)

        response = staff_client.get("/django-admin/reports/export/sales.xlsx")
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        values = [row[1] for row in sheet.iter_rows(min_row=1, max_col=2, values_only=True) if row[0] == "Gross Sales (UZS)"]
        assert values and values[0] >= 77000.0


class TestReportSectionExcelExport:
    """Real, closed gap: only the Sales section had an export -- Products, Inventory, Marketing, Loyalty, Wallet, Reviews, and Fulfillment had no way to download their data at all."""

    def test_non_staff_denied(self, non_staff_client):
        response = non_staff_client.get("/django-admin/reports/export/inventory.xlsx")
        assert response.status_code in (302, 403)

    def test_unknown_section_returns_404(self, staff_client):
        response = staff_client.get("/django-admin/reports/export/not-a-real-section.xlsx")
        assert response.status_code == 404

    def test_each_real_section_produces_a_genuinely_valid_workbook(self, staff_client):
        import io

        from openpyxl import load_workbook

        for section in ("products", "inventory", "marketing", "reviews"):
            response = staff_client.get(f"/django-admin/reports/export/{section}.xlsx")
            assert response.status_code == 200, f"{section} export failed"
            assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            assert "attachment" in response["Content-Disposition"]

            workbook = load_workbook(io.BytesIO(response.content))
            sheet = workbook.active
            assert sheet.max_row > 4, f"{section} export has no real data rows"

    def test_inventory_export_reflects_a_real_low_stock_item(self, staff_client):
        """Proves the export is not a static template -- a WMS product below its reorder minimum shows up (the same source as the dashboard home)."""
        import io

        from openpyxl import load_workbook

        from apps.stores.infrastructure.models import ReorderRule
        from apps.stores.tests.interfaces.test_wms_planning import _put, _store, _variant

        shop, variant = _store(), _variant()
        _put(shop, variant, 2)
        ReorderRule.objects.create(store=shop, product_variant_id=variant.id, min_quantity=10, target_quantity=20)

        response = staff_client.get("/django-admin/reports/export/inventory.xlsx")
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        rows = list(sheet.iter_rows(min_row=1, max_col=2, values_only=True))
        low_stock_values = [row[1] for row in rows if row[0] and "low_stock_total" in str(row[0])]
        assert low_stock_values and low_stock_values[0] >= 1


class TestWalletAndOrderStatusVisibility:
    """
    Real, live-verified regression test for a confirmed gap:
    order_status_distribution (part of sales_report()) was only ever
    exposed via the Excel export, with no visual chart in the HTML
    report itself.
    """

    def test_order_status_chart_canvas_and_data_are_present(self, staff_client):
        order_id = uuid.uuid4()
        order = Order.create_from_cart_data(
            order_id=order_id, order_number=OrderNumber(_next_order_number()), cart_id=None, customer_id=uuid.uuid4(),
            shipping_address=make_address(), billing_address=make_address(), currency="UZS",
            lines=[make_line_item(unit_price=Money(Decimal("30000.00"), "UZS"), quantity=1)],
            coupon_ids=set(), channel=OrderChannel.WEB, at=datetime.now(),
        )[0]
        ordering_composition.build_order_repository().save(order)

        response = staff_client.get("/django-admin/reports/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "nz-order-status-chart" in content
        assert "nz-order-status-series-data" in content

    def test_order_status_chart_data_reflects_a_real_persisted_order(self, staff_client):
        import json

        order_id = uuid.uuid4()
        order = Order.create_from_cart_data(
            order_id=order_id, order_number=OrderNumber(_next_order_number()), cart_id=None, customer_id=uuid.uuid4(),
            shipping_address=make_address(), billing_address=make_address(), currency="UZS",
            lines=[make_line_item(unit_price=Money(Decimal("50000.00"), "UZS"), quantity=1)],
            coupon_ids=set(), channel=OrderChannel.WEB, at=datetime.now(),
        )[0]
        order.status = OrderStatus.COMPLETED
        ordering_composition.build_order_repository().save(order)

        response = staff_client.get("/django-admin/reports/")
        content = response.content.decode()
        start = content.index('id="nz-order-status-series-data"')
        json_start = content.index(">", start) + 1
        json_end = content.index("</script>", json_start)
        distribution = json.loads(content[json_start:json_end])
        assert distribution.get("completed", 0) >= 1

    def test_top_products_chart_canvas_and_data_reflect_a_real_order(self, staff_client):
        """Real, closed gap: Top Products previously showed only a plain ranked text list -- now a real bar chart, built from the same, already-certified sales-report data (no separate/fabricated data source)."""
        import json

        from apps.products.infrastructure import models as products_orm

        product = products_orm.Product.objects.create(id=uuid.uuid4(), status="published")
        products_orm.ProductTranslation.objects.create(id=uuid.uuid4(), product=product, locale_code="en", name="Top Chart Product", slug=f"tcp-{uuid.uuid4().hex[:8]}")
        variant = products_orm.ProductVariant.objects.create(id=uuid.uuid4(), product=product, sku=f"SKU-{uuid.uuid4().hex[:8].upper()}", is_default=True)

        line = make_line_item(unit_price=Money(amount=Decimal("75000.00"), currency="UZS"), quantity=3, product_variant_id=variant.id)
        order = Order(
            id=uuid.uuid4(), order_number=OrderNumber(_next_order_number()), customer_id=uuid.uuid4(), cart_id=uuid.uuid4(),
            shipping_address=make_address(), billing_address=make_address(), currency="UZS", channel=OrderChannel.WEB,
            status=OrderStatus.COMPLETED, lines={line.id: line}, placed_at=datetime.utcnow(),
        )
        ordering_composition.build_order_repository().save(order)

        response = staff_client.get("/django-admin/reports/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "nz-top-products-chart" in content
        assert "nz-top-products-series-data" in content

        start = content.index('id="nz-top-products-series-data"')
        json_start = content.index(">", start) + 1
        json_end = content.index("</script>", json_start)
        top_products = json.loads(content[json_start:json_end])
        assert any(p["name"] == "Top Chart Product" and p["quantity_sold"] == 3 for p in top_products)
