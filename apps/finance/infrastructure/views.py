"""Moliya pages: overview, P&L, cash flow, budget, debts, account statement (+ Excel). URLs: install() below."""
from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.db.models import Sum
from django.http import HttpResponseRedirect
from django.shortcuts import get_object_or_404
from django.template.response import TemplateResponse
from django.urls import reverse
from django.utils import timezone

from apps.finance.infrastructure import ledger, pnl
from apps.finance.infrastructure.models import (
    SOURCES, BudgetLine, Category, CategoryKind, MoneyAccount, PnlSection, Transaction, TxKind,
)

ZERO = Decimal(0)
TABS = (
    ("overview", "Moliya paneli", "fa-gauge-high", "admin:nz_finance"),
    ("transactions", "Operatsiyalar", "fa-right-left", "admin:finance_transaction_changelist"),
    ("accounts", "Hisoblar", "fa-wallet", "admin:finance_moneyaccount_changelist"),
    ("pnl", "Foyda va zarar", "fa-scale-balanced", "admin:nz_finance_pnl"),
    ("cashflow", "Pul oqimi", "fa-water", "admin:nz_finance_cashflow"),
    ("budget", "Budjet", "fa-bullseye", "admin:nz_finance_budget"),
    ("debts", "Qarzlar", "fa-file-invoice-dollar", "admin:nz_finance_debts"),
    ("categories", "Moddalar", "fa-tags", "admin:finance_category_changelist"),
)


def money(v) -> str:
    if v is None:
        return "—"
    v = Decimal(v).quantize(Decimal(1))
    s = f"{abs(v):,.0f}".replace(",", " ")
    return f"−{s}" if v < 0 else s


def tr(text: str) -> str:
    """A dashboard phrase in the staff member's language, for text that never passes through the HTML translator (chart JSON)."""
    from django.utils.translation import get_language

    from apps.platform.infrastructure.admin_i18n import translate_text

    lang = (get_language() or "uz")[:2]
    if lang == "uz" or not text:
        return text
    return translate_text(text, lang) or text


UNITS = {"uz": ("ming", "mln", "mlrd"), "ru": ("тыс", "млн", "млрд"), "en": ("K", "M", "B")}


def chart_units() -> dict:
    from django.utils.translation import get_language

    k, m, b = UNITS.get((get_language() or "uz")[:2], UNITS["uz"])
    return {"k": k, "m": m, "b": b}


def tabs(active: str) -> list[dict]:
    return [{"key": k, "label": label, "icon": icon, "url": reverse(name), "on": k == active} for k, label, icon, name in TABS]


def _check(request, perm: str = "finance.view_transaction"):
    if not request.user.is_active or not request.user.is_staff or not request.user.has_perm(perm):
        raise PermissionDenied


def _ctx(request, active: str, title: str, **extra) -> dict:
    return {**admin.site.each_context(request), "title": title, "nz_tabs": tabs(active), "nz_active": active,
            "can_add": request.user.has_perm("finance.add_transaction"), **extra}


def _months_param(request, default: int = 6) -> tuple[int, list[date]]:
    try:
        n = int(request.GET.get("months") or default)
    except ValueError:
        n = default
    n = n if n in (3, 6, 12) else default
    return n, pnl.months_back(timezone.localdate(), n)


# ---------------------------------------------------------------- overview


def overview_view(request):
    _check(request)
    today = timezone.localdate()
    m0 = today.replace(day=1)
    accounts = list(MoneyAccount.objects.filter(is_active=True).select_related("store"))
    bal = ledger.balances(today, accounts)
    total = sum((b["balance"] for b in bal.values()), ZERO)
    cards = [{"a": a, "balance": money(bal[a.id]["balance"]), "neg": bal[a.id]["balance"] < 0, "in": money(bal[a.id]["in_month"]), "out": money(bal[a.id]["out_month"]),
              "url": reverse("admin:nz_finance_statement", args=[a.id])} for a in accounts]
    six = pnl.months_back(today, 6)
    report = pnl.pnl(six)
    cur = len(six) - 1
    flow = pnl.cashflow(six)
    # money the system took/paid but no account is set up to receive -- shown so balances can be trusted
    first = min((a.opening_date for a in accounts), default=m0)
    _routed, orphans = ledger.route(ledger.auto_flows(max(first, m0), today), MoneyAccount.objects.filter(is_active=True))
    orphan_by_source: dict = {}
    for f in orphans:
        orphan_by_source[f.source] = orphan_by_source.get(f.source, ZERO) + f.amount
    expense_rows = [r for r in report["rows"] if r["kind"] == "line" and r["key"] not in ("pos", "online", "coins", "cogs") and r["values"][cur]]
    structure = sorted(({"label": r["label"], "value": float(-r["values"][cur]), "color": r["color"]} for r in expense_rows if r["values"][cur] < 0), key=lambda x: -x["value"])
    debts = _debt_summary()
    budget = _budget_rows(m0)
    over = [b for b in budget if b["plan"] and b["actual"] > b["plan_raw"]]
    recent = Transaction.objects.select_related("account", "to_account", "category").order_by("-occurred_on", "-created_at")[:8]
    month_in = flow["total_in"][cur]
    month_out = flow["total_out"][cur]
    context = _ctx(
        request, "overview", "Moliya paneli",
        total=money(total), cards=cards, month_in=money(month_in), month_out=money(month_out), month_net=money(month_in - month_out),
        profit=money(report["net"][cur]), profit_neg=report["net"][cur] < 0, margin=report["margins"][cur], gross_margin=report["gross_margins"][cur],
        revenue=money(report["revenue"][cur]), coverage=report["months"][cur]["coverage"],
        chart={"labels": [m["label"] for m in flow["months"]], "in": [float(v) for v in flow["total_in"]], "out": [float(v) for v in flow["total_out"]],
               "profit": [float(v) for v in report["net"]], "structure": [{**x, "label": tr(x["label"])} for x in structure],
               "t": {"in": tr("Kirim"), "out": tr("Chiqim"), "profit": tr("Sof foyda"), "som": tr("so'm"), **chart_units()}},
        orphans=[{"label": str(SOURCES[s][0]), "amount": money(v)} for s, v in orphan_by_source.items() if v],
        debts=debts, over_budget=over, recent=recent, month_label=pnl.month_label(m0), today=today,
        negative=[c for c in cards if c["neg"]], no_accounts=not accounts,
        add_url=reverse("admin:finance_transaction_add"),
    )
    return TemplateResponse(request, "admin/finance/overview.html", context)


# ---------------------------------------------------------------- P&L and cash flow


def _fmt_rows(rows):
    for r in rows:
        if "values" in r:
            r["cells"] = [{"v": money(v), "neg": v < 0} for v in r["values"]]
            r["total_fmt"] = money(r["total"])
            r["total_neg"] = r["total"] < 0
    return rows


def pnl_view(request):
    _check(request)
    n, months = _months_param(request)
    report = pnl.pnl(months)
    if request.GET.get("export") == "xlsx":
        return _pnl_excel(report)
    _fmt_rows(report["rows"])
    return TemplateResponse(request, "admin/finance/pnl.html", _ctx(
        request, "pnl", "Foyda va zarar", r=report, n=n,
        total_margin=round(float(sum(report["net"], ZERO) / sum(report["revenue"], ZERO) * 100), 1) if sum(report["revenue"], ZERO) else None,
        chart={"labels": [m["label"] for m in report["months"]], "revenue": [float(v) for v in report["revenue"]], "gross": [float(v) for v in report["gross"]],
               "net": [float(v) for v in report["net"]],
               "t": {"revenue": tr("Sof tushum"), "gross": tr("Yalpi foyda"), "net": tr("Sof foyda"), "som": tr("so'm"), **chart_units()}},
    ))


def cashflow_view(request):
    _check(request)
    n, months = _months_param(request)
    report = pnl.cashflow(months)
    if request.GET.get("export") == "xlsx":
        return _cashflow_excel(report)
    _fmt_rows(report["inflow"])
    _fmt_rows(report["outflow"])
    for key in ("total_in", "total_out", "net", "closing"):
        report[f"{key}_fmt"] = [{"v": money(v), "neg": v < 0} for v in report[key]]
    report["sum_in"], report["sum_out"] = money(sum(report["total_in"], ZERO)), money(sum(report["total_out"], ZERO))
    report["sum_net"] = money(sum(report["net"], ZERO))
    return TemplateResponse(request, "admin/finance/cashflow.html", _ctx(request, "cashflow", "Pul oqimi", r=report, n=n))


def _xlsx_rows(title, months, blocks):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"NutriZone · {timezone.localtime():%d.%m.%Y %H:%M}"])
    ws.append([])
    ws.append(["Modda", *months, "Jami"])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E293B")
    for label, values, total, bold in blocks:
        ws.append([label, *[float(v) for v in values], float(total) if total is not None else None])
        if bold:
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, min_col=2):
        for c in row:
            c.number_format = "# ##0"
            c.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 42
    for i in range(len(months) + 1):
        ws.column_dimensions[chr(66 + i)].width = 16
    from io import BytesIO

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pnl_excel(report):
    from apps.platform.infrastructure.excel_export import response

    blocks = []
    for r in report["rows"]:
        if r["kind"] == "head":
            blocks.append((r["label"], [], None, True))
        else:
            blocks.append((r["label"], r["values"], r["total"], r["kind"] in ("total", "grand", "sub")))
    return response(_xlsx_rows("Foyda va zarar", [m["label"] for m in report["months"]], blocks), "foyda-va-zarar")


def _cashflow_excel(report):
    from apps.platform.infrastructure.excel_export import response

    blocks = [("KIRIM", [], None, True)] + [(r["label"], r["values"], r["total"], False) for r in report["inflow"]]
    blocks += [("Jami kirim", report["total_in"], sum(report["total_in"], ZERO), True), ("CHIQIM", [], None, True)]
    blocks += [(r["label"], r["values"], r["total"], False) for r in report["outflow"]]
    blocks += [("Jami chiqim", report["total_out"], sum(report["total_out"], ZERO), True), ("Sof pul oqimi", report["net"], sum(report["net"], ZERO), True),
               ("Oy oxiridagi qoldiq", report["closing"], None, True)]
    return response(_xlsx_rows("Pul oqimi", [m["label"] for m in report["months"]], blocks), "pul-oqimi")


# ---------------------------------------------------------------- budget


def _budget_rows(month: date) -> list[dict]:
    end = min(pnl.next_month(month) - timedelta(days=1), timezone.localdate())
    actual = ledger.manual_totals(month, end) if month <= timezone.localdate() else {}
    plans = {b.category_id: b.amount for b in BudgetLine.objects.filter(month=month)}
    out = []
    for c in Category.objects.filter(kind=CategoryKind.EXPENSE, is_active=True).exclude(section=PnlSection.CAPITAL):
        a = actual.get((TxKind.EXPENSE, c.id), ZERO)
        p = plans.get(c.id)
        pct = int(round(a / p * 100)) if p else None
        out.append({"c": c, "actual": a, "actual_fmt": money(a), "plan_raw": p, "plan": f"{p:.0f}" if p is not None else "", "plan_fmt": money(p) if p is not None else "—",
                    "pct": pct, "bar": min(pct or 0, 100), "state": "over" if pct and pct > 100 else "warn" if pct and pct >= 85 else "ok", "left": money(p - a) if p is not None else ""})
    return out


def budget_view(request):
    _check(request)
    today = timezone.localdate()
    try:
        month = date.fromisoformat((request.GET.get("month") or request.POST.get("month") or "") + "-01")
    except ValueError:
        month = today.replace(day=1)
    if request.method == "POST":
        _check(request, "finance.add_budgetline")
        if request.POST.get("copy_prev"):
            prev = pnl.months_back(month, 2)[0]
            n = 0
            for b in BudgetLine.objects.filter(month=prev):
                _o, created = BudgetLine.objects.get_or_create(month=month, category_id=b.category_id, defaults={"amount": b.amount})
                n += created
            messages.success(request, f"O'tgan oydan {n} ta reja ko'chirildi.")
        else:
            saved = 0
            for c in Category.objects.filter(kind=CategoryKind.EXPENSE):
                raw = (request.POST.get(f"plan_{c.id}") or "").replace(" ", "").strip()
                if raw == "":
                    BudgetLine.objects.filter(month=month, category=c).delete()
                    continue
                try:
                    amount = Decimal(raw)
                except InvalidOperation:
                    messages.error(request, f"{c.name}: summa noto'g'ri.")
                    continue
                if amount < 0:
                    messages.error(request, f"{c.name}: summa manfiy bo'lmaydi.")
                    continue
                BudgetLine.objects.update_or_create(month=month, category=c, defaults={"amount": amount})
                saved += 1
            messages.success(request, f"Budjet saqlandi ({saved} ta modda).")
        return HttpResponseRedirect(f"{reverse('admin:nz_finance_budget')}?month={month:%Y-%m}")
    rows = _budget_rows(month)
    plan_total = sum((r["plan_raw"] or ZERO for r in rows), ZERO)
    actual_total = sum((r["actual"] for r in rows), ZERO)
    prev, nxt = pnl.months_back(month, 2)[0], pnl.next_month(month)
    return TemplateResponse(request, "admin/finance/budget.html", _ctx(
        request, "budget", "Budjet", rows=rows, month=month, month_label=pnl.month_label(month), prev=prev, next=nxt,
        plan_total=money(plan_total), actual_total=money(actual_total), left_total=money(plan_total - actual_total),
        pct_total=int(round(actual_total / plan_total * 100)) if plan_total else None,
        can_edit=request.user.has_perm("finance.add_budgetline"),
    ))


# ---------------------------------------------------------------- debts


def _debt_summary() -> dict:
    from apps.loyalty.infrastructure.models import Wallet
    from apps.ordering.infrastructure.models import OrderLineItemModel, OrderModel
    from apps.payment.infrastructure.models import PaymentModel
    from apps.stores.infrastructure import supplier_service
    from apps.stores.infrastructure.models import Supplier

    bal = supplier_service.balances()
    names = dict(Supplier.objects.filter(id__in=list(bal)).values_list("id", "name"))
    suppliers = []
    for sid, b in bal.items():
        due = b.purchased - b.paid
        if due > 0 or b.overdue > 0:
            suppliers.append({"id": sid, "name": names.get(sid, "—"), "due": due, "due_fmt": money(due), "overdue": b.overdue, "overdue_fmt": money(b.overdue),
                              "oldest": b.oldest_overdue, "url": reverse("admin:stores_supplier_change", args=[sid])})
    suppliers.sort(key=lambda s: (-s["overdue"], -s["due"]))
    road = list(OrderModel.objects.filter(status="shipped", is_archived=False).values_list("id", flat=True))
    paid = set(PaymentModel.objects.filter(order_id__in=road, status="succeeded").exclude(provider="cash").values_list("order_id", flat=True))
    cod = [o for o in road if o not in paid]
    courier = ZERO
    for l in OrderLineItemModel.objects.filter(order_id__in=cod, status="active").values("unit_price_amount", "quantity", "discount_amount", "tax_amount"):
        courier += l["unit_price_amount"] * l["quantity"] - l["discount_amount"] + l["tax_amount"]
    coins = Wallet.objects.aggregate(s=Sum("balance"))["s"] or 0
    payable = sum((s["due"] for s in suppliers), ZERO)
    overdue = sum((s["overdue"] for s in suppliers), ZERO)
    return {"suppliers": suppliers, "payable": money(payable), "payable_raw": payable, "overdue": money(overdue), "overdue_raw": overdue,
            "courier": money(courier), "courier_orders": len(cod), "coins": money(coins), "coins_raw": Decimal(coins)}


def debts_view(request):
    _check(request)
    return TemplateResponse(request, "admin/finance/debts.html", _ctx(request, "debts", "Qarzlar", d=_debt_summary()))


# ---------------------------------------------------------------- account statement


def statement_view(request, pk):
    _check(request)
    account = get_object_or_404(MoneyAccount, pk=pk)
    today = timezone.localdate()
    try:
        start = date.fromisoformat(request.GET.get("from") or "")
    except ValueError:
        start = today.replace(day=1)
    try:
        end = date.fromisoformat(request.GET.get("to") or "")
    except ValueError:
        end = today
    end = min(end, today)
    start = min(start, end)
    st = ledger.statement(account, start, end)
    if request.GET.get("export") == "xlsx":
        from apps.platform.infrastructure.excel_export import Col, response, workbook

        cols = [Col("Sana", "day", 12), Col("Operatsiya", "title", 34), Col("Tafsilot", "detail", 40), Col("Summa", "amount", 16, money=True), Col("Qoldiq", "balance", 16)]
        rows = [type("R", (), r) for r in reversed(st["rows"])]
        return response(workbook(f"{account.name} — ko'chirma".replace(":", " ")[:31], cols, rows, f"{st['start']:%d.%m.%Y} — {st['end']:%d.%m.%Y}"), "hisob-kochirma")
    for r in st["rows"]:
        r["amount_fmt"], r["balance_fmt"] = money(r["amount"]), money(r["balance"])
    return TemplateResponse(request, "admin/finance/statement.html", _ctx(
        request, "accounts", account.name, a=account, st=st, opening=money(st["opening"]), closing=money(st["closing"]), inn=money(st["in"]), out=money(st["out"]),
        start=start, end=end,
    ))


# ---------------------------------------------------------------- void an operation


def void_view(request, pk):
    _check(request, "finance.void_transaction")
    tx = get_object_or_404(Transaction, pk=pk)
    back = reverse("admin:finance_transaction_change", args=[tx.pk])
    if request.method != "POST":
        return HttpResponseRedirect(back)
    reason = (request.POST.get("reason") or "").strip()[:255]
    if tx.voided_at:
        messages.info(request, "Operatsiya allaqachon bekor qilingan.")
    elif not reason:
        messages.error(request, "Bekor qilish sababini yozing.")
    else:
        tx.voided_at = timezone.now()
        tx.voided_by = request.user.get_full_name() or request.user.get_username()
        tx.void_reason = reason
        tx.save(update_fields=["voided_at", "voided_by", "void_reason"])
        messages.success(request, f"{tx.number} bekor qilindi — endi hech qayerda hisoblanmaydi.")
    return HttpResponseRedirect(back)


def urls():
    from django.urls import path

    v = admin.site.admin_view
    return [
        path("finance/overview/", v(overview_view), name="nz_finance"),
        path("finance/pnl/", v(pnl_view), name="nz_finance_pnl"),
        path("finance/cashflow/", v(cashflow_view), name="nz_finance_cashflow"),
        path("finance/budget/", v(budget_view), name="nz_finance_budget"),
        path("finance/debts/", v(debts_view), name="nz_finance_debts"),
        path("finance/account/<uuid:pk>/", v(statement_view), name="nz_finance_statement"),
        path("finance/operation/<uuid:pk>/void/", v(void_view), name="nz_finance_void"),
    ]


def install() -> None:
    """Prepends the finance pages to the admin site's URLs (before the `finance/` app index catch-all)."""
    if getattr(admin.site, "_nz_finance_installed", False):
        return
    previous = admin.site.get_urls

    def get_urls():
        return urls() + previous()

    admin.site.get_urls = get_urls
    admin.site._nz_finance_installed = True
