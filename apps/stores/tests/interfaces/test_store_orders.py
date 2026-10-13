"""
Branch online orders: pickup at checkout, the dispatcher handing a delivery to a branch, the branch owner moving it
(stock leaves THAT branch), owner / customer notifications, and the dashboard kassa PIN errors being shown.
"""
from __future__ import annotations

import uuid
from decimal import Decimal

import pytest
from django.contrib.auth.models import User as AuthUser
from django.test import Client

from apps.cart.tests.interfaces.test_checkout_order_creation import _client_as, _make_address, _make_geo, _make_user, _priced_cart
from apps.products.tests.stock_helpers import location, put_stock
from apps.stores.tests.interfaces.test_my_store_products import _make_variant

pytestmark = [pytest.mark.django_db(transaction=True), pytest.mark.integration]  # owner / customer alerts go out on commit


def _priced_variant(amount=Decimal("40000.00")):
    """A real catalogue variant (stock rows need one) with an active price."""
    from apps.pricing.infrastructure import models as pricing_orm

    variant = _make_variant(barcode=f"479{uuid.uuid4().int % 10**10:010d}", status="active")
    price_list, _ = pricing_orm.PriceListModel.objects.get_or_create(code="DEFAULT-UZS", defaults={"name": "Default", "currency": "UZS", "status": "active"})
    pricing_orm.PriceListEntryModel.objects.create(price_list=price_list, product_variant_id=variant.id, base_price_amount=amount)
    return variant.id


def _owner_store(name="Chilonzor", lat=41.2856, lng=69.2034):
    from apps.identity.infrastructure.models import User

    owner = User.objects.create(phone_number=f"+99893{uuid.uuid4().int % 10_000_000:07d}", full_name="Do'kon Egasi", status="active")
    store = location(name)
    store.owner = owner
    store.latitude, store.longitude, store.address = lat, lng, "Bunyodkor 1"
    store.save()
    return owner, store


def _checkout(client, cart, **body):
    return client.post(f"/api/v1/cart/{cart['code']}/checkout", body, format="json", HTTP_X_CART_TOKEN=cart["token"], HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4()))


class TestPickupCheckout:
    def test_pickup_goes_to_the_branch_and_tells_its_owner(self, monkeypatch):
        from apps.notifications.infrastructure.models import NotificationModel
        from apps.ordering.infrastructure.models import OrderModel

        owner, store = _owner_store()
        variant = _priced_variant(Decimal("50000.00"))
        put_stock(variant, 5, store)
        user_id = uuid.uuid4()
        _make_user(user_id, 31)
        client = _client_as(monkeypatch, user_id)
        cart, *_ = _priced_cart(client, variant, quantity=2)
        r = _checkout(client, cart, fulfilment_method="pickup", pickup_store_id=str(store.id))
        assert r.status_code == 201, r.json()
        order = OrderModel.objects.get(pk=r.json()["data"]["id"])
        assert order.fulfilment_method == "pickup" and order.fulfilment_store_id == store.id and order.shipping_line1.startswith("Olib ketish:")
        note = NotificationModel.objects.get(recipient_id=owner.id, kind="store_order_new")
        assert order.order_number in note.title and note.link.endswith(order.order_number)

    def test_pickup_refused_when_the_branch_lacks_stock_and_inputs_are_checked(self, monkeypatch):
        _owner, store = _owner_store()
        variant = _priced_variant()
        put_stock(variant, 1, store)
        user_id = uuid.uuid4()
        _make_user(user_id, 32)
        client = _client_as(monkeypatch, user_id)
        cart, *_ = _priced_cart(client, variant, quantity=3)
        r = _checkout(client, cart, fulfilment_method="pickup", pickup_store_id=str(store.id))
        assert r.status_code == 409 and r.json()["error"]["code"] == "PICKUP_STORE_OUT_OF_STOCK"
        assert _checkout(client, cart, fulfilment_method="pickup").status_code == 422  # no store
        assert _checkout(client, cart).status_code == 422  # delivery without an address

    def test_options_rank_in_stock_then_nearest(self):
        from apps.stores.infrastructure.store_orders import pickup_options

        variant = _priced_variant()
        _o1, far_full = _owner_store("Uzoq", 41.36, 69.33)
        _o2, near_full = _owner_store("Yaqin", 41.312, 69.28)
        _o3, near_empty = _owner_store("Bo'sh", 41.3112, 69.2798)
        put_stock(variant, 4, far_full)
        put_stock(variant, 4, near_full)
        names = [o["name"] for o in pickup_options([(variant, 2)], 41.3111, 69.2797) if o["id"] in {str(far_full.id), str(near_full.id), str(near_empty.id)}]
        assert names == [near_full.name, far_full.name, near_empty.name]


class TestBranchOwnerMovesTheOrder:
    def _pickup_order(self, monkeypatch):
        from apps.ordering.infrastructure.models import OrderModel

        owner, store = _owner_store()
        variant = _priced_variant(Decimal("30000.00"))
        put_stock(variant, 5, store)
        user_id = uuid.uuid4()
        _make_user(user_id, 33)
        client = _client_as(monkeypatch, user_id)
        cart, *_ = _priced_cart(client, variant, quantity=2)
        r = _checkout(client, cart, fulfilment_method="pickup", pickup_store_id=str(store.id))
        return owner, store, variant, user_id, OrderModel.objects.get(pk=r.json()["data"]["id"])

    def test_accept_pack_ready_picked_up(self, monkeypatch):
        from apps.notifications.infrastructure.models import NotificationModel
        from apps.stores.infrastructure.models import StoreProductStock

        owner, store, variant, customer, order = self._pickup_order(monkeypatch)
        c = _client_as(monkeypatch, owner.id)
        listed = c.get("/api/v1/my-store/orders").json()["data"]
        assert [o["number"] for o in listed["orders"]] == [order.order_number] and listed["orders"][0]["steps"][0]["action"] == "confirm"
        for action, status in (("confirm", "confirmed"), ("start_processing", "processing"), ("ship", "shipped"), ("deliver", "delivered")):
            r = c.post(f"/api/v1/my-store/orders/{order.order_number}/action", {"action": action}, format="json")
            assert r.status_code == 200, r.json()
            assert r.json()["data"]["status"] == status
        stock = StoreProductStock.objects.get(store=store, product_variant_id=variant)
        assert stock.quantity == 3 and stock.reserved_quantity == 0  # sold from this branch, nothing left reserved
        kinds = set(NotificationModel.objects.filter(recipient_id=customer).values_list("kind", flat=True))
        assert {"order_ready_pickup", "order_picked_up"} <= kinds and "order_shipped" not in kinds

    def test_other_branches_and_cashiers_are_limited(self, monkeypatch):
        from apps.stores.infrastructure.models import StoreCashier

        owner, store, _v, _c, order = self._pickup_order(monkeypatch)
        stranger, _other = _owner_store("Boshqa")
        r = _client_as(monkeypatch, stranger.id).post(f"/api/v1/my-store/orders/{order.order_number}/action", {"action": "confirm"}, format="json")
        assert r.status_code == 409 and r.json()["error"]["code"] == "ORDER_NOT_FOUND"
        cashier, _s = _owner_store("Kassir")
        _s.owner = None
        _s.save()
        StoreCashier.objects.create(store=store, user=cashier, is_active=True)
        r = _client_as(monkeypatch, cashier.id).post(f"/api/v1/my-store/orders/{order.order_number}/action", {"action": "hand_back"}, format="json")
        assert r.status_code == 403
        r = _client_as(monkeypatch, owner.id).post(f"/api/v1/my-store/orders/{order.order_number}/action", {"action": "deliver"}, format="json")
        assert r.status_code == 409 and r.json()["error"]["code"] == "STEP_NOT_ALLOWED"

    def test_hand_back_returns_it_to_the_warehouse(self, monkeypatch):
        owner, _store, _v, _c, order = self._pickup_order(monkeypatch)
        r = _client_as(monkeypatch, owner.id).post(f"/api/v1/my-store/orders/{order.order_number}/action", {"action": "hand_back", "reason": "Mahsulot shikastlangan"}, format="json")
        assert r.status_code == 200
        order.refresh_from_db()
        assert order.fulfilment_store_id is None and "omborga qaytardi" in order.note


class TestDispatcherAssigns:
    def test_assign_moves_the_reservation_and_notifies(self, monkeypatch):
        from apps.notifications.infrastructure.models import NotificationModel
        from apps.ordering.infrastructure.models import OrderModel
        from apps.stores.infrastructure.models import StoreProductStock

        warehouse = location("Sklad", kind="warehouse")
        owner, branch = _owner_store()
        variant = _priced_variant()
        put_stock(variant, 10, warehouse)
        put_stock(variant, 10, branch)
        country, city = _make_geo(city_code="ASGCITY1")
        user_id = uuid.uuid4()
        _make_user(user_id, 34)
        address = _make_address(user_id, country, city, latitude=41.30, longitude=69.25)
        client = _client_as(monkeypatch, user_id)
        cart, *_ = _priced_cart(client, variant, quantity=2)
        order = OrderModel.objects.get(pk=_checkout(client, cart, shipping_address_id=str(address.id)).json()["data"]["id"])
        staff = Client()
        staff.force_login(AuthUser.objects.create_superuser(username=f"d_{uuid.uuid4().hex[:6]}", password="x-pass-1234"))
        staff.post(f"/django-admin/ordering/ordermodel/{order.pk}/act/", {"action": "confirm"})
        assert StoreProductStock.objects.get(store=warehouse, product_variant_id=variant).reserved_quantity == 2
        r = staff.post(f"/django-admin/ordering/ordermodel/{order.pk}/act/", {"action": "assign_store", "assign_store_id": str(branch.id)})
        assert r.status_code == 302
        assert StoreProductStock.objects.get(store=warehouse, product_variant_id=variant).reserved_quantity == 0
        assert StoreProductStock.objects.get(store=branch, product_variant_id=variant).reserved_quantity == 2
        assert NotificationModel.objects.filter(recipient_id=owner.id, kind="store_order_new").exists()
        html = staff.get(f"/django-admin/ordering/ordermodel/{order.pk}/change/").content.decode()
        assert "Do'konga biriktirish" in html and branch.name in html and "maps.apple.com" in html


class TestKassaPinErrors:
    def test_rejected_pin_comes_back_with_its_reason(self):
        store = location("Kassa")
        staff = Client()
        staff.force_login(AuthUser.objects.create_superuser(username=f"p_{uuid.uuid4().hex[:6]}", password="x-pass-1234"))
        r = staff.post(f"/django-admin/stores/storecashier/pin/{store.id}/", {"pin": "1111", "pin_again": "1111"})
        assert "pin_error=PIN_TOO_SIMPLE" in r["Location"] and f"pin_kassa={store.id}" in r["Location"]
        r = staff.post(f"/django-admin/stores/storecashier/pin/{store.id}/", {"pin": "5821", "pin_again": "5822"})
        assert "pin_error=PIN_MISMATCH" in r["Location"]
        page = staff.get(f"/django-admin/stores/storecashier/?pin_error=PIN_MISMATCH&pin_kassa={store.id}").content.decode()
        assert "Ikki PIN bir xil emas" in page
        r = staff.post(f"/django-admin/stores/storecashier/pin/{store.id}/", {"pin": "5821", "pin_again": "5821"})
        assert f"pin_ok={store.id}" in r["Location"]
        store.refresh_from_db()
        assert store.pos_pin_hash


class TestBranchReport:
    def test_counts_speed_and_hand_backs(self, monkeypatch):
        from apps.stores.infrastructure.branch_report import branch_report

        owner, store = _owner_store("Hisobot")
        variant = _priced_variant(Decimal("50000.00"))
        put_stock(variant, 10, store)
        user_id = uuid.uuid4()
        _make_user(user_id, 41)
        client = _client_as(monkeypatch, user_id)
        numbers = []
        for _ in range(2):
            cart, *_ = _priced_cart(client, variant, quantity=1)
            numbers.append(_checkout(client, cart, fulfilment_method="pickup", pickup_store_id=str(store.id)).json()["data"]["order_number"])
            client = _client_as(monkeypatch, user_id)
        c = _client_as(monkeypatch, owner.id)
        for action in ("confirm", "start_processing", "ship", "deliver"):
            assert c.post(f"/api/v1/my-store/orders/{numbers[0]}/action", {"action": action}, format="json").status_code == 200
        assert c.post(f"/api/v1/my-store/orders/{numbers[1]}/action", {"action": "hand_back", "reason": "Qadoq shikastlangan"}, format="json").status_code == 200
        row = next(r for r in branch_report(7)["rows"] if r["store"].id == store.id)
        assert row["handled"] == 2 and row["done"] == 1 and row["handed_back"] == 1 and row["handed_back_pct"] == 50
        assert row["accept_median"] is not None and row["ready_median"] is not None and row["revenue"] == Decimal("50000.00")
        staff = Client()
        staff.force_login(AuthUser.objects.create_superuser(username=f"r_{uuid.uuid4().hex[:6]}", password="x-pass-1234"))
        html = staff.get("/django-admin/stores/branch-report/?days=7").content.decode()
        assert "Filial hisoboti" in html and store.name in html and "Qadoq shikastlangan" in html
        # Excel export: the same numbers, four sheets
        from io import BytesIO

        from openpyxl import load_workbook
        resp = staff.get("/django-admin/stores/branch-report/export.xlsx?days=7")
        assert resp.status_code == 200 and "filial-hisoboti-7kun" in resp["Content-Disposition"]
        wb = load_workbook(BytesIO(resp.content))
        assert wb.sheetnames == ["Filiallar", "Buyurtmalar", "Kunlar", "Omborga qaytarishlar"]
        data_row = next(r for r in wb["Filiallar"].iter_rows(values_only=True) if r[0] == store.name)
        assert data_row[1] == 2 and data_row[7] == 1 and data_row[12] == 1 and data_row[14] == 50000
        assert any(r[3] == "Qadoq shikastlangan" for r in wb["Omborga qaytarishlar"].iter_rows(values_only=True))
        # the owner's own report: own numbers + the network, owner only
        mine = c.get("/api/v1/my-store/report?days=7").json()["data"]
        assert mine["me"]["handled"] == 2 and mine["me"]["handed_back"] == 1 and mine["me"]["revenue"] == "50000.00"
        assert mine["network"]["branches"] >= 1 and len(mine["orders"]) == 1 and mine["hand_backs"][0]["reason"] == "Qadoq shikastlangan"
        assert numbers[0] in {o["number"] for o in mine["orders"]}

    def test_report_is_owner_only(self, monkeypatch):
        owner, store = _owner_store("Kassir hisobot")
        cashier_id = uuid.uuid4()
        _make_user(cashier_id, 42)
        from apps.stores.infrastructure.models import StoreCashier
        StoreCashier.objects.create(store=store, user_id=cashier_id)
        assert _client_as(monkeypatch, cashier_id).get("/api/v1/my-store/report").status_code == 403
        assert _client_as(monkeypatch, owner.id).get("/api/v1/my-store/report?days=90").json()["data"]["me"]["handled"] == 0
