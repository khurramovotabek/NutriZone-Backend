"""
Branch receiving (my-store/transfers API), the shared transfer read model (batch rows) and the .xlsx waybill:
only the branch's own transfers, scan counts -> partial / final receipt through transfer_service, batches kept apart.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, timedelta

import pytest
from django.db import transaction
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.stores.infrastructure import stock_service, transfer_export, transfer_read, transfer_service
from apps.stores.infrastructure.models import StockTransfer, Store, StoreProductStock
from apps.stores.tests.interfaces.test_my_store_products import _authenticate_as, _make_owner, _make_variant

pytestmark = [pytest.mark.django_db, pytest.mark.integration]
OLD, NEW = date.today() + timedelta(days=200), date.today() + timedelta(days=400)


def _put(store, v, qty, batch="", expiry=None):
    with transaction.atomic():
        stock_service.move(store, v.id, qty, "receipt", reference="T", mirror=False, batch_number=batch, expiry_date=expiry)


def _setup(two_batches=True):
    owner = _make_owner(full_name="Dilshod Karimov")
    shop = Store.objects.create(id=uuid.uuid4(), owner=owner, name="Olmazor", slug=f"o-{uuid.uuid4().hex[:6]}")
    wh = Store.objects.create(id=uuid.uuid4(), name="Sklad", slug=f"w-{uuid.uuid4().hex[:6]}", kind="warehouse")
    v = _make_variant(barcode=f"479{uuid.uuid4().int % 10**10:010d}", status="active")
    _put(wh, v, 10, "APP-250", OLD)
    if two_batches:
        _put(wh, v, 10, "APP-260", NEW)
    lines = [{"variant_id": str(v.id), "quantity": 10, "batch": "APP-250", "expiry": OLD.isoformat()}]
    if two_batches:
        lines.append({"variant_id": str(v.id), "quantity": 10, "batch": "APP-260", "expiry": NEW.isoformat()})
    t = transfer_service.create(wh.id, shop.id, "Sklad")
    transfer_service.update_draft(t.pk, "Sklad", lines)
    transfer_service.mark_ready(t.pk, "Sklad")
    picked = {str(i.id): i.expected_quantity for i in t.items.all()}
    transfer_service.dispatch(t.pk, "Sklad", picked=picked)
    return owner, shop, wh, v, StockTransfer.objects.get(pk=t.pk)


class TestReadModel:
    def test_batch_lines_kept_apart(self):
        _owner, _shop, _wh, _v, t = _setup()
        rows = transfer_read.batch_rows(t)
        assert [(r["batch"], r["sent"], r["received"]) for r in rows] == [("APP-250", 10, None), ("APP-260", 10, None)]
        p = transfer_read.payload(t)
        assert p["summary"]["batches"] == 2 and p["summary"]["products"] == 1 and len(p["items"]) == 2

    def test_lookup_returns_source_batches(self):
        owner = _make_owner()
        wh = Store.objects.create(id=uuid.uuid4(), owner=owner, name="W", slug=f"w-{uuid.uuid4().hex[:6]}", kind="warehouse")
        v = _make_variant(barcode=f"479{uuid.uuid4().int % 10**10:010d}", status="active")
        _put(wh, v, 5, "B-NEW", NEW)
        _put(wh, v, 3, "B-OLD", OLD)
        data = transfer_read.lookup(v.barcode, wh)
        assert [b["batch"] for b in data["batches"]] == ["B-OLD", "B-NEW"] and data["available"] == 8


class TestStoreApi:
    def test_incoming_list_detail_and_strangers_get_404(self, monkeypatch):
        owner, shop, _wh, _v, t = _setup()
        c = APIClient()
        _authenticate_as(monkeypatch, owner)
        r = c.get("/api/v1/my-store/transfers")
        assert r.status_code == 200
        body = r.json()["data"]
        assert body["counts"]["incoming"] == 1 and body["results"][0]["number"] == t.number and body["results"][0]["sent"] == 20
        d = c.get(f"/api/v1/my-store/transfers/{t.pk}").json()["data"]
        assert d["can_receive"] is True and d["direction"] == "incoming" and len(d["reasons"]) >= 5
        stranger = _make_owner()
        Store.objects.create(id=uuid.uuid4(), owner=stranger, name="Other", slug=f"x-{uuid.uuid4().hex[:6]}")
        _authenticate_as(monkeypatch, stranger)
        assert c.get(f"/api/v1/my-store/transfers/{t.pk}").status_code == 404
        assert c.post(f"/api/v1/my-store/transfers/{t.pk}/receive", {"mode": "final", "counts": {}}, format="json").status_code == 404

    def test_partial_then_final_with_shortage_needs_reason_and_posts_per_batch(self, monkeypatch):
        owner, shop, _wh, v, t = _setup()
        c = APIClient()
        _authenticate_as(monkeypatch, owner)
        old, new = sorted(t.items.all(), key=lambda i: i.batch_number)
        r = c.post(f"/api/v1/my-store/transfers/{t.pk}/receive", {"mode": "partial", "counts": {str(old.id): 10}}, format="json")
        assert r.status_code == 200 and r.json()["data"]["status"] == "partial"
        assert StoreProductStock.objects.get(store=shop, product_variant=v).quantity == 10
        counts = {str(old.id): 10, str(new.id): 8}
        r = c.post(f"/api/v1/my-store/transfers/{t.pk}/receive", {"mode": "final", "counts": counts}, format="json")
        assert r.status_code == 409 and r.json()["error"]["code"] == "DISCREPANCY_REASON_REQUIRED"
        r = c.post(f"/api/v1/my-store/transfers/{t.pk}/receive", {"mode": "final", "counts": counts, "reasons": {str(new.id): {"reason": "missing", "note": "2 ta quti"}}}, format="json")
        data = r.json()["data"]
        assert r.status_code == 200 and data["discrepancy_open"] is True and data["can_receive"] is False
        assert "Dilshod Karimov" in data["received_by"]
        assert StoreProductStock.objects.get(store=shop, product_variant=v).quantity == 18
        rows = transfer_read.batch_rows(StockTransfer.objects.get(pk=t.pk))
        assert [(r["batch"], r["sent"], r["received"], r["difference"]) for r in rows] == [("APP-250", 10, 10, 0), ("APP-260", 10, 8, -2)]

    def test_bad_input(self, monkeypatch):
        owner, _shop, _wh, _v, t = _setup(two_batches=False)
        c = APIClient()
        _authenticate_as(monkeypatch, owner)
        assert c.post(f"/api/v1/my-store/transfers/{t.pk}/receive", {"mode": "x"}, format="json").status_code == 422
        assert c.post(f"/api/v1/my-store/transfers/{t.pk}/receive", {"mode": "save", "counts": {"a": -1}}, format="json").status_code == 422
        assert c.get("/api/v1/my-store/transfers/not-a-uuid").status_code == 404


class TestWaybill:
    def test_xlsx_one_row_per_batch_with_totals_and_signatures(self, monkeypatch):
        owner, _shop, _wh, _v, t = _setup()
        c = APIClient()
        _authenticate_as(monkeypatch, owner)
        r = c.get(f"/api/v1/my-store/transfers/{t.pk}/document.xlsx")
        assert r.status_code == 200 and t.number in r["Content-Disposition"]
        ws = load_workbook(io.BytesIO(r.content)).active
        values = [[c for c in row if c is not None] for row in ws.iter_rows(values_only=True)]
        flat = [x for row in values for x in row]
        assert "APP-250" in flat and "APP-260" in flat and "ИТОГО" in flat and "Отпустил" in flat and "Получил" in flat
        assert ws.print_title_rows and ws.page_setup.orientation == "landscape"
        assert transfer_export.filename(t) == f"{t.number}.xlsx"
