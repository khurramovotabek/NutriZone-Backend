"""
Dashboard → Ombor (WMS): inventory transfers, goods receipts (Kirim), stock by location, and the stock ledger.

Every page is a staff-only admin view; every change goes through transfer_service / stock_service
(database transactions, row locks, validated state machine, ledger rows). Permissions are the
model permissions of StockTransfer / StockReceipt / StockMovement (see the role groups created by
migration 0008: Sklad operatori, Sklad menejeri, Filial qabul qiluvchi, Inventar menejeri).
"""
from __future__ import annotations

import io
import json
import uuid
from datetime import datetime, time, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.db.models import Count, F, Q, Sum
from django.http import HttpResponse, HttpResponseRedirect, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.urls import path, reverse
from django.utils import timezone
from django.utils.html import format_html
from django.utils.translation import gettext as _t
from django.utils.translation import gettext_lazy as _

from apps.stores.infrastructure import stock_service, transfer_service
from apps.stores.infrastructure.models import (
    DiscrepancyReason,
    StockMovement,
    StockMovementType,
    StockReceipt,
    StockReceiptItem,
    StockTransfer,
    Store,
    Supplier,
    StoreKind,
    StoreProductStock,
    TransferStatus,
)
from common.exceptions import DomainException

TZ = ZoneInfo("Asia/Tashkent")
STATUS_TONE = {
    "draft": "#64748b", "ready": "#0ea5e9", "in_transit": "#6366f1", "receiving": "#f59e0b", "partial": "#f97316",
    "received": "#10b981", "completed": "#16a34a", "cancelled": "#ef4444",
}
MOVEMENT_TONE = {
    "receipt": "#10b981", "sale": "#3b82f6", "return": "#a855f7", "transfer_out": "#f97316", "transfer_in": "#14b8a6",
    "adjustment": "#64748b", "damage": "#ef4444", "expired": "#b91c1c", "reservation": "#0ea5e9", "release": "#94a3b8", "write_off": "#be123c",
}


def _brand() -> str:
    from apps.platform.infrastructure.branding import brand_name

    return brand_name()


def _is_uuid(value) -> bool:
    try:
        uuid.UUID(str(value))
    except (TypeError, ValueError):
        return False
    return True


def _slug() -> str:
    from apps.platform.infrastructure.branding import file_slug

    return file_slug()


def _actor(request) -> str:
    return request.user.get_full_name() or request.user.get_username()


def _local(dt, fmt="%d.%m.%Y %H:%M") -> str | None:
    return timezone.localtime(dt, TZ).strftime(fmt) if dt else None


def _som(v) -> str:
    return f"{int(round(Decimal(v or 0))):,}".replace(",", " ")


def _error(exc: Exception, status: int = 409) -> JsonResponse:
    code = getattr(exc, "code", None) or exc.__class__.__name__
    details = getattr(exc, "details", None) or {}
    return JsonResponse({"ok": False, "error": {"code": code, "message": str(exc), "details": details}}, status=getattr(exc, "http_status", status))


def _body(request) -> dict:
    try:
        return json.loads(request.body or b"{}")
    except ValueError:
        return {}


def _location_payload(s: Store) -> dict:
    return {"id": str(s.id), "name": s.name, "kind": s.kind, "color": s.color or "#94a3b8", "city": s.city}


def _stock_rows(store: Store, variant_ids=None, q: str = "", only_available: bool = False, limit: int = 60) -> list[dict]:
    qs = StoreProductStock.objects.filter(store=store)
    if variant_ids is not None:
        qs = qs.filter(product_variant_id__in=list(variant_ids))
    if q:
        from apps.products.infrastructure.barcodes import barcode_contains_q

        qs = qs.filter(Q(product_variant__sku__icontains=q) | barcode_contains_q(q, "product_variant__") | Q(product_variant__product__translations__name__icontains=q)).distinct()
    rows = list(qs.order_by("-quantity")[:limit]) if variant_ids is None else list(qs)
    infos = stock_service.variant_infos([r.product_variant_id for r in rows])
    out = []
    for r in rows:
        info = infos.get(r.product_variant_id)
        if info is None or (only_available and r.available <= 0):
            continue
        out.append({
            "variant_id": str(r.product_variant_id), "name": info.name, "sku": info.sku, "barcode": info.barcode, "size": info.size,
            "image_url": info.image_url, "price": str(info.price) if info.price is not None else None, "active": info.active,
            "quantity": r.quantity, "reserved": r.reserved_quantity, "available": r.available,
        })
    return out


def _costs(variant_ids) -> dict:
    from apps.stores.infrastructure.count_service import cost_map

    return cost_map(variant_ids) if variant_ids else {}


def _cost_fields(variant_ids) -> dict:
    """
    variant_id -> {"cost", "cost_source", "prefill_cost", "last_cost"} for the scanning pages. A new goods receipt is
    pre-filled with the latest receipt cost, else the catalog cost price -- never with the selling price.
    """
    from apps.stores.infrastructure.count_service import COST_PRICE, cost_info

    info, last = cost_info(variant_ids), _last_costs(variant_ids)
    out = {}
    for vid in variant_ids:
        cost, src = info.get(vid, (None, None))
        out[vid] = {"cost": str(cost) if cost else "", "cost_source": src or "", "last_cost": str(last.get(vid) or ""),
                    "prefill_cost": str(cost) if cost and src and src != COST_PRICE else ""}
    return out


def _last_costs(variant_ids) -> dict:
    """The unit cost of the latest goods receipt only (no price fallback) -- pre-fills the cost on a new receipt."""
    out: dict = {}
    for row in StockReceiptItem.objects.filter(product_variant_id__in=list(variant_ids), unit_cost__isnull=False).order_by("product_variant_id", "-receipt__created_at").values("product_variant_id", "unit_cost"):
        out.setdefault(row["product_variant_id"], row["unit_cost"])
    return out


# ================================================================ transfers


class TransferStatusFilter(admin.SimpleListFilter):
    title = _("Status")
    parameter_name = "status"

    def lookups(self, request, model_admin):
        return [("awaiting", _t("Qabul kutilmoqda")), *TransferStatus.choices]

    def queryset(self, request, queryset):
        if self.value() == "awaiting":
            return queryset.filter(status__in=["in_transit", "receiving", "partial"])
        return queryset.filter(status=self.value()) if self.value() else queryset


class CreatedByFilter(admin.SimpleListFilter):
    title = _("Created by")
    parameter_name = "author"

    def lookups(self, request, model_admin):
        return [(n, n) for n in StockTransfer.objects.exclude(created_by="").values_list("created_by", flat=True).distinct().order_by("created_by")]

    def queryset(self, request, queryset):
        return queryset.filter(created_by=self.value()) if self.value() else queryset


@admin.register(StockTransfer)
class StockTransferAdmin(admin.ModelAdmin):
    class Media:
        css = {"all": ("admin/css/nutrizone-wms.css",)}

    list_display = ("number_col", "created_col", "route_col", "lines_col", "units_col", "status_col", "created_by", "sent_col", "received_col", "open_col")
    list_display_links = None
    list_filter = (TransferStatusFilter, "source", "destination", CreatedByFilter, ("created_at", admin.DateFieldListFilter))
    search_fields = ("number",)
    search_help_text = _("Transfer raqami (TR-000001) yoki mahsulot nomi / SKU / shtrix-kod")
    list_select_related = ("source", "destination")
    list_per_page = 30
    ordering = ("-created_at",)
    change_list_template = "admin/wms/transfer_list.html"

    def has_delete_permission(self, request, obj=None):
        return False

    def get_queryset(self, request):
        return super().get_queryset(request).annotate(
            _lines=Count("items"), _units=Sum("items__expected_quantity"), _sent=Sum("items__dispatched_quantity"), _recv=Sum("items__received_posted"),
            _value=Sum(F("items__unit_cost") * F("items__expected_quantity")),
        )

    def get_changelist_instance(self, request):
        cl = super().get_changelist_instance(request)
        now = timezone.now()
        for t in cl.result_list:
            sent, recv = t._sent or 0, t._recv or 0
            t.nz = {
                "color": STATUS_TONE.get(t.status, "#64748b"), "lines": t._lines, "units": sent or t._units or 0, "received": recv,
                "progress": round(recv * 100 / sent) if sent else 0, "created": _local(t.created_at), "sent": _local(t.dispatched_at),
                "received_at": _local(t.received_at), "age_days": (now - t.created_at).days,
                "stuck": t.status == "in_transit" and t.dispatched_at and (now - t.dispatched_at).days >= 2,
                "value": _som(t._value) if t._value else "",
            }
        return cl

    def get_search_results(self, request, queryset, search_term):
        term = search_term.strip()
        if not term:
            return queryset, False
        product_hit = Q(items__name__icontains=term) | Q(items__sku__iexact=term) | Q(items__barcode=term)
        return queryset.filter(Q(number__icontains=term) | product_hit).distinct(), True

    # --- columns
    @admin.display(description=_("Transfer No."), ordering="number")
    def number_col(self, obj):
        return format_html('<a class="wms-num" href="{}">{}</a>', reverse("admin:stores_stocktransfer_workspace", args=[obj.pk]), obj.number)

    @admin.display(description=_("Created"), ordering="created_at")
    def created_col(self, obj):
        return _local(obj.created_at)

    @admin.display(description=_("Route"))
    def route_col(self, obj):
        return format_html(
            '<span class="wms-route"><i class="wms-dot" style="background:{}"></i>{}{}<span class="wms-arrow">→</span><i class="wms-dot" style="background:{}"></i>{}</span>',
            obj.source.color or "#94a3b8", obj.source.name, format_html(' <small class="wms-tag">{}</small>', _t("sklad")) if obj.source.kind == StoreKind.WAREHOUSE else "",
            obj.destination.color or "#94a3b8", obj.destination.name,
        )

    @admin.display(description=_("Products"))
    def lines_col(self, obj):
        return obj._lines

    @admin.display(description=_("Units"))
    def units_col(self, obj):
        return obj._sent or obj._units or 0

    @admin.display(description=_("Status"), ordering="status")
    def status_col(self, obj):
        extra = format_html(' <span class="wms-pill" style="--t:#ef4444">{}</span>', _t("tafovut")) if obj.discrepancy_open else ""
        return format_html('<span class="wms-pill" style="--t:{}">{}</span>{}', STATUS_TONE.get(obj.status, "#64748b"), obj.get_status_display(), extra)

    @admin.display(description=_("Sent"), ordering="dispatched_at")
    def sent_col(self, obj):
        return _local(obj.dispatched_at) or "—"

    @admin.display(description=_("Received"), ordering="received_at")
    def received_col(self, obj):
        return _local(obj.received_at) or "—"

    @admin.display(description="")
    def open_col(self, obj):
        return format_html('<a class="btn btn-sm btn-outline-primary" href="{}">{}</a>', reverse("admin:stores_stocktransfer_workspace", args=[obj.pk]), _t("Ochish"))

    # --- list page KPIs
    def changelist_view(self, request, extra_context=None):
        today = datetime.combine(timezone.localtime(timezone.now(), TZ).date(), time.min, tzinfo=TZ)
        agg = StockTransfer.objects.aggregate(
            today=Count("id", filter=Q(created_at__gte=today)),
            in_transit=Count("id", filter=Q(status="in_transit")),
            receiving=Count("id", filter=Q(status__in=["receiving", "partial"])),
            partial=Count("id", filter=Q(status="partial")),
            discrepancy=Count("id", filter=Q(discrepancy_open=True)),
            completed=Count("id", filter=Q(status="completed")),
            drafts=Count("id", filter=Q(status__in=["draft", "ready"])),
        )
        units = StockTransfer.objects.filter(status__in=["in_transit", "receiving", "partial"]).aggregate(
            d=Sum("items__dispatched_quantity"), r=Sum("items__received_posted"))
        counts = dict(StockTransfer.objects.values_list("status").annotate(n=Count("id")))
        current = request.GET.get("status", "")
        tabs = [{"key": "", "label": _t("All"), "n": sum(counts.values()), "color": "#64748b", "on": not current},
                {"key": "awaiting", "label": _t("Qabul kutilmoqda"), "n": sum(counts.get(k, 0) for k in ("in_transit", "receiving", "partial")), "color": "#f59e0b", "on": current == "awaiting"}]
        for key, label in TransferStatus.choices:
            if counts.get(key) or key in ("draft", "in_transit", "completed"):
                tabs.append({"key": key, "label": label, "n": counts.get(key, 0), "color": STATUS_TONE.get(key, "#64748b"), "on": current == key})
        extra_context = {
            **(extra_context or {}),
            "wms_kpis": {**agg, "units_in_transit": max(0, (units["d"] or 0) - (units["r"] or 0))},
            "can_create": request.user.has_perm("stores.add_stocktransfer"), "nz_tabs": tabs, "nz_status": current,
            "export_query": request.GET.urlencode(),
            "title": _t("Transferlar"),
        }
        return super().changelist_view(request, extra_context)

    # --- custom pages & API
    def get_urls(self):
        v = self.admin_site.admin_view
        return [
            path("new/", v(self.new_view), name="stores_stocktransfer_new"),
            path("api/lookup/", v(self.api_lookup), name="stores_stocktransfer_lookup"),
            path("api/search/", v(self.api_search), name="stores_stocktransfer_search"),
            path("api/suggest/", v(self.api_suggest), name="stores_stocktransfer_suggest"),
            path("api/attach-barcode/", v(self.api_attach_barcode), name="stores_stocktransfer_attach_barcode"),
            path("api/batches/", v(self.api_batches), name="stores_stocktransfer_batches"),
            path("export.xlsx", v(self.export_view), name="stores_stocktransfer_export"),
            path("<uuid:pk>/document.xlsx", v(self.document_xlsx_view), name="stores_stocktransfer_document_xlsx"),
            path("<uuid:pk>/work/", v(self.workspace_view), name="stores_stocktransfer_workspace"),
            path("<uuid:pk>/data/", v(self.api_data), name="stores_stocktransfer_data"),
            path("<uuid:pk>/do/<str:action>/", v(self.api_action), name="stores_stocktransfer_action"),
            path("<uuid:pk>/document/", v(self.document_view), name="stores_stocktransfer_document"),
            *super().get_urls(),
        ]

    def change_view(self, request, object_id, form_url="", extra_context=None):
        return HttpResponseRedirect(reverse("admin:stores_stocktransfer_workspace", args=[object_id]))

    def add_view(self, request, form_url="", extra_context=None):
        return HttpResponseRedirect(reverse("admin:stores_stocktransfer_new"))

    def _ctx(self, request, title):
        return {**self.admin_site.each_context(request), "opts": self.model._meta, "title": title}

    def new_view(self, request):
        if not request.user.has_perm("stores.add_stocktransfer"):
            raise PermissionDenied
        locations = list(transfer_service.locations())
        if request.method == "POST":
            try:
                t = transfer_service.create(request.POST.get("source"), request.POST.get("destination"), _actor(request), request.POST.get("notes", ""))
            except DomainException as exc:
                messages.error(request, str(exc))
            else:
                lines = self._posted_lines(request)
                if lines:  # opened from the stock pages or picked from the suggestions: the products are already on the list
                    try:
                        transfer_service.update_draft(t.pk, _actor(request), lines)
                    except DomainException as exc:
                        messages.warning(request, str(exc))
                    else:
                        messages.success(request, _t("Transfer qoralamasi yaratildi: %(n)s ta mahsulot qo'shildi. Tekshirib «Tayyor» qiling.") % {"n": len(lines)})
                return HttpResponseRedirect(reverse("admin:stores_stocktransfer_workspace", args=[t.pk]))
        by_id = {str(s.pk): s for s in locations}
        default_source = by_id.get(request.GET.get("source", "")) or next((s for s in locations if s.kind == StoreKind.WAREHOUSE), None)
        default_destination = by_id.get(request.GET.get("destination", ""))
        prefill = None
        if request.GET.get("variant"):
            info = stock_service.variant_infos([request.GET["variant"]]).get(uuid.UUID(request.GET["variant"])) if _is_uuid(request.GET["variant"]) else None
            qty = request.GET.get("quantity", "")
            if info is not None:
                prefill = {"info": info, "quantity": qty if qty.isdigit() and int(qty) > 0 else "1"}
        stats = self._location_stats(locations)
        for s in locations:
            s.nz = stats[str(s.pk)]
        return render(request, "admin/wms/transfer_new.html", {
            **self._ctx(request, _t("Yangi transfer")), "locations": locations, "default_source": default_source,
            "default_destination": default_destination, "prefill": prefill,
            "suggest_url": reverse("admin:stores_stocktransfer_suggest"),
        })

    @staticmethod
    def _posted_lines(request) -> list[dict]:
        """`lines` (JSON [{variant_id, quantity}] from the suggestions) and/or the single prefilled `variant` + `quantity`."""
        lines = []
        try:
            raw = json.loads(request.POST.get("lines") or "[]")
        except ValueError:
            raw = []
        for ln in raw if isinstance(raw, list) else []:
            if isinstance(ln, dict) and _is_uuid(ln.get("variant_id")) and str(ln.get("quantity", "")).isdigit() and int(ln["quantity"]) > 0:
                lines.append({"variant_id": str(ln["variant_id"]), "quantity": int(ln["quantity"])})
        variant, qty = request.POST.get("variant"), request.POST.get("quantity")
        if variant and qty and _is_uuid(variant) and not any(ln["variant_id"] == str(variant) for ln in lines):
            lines.append({"variant_id": variant, "quantity": qty})
        return lines

    @staticmethod
    def _location_stats(locations) -> dict:
        """Location id -> what it holds (sellable units, products in stock), how many products are under their minimum, transfers on the way in."""
        from django.db.models import F

        from apps.stores.infrastructure import planning_service

        held = {r["store_id"]: r for r in StoreProductStock.objects.filter(store__in=locations).values("store_id").annotate(
            units=Sum(F("quantity") - F("reserved_quantity")), skus=Count("id", filter=Q(quantity__gt=F("reserved_quantity"))))}
        need = planning_service.overview()
        coming = dict(StockTransfer.objects.filter(destination__in=locations, status__in=planning_service.OPEN_TRANSFERS).values_list("destination_id").annotate(n=Count("id")))
        return {str(s.pk): {"units": (held.get(s.pk) or {}).get("units") or 0, "skus": (held.get(s.pk) or {}).get("skus") or 0,
                            "need": (need.get(s.pk) or {}).get("need", 0), "coming": coming.get(s.pk, 0)} for s in locations}

    def api_suggest(self, request):
        """New transfer: what the destination is short of that the source can send (planning_service.transfer_suggestions)."""
        if not request.user.has_perm("stores.add_stocktransfer"):
            raise PermissionDenied
        from apps.stores.infrastructure import planning_service

        src = Store.objects.filter(pk=request.GET.get("source")).first() if _is_uuid(request.GET.get("source")) else None
        dst = Store.objects.filter(pk=request.GET.get("destination")).first() if _is_uuid(request.GET.get("destination")) else None
        if src is None or dst is None or src.pk == dst.pk:
            return JsonResponse({"ok": True, "data": {"lines": []}})
        lines = planning_service.transfer_suggestions(dst, src)
        costs = _costs([uuid.UUID(ln["variant_id"]) for ln in lines])
        for ln in lines:
            ln["cost"] = str(costs.get(uuid.UUID(ln["variant_id"])) or "")
        return JsonResponse({"ok": True, "data": {"lines": lines}})

    def workspace_view(self, request, pk):
        t = get_object_or_404(StockTransfer, pk=pk)
        perms = self._perms(request)
        return render(request, "admin/wms/transfer_workspace.html", {
            **self._ctx(request, f"{_t('Transfer')} {t.number}"), "transfer": t,
            "config": {
                "id": str(t.pk), "number": t.number,
                "urls": {
                    "data": reverse("admin:stores_stocktransfer_data", args=[t.pk]),
                    "action": reverse("admin:stores_stocktransfer_action", args=[t.pk, "__"]),
                    "lookup": reverse("admin:stores_stocktransfer_lookup"),
                    "search": reverse("admin:stores_stocktransfer_search"),
                    "document": reverse("admin:stores_stocktransfer_document", args=[t.pk]),
                    "xlsx": reverse("admin:stores_stocktransfer_document_xlsx", args=[t.pk]),
                    "batches": reverse("admin:stores_stocktransfer_batches"),
                    "list": reverse("admin:stores_stocktransfer_changelist"),
                },
                "perms": perms,
                "reasons": DiscrepancyReason.choices,
                # the document's own locations stay selectable even if one was deactivated since
                "locations": [_location_payload(s) for s in transfer_service.locations()]
                + [_location_payload(s) for s in (t.source, t.destination) if s.status == "inactive"],
            },
        })

    @staticmethod
    def _perms(request) -> dict:
        u = request.user
        return {
            "edit": u.has_perm("stores.change_stocktransfer"), "send": u.has_perm("stores.send_stocktransfer"),
            "receive": u.has_perm("stores.receive_stocktransfer"), "approve": u.has_perm("stores.approve_discrepancy"),
            "cancel": u.has_perm("stores.cancel_stocktransfer"),
        }

    def _payload(self, t: StockTransfer) -> dict:
        from apps.stores.infrastructure import transfer_read

        return transfer_read.payload(t)

    def api_data(self, request, pk):
        return JsonResponse({"ok": True, "data": self._payload(get_object_or_404(StockTransfer, pk=pk))})

    ACTION_PERMS = {
        "save_draft": "change_stocktransfer", "ready": "change_stocktransfer", "back_to_draft": "change_stocktransfer", "pick": "change_stocktransfer",
        "dispatch": "send_stocktransfer", "cancel": "cancel_stocktransfer",
        "receive_save": "receive_stocktransfer", "receive_partial": "receive_stocktransfer", "receive_final": "receive_stocktransfer",
        "approve": "approve_discrepancy", "complete": "receive_stocktransfer",
    }

    def api_action(self, request, pk, action):
        if request.method != "POST":
            return JsonResponse({"ok": False}, status=405)
        perm = self.ACTION_PERMS.get(action)
        if perm is None:
            return JsonResponse({"ok": False, "error": {"code": "UNKNOWN_ACTION", "message": action}}, status=404)
        if not request.user.has_perm(f"stores.{perm}"):
            return JsonResponse({"ok": False, "error": {"code": "FORBIDDEN", "message": _t("Bu amal uchun sizda ruxsat yo'q.")}}, status=403)
        b, a, s = _body(request), _actor(request), transfer_service
        try:
            if action == "save_draft":
                s.update_draft(pk, a, b.get("lines") or [], notes=b.get("notes"), source_id=b.get("source"), destination_id=b.get("destination"))
            elif action == "ready":
                if "lines" in b:
                    s.update_draft(pk, a, b.get("lines") or [], notes=b.get("notes"))
                s.mark_ready(pk, a)
            elif action == "back_to_draft":
                s.back_to_draft(pk, a)
            elif action == "pick":
                s.save_pick(pk, a, b.get("picked") or {})
            elif action == "dispatch":
                s.dispatch(pk, a, picked=b.get("picked"), confirm_mismatch=bool(b.get("confirm")))
            elif action == "cancel":
                s.cancel(pk, a, str(b.get("reason") or ""))
            elif action == "receive_save":
                s.save_receiving(pk, a, b.get("counts") or {}, b.get("reasons"))
            elif action in ("receive_partial", "receive_final"):
                s.confirm_receipt(pk, a, counts=b.get("counts") or {}, reasons=b.get("reasons"), final=action == "receive_final")
            elif action == "approve":
                s.approve_discrepancies(pk, a, b.get("resolutions") or {})
            elif action == "complete":
                s.complete(pk, a)
        except DomainException as exc:
            return _error(exc)
        return JsonResponse({"ok": True, "data": self._payload(StockTransfer.objects.get(pk=pk))})

    def api_lookup(self, request):
        """A scanned barcode / typed SKU → the product, what the chosen location has of it and its batches there (FEFO order)."""
        from apps.stores.infrastructure import transfer_read

        location = Store.objects.filter(id=request.GET.get("location")).first() if _is_uuid(request.GET.get("location")) else None
        data = transfer_read.lookup(request.GET.get("code", ""), location)
        if data is None:
            return JsonResponse({"ok": False, "error": {"code": "NOT_FOUND", "message": _t("Bu shtrix-kod bo'yicha mahsulot topilmadi.")}}, status=404)
        data.update(_cost_fields([uuid.UUID(data["variant_id"])])[uuid.UUID(data["variant_id"])])
        return JsonResponse({"ok": True, "data": data})

    def api_batches(self, request):
        """A location's batches with stock for some products (the draft's batch picker)."""
        from apps.stores.infrastructure import transfer_read

        location = Store.objects.filter(id=request.GET.get("location")).first() if _is_uuid(request.GET.get("location")) else None
        ids = [uuid.UUID(v) for v in request.GET.get("variants", "").split(",") if _is_uuid(v)][:200]
        if location is None:
            return JsonResponse({"ok": True, "data": {}})
        return JsonResponse({"ok": True, "data": {str(k): v for k, v in transfer_read.batches(location, ids).items()}})

    def document_xlsx_view(self, request, pk):
        """The waybill as .xlsx (draft before dispatch, final with sent vs received after receiving)."""
        from apps.stores.infrastructure import transfer_export

        t = get_object_or_404(StockTransfer.objects.select_related("source", "destination"), pk=pk)
        response = HttpResponse(transfer_export.workbook(t, _brand()), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{transfer_export.filename(t)}"'
        return response

    def export_view(self, request):
        """The filtered transfer list (same filters / search as the page) as .xlsx."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        cl = self.get_changelist_instance(request)
        qs = cl.get_queryset(request)
        wb = Workbook()
        ws = wb.active
        ws.title = "Transfers"
        head = ["Документ №", "Создан", "Отправитель", "Получатель", "Товаров", "Кол-во (план)", "Отправлено", "Принято", "Сумма (себест.), UZS",
                "Создал", "Отправил", "Принял", "Статус", "Расхождение"]
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E293B")
        for t in qs.select_related("source", "destination")[:5000]:
            ws.append([t.number, _local(t.created_at), t.source.name, t.destination.name, t._lines, t._units or 0, t._sent or 0, t._recv or 0,
                       float(t._value or 0), t.created_by, t.dispatched_by, t.received_by, t.get_status_display(), "да" if t.discrepancy_open else ""])
        for i, w in enumerate((14, 17, 26, 26, 9, 12, 12, 10, 18, 20, 20, 20, 18, 12), 1):
            ws.column_dimensions[chr(64 + i)].width = w
        for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
            row[0].number_format = '#,##0" UZS"'
        ws.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="transfers-{timezone.localdate():%Y-%m-%d}.xlsx"'
        return response

    def api_attach_barcode(self, request):
        """Kirim: an unknown barcode (a new batch's label) is attached to an existing product as an extra barcode."""
        u = request.user
        if request.method != "POST" or not (u.has_perm("products.change_productvariant") or u.has_perm("stores.add_stockreceipt")):
            raise PermissionDenied
        from datetime import date

        from apps.products.infrastructure import barcodes
        from apps.products.infrastructure.models import ProductVariant

        b = _body(request)
        variant = ProductVariant.objects.filter(pk=b.get("variant_id")).first()
        if variant is None:
            return JsonResponse({"ok": False, "error": {"message": _t("Mahsulotni tanlang.")}}, status=400)
        expiry = None
        if b.get("expiry"):
            try:
                expiry = date.fromisoformat(str(b["expiry"])[:10])
            except ValueError:
                return JsonResponse({"ok": False, "error": {"message": _t("Yaroqlilik sanasi noto'g'ri.")}}, status=400)
        try:
            row = barcodes.attach(variant, b.get("code", ""), batch=str(b.get("batch") or ""), expires_at=expiry, note=str(b.get("note") or ""), actor=_actor(request))
        except ValueError as exc:
            return JsonResponse({"ok": False, "error": {"message": str(exc)}}, status=409)
        return JsonResponse({"ok": True, "data": {"barcode": row.barcode, "variant_id": str(variant.id)}})

    def api_search(self, request):
        """Search a location's stock (name / SKU / barcode); with `all=1` search the whole catalogue (goods receipt)."""
        q = (request.GET.get("q") or "").strip()
        location = Store.objects.filter(id=request.GET.get("location")).first() if request.GET.get("location") else None
        if request.GET.get("all") == "1" or location is None:
            variants = stock_service.search_variants(q, limit=25)
            infos = stock_service.variant_infos([v.id for v in variants])
            avail = stock_service.available_map(location, [v.id for v in variants]) if location else {}
            data = [{"variant_id": str(v.id), "name": infos[v.id].name, "sku": v.sku, "barcode": v.barcode or "", "size": infos[v.id].size,
                     "image_url": infos[v.id].image_url, "price": str(infos[v.id].price) if infos[v.id].price is not None else None,
                     "active": infos[v.id].active, "available": avail.get(v.id, 0)} for v in variants if v.id in infos]
        else:
            data = _stock_rows(location, q=q, only_available=True, limit=40)
        ids = [uuid.UUID(d["variant_id"]) for d in data]
        costs = _cost_fields(ids)
        for d in data:
            d.update(costs[uuid.UUID(d["variant_id"])])
        return JsonResponse({"ok": True, "data": data})

    def document_view(self, request, pk):
        t = get_object_or_404(StockTransfer.objects.select_related("source", "destination"), pk=pk)
        items = list(t.items.all())
        rows = []
        for n, i in enumerate(items, 1):
            qty = i.dispatched_quantity or i.expected_quantity
            rows.append({"n": n, "item": i, "qty": qty, "total": i.unit_cost * qty})
        return render(request, "admin/wms/transfer_document.html", {
            **self._ctx(request, f"{t.number}"), "t": t, "rows": rows, "summary": transfer_service.summary(t),
            "created": _local(t.created_at), "sent": _local(t.dispatched_at), "received": _local(t.received_at),
            "total_value": sum((r["total"] for r in rows), Decimal("0")),
        })


# ================================================================ goods receipts (Kirim)


@admin.register(StockReceipt)
class StockReceiptAdmin(admin.ModelAdmin):
    class Media:
        css = {"all": ("admin/css/nutrizone-wms.css",)}

    list_display = ("number", "created_col", "store_col", "supplier", "invoice_number", "lines_col", "total_units", "value_col", "created_by")
    list_filter = ("store", ("created_at", admin.DateFieldListFilter))
    search_fields = ("number", "supplier", "invoice_number", "items__product_variant__sku", "items__product_variant__barcode")
    list_select_related = ("store",)
    ordering = ("-created_at",)
    change_list_template = "admin/wms/receipt_list.html"
    readonly_fields = ("number", "store", "supplier", "invoice_number", "note", "created_by", "created_at", "total_units", "total_value", "items_table")
    fields = readonly_fields

    def has_add_permission(self, request):
        return False  # a receipt is posted from the scanning page, never typed into a raw form

    def has_change_permission(self, request, obj=None):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def get_queryset(self, request):
        return super().get_queryset(request).annotate(_lines=Count("items"))

    def get_changelist_instance(self, request):
        cl = super().get_changelist_instance(request)
        for r in cl.result_list:
            r.nz = {"lines": r._lines, "created": _local(r.created_at), "value": _som(r.total_value) if r.total_value else "",
                    "supplier_url": reverse("admin:stores_supplier_card", args=[r.supplier_ref_id]) if r.supplier_ref_id else ""}
        return cl

    def change_view(self, request, object_id, form_url="", extra_context=None):
        """A printable goods-receipt document instead of a read-only form."""
        r = get_object_or_404(StockReceipt.objects.select_related("store"), pk=object_id)
        items = list(r.items.all())
        infos = stock_service.variant_infos([i.product_variant_id for i in items])
        rows, revenue, costed = [], Decimal(0), Decimal(0)
        for n, i in enumerate(items, 1):
            info = infos.get(i.product_variant_id)
            price = info.price if info and info.price else None
            margin = round((price - i.unit_cost) * 100 / price, 1) if price and i.unit_cost else None
            if margin is not None:
                revenue, costed = revenue + price * i.quantity, costed + i.unit_cost * i.quantity
            rows.append({"n": n, "i": i, "info": info, "total": _som(i.unit_cost * i.quantity) if i.unit_cost is not None else "",
                         "price": _som(price) if price else "", "margin": margin,
                         "tone": "bad" if margin is not None and margin < 0 else "warn" if margin is not None and margin < 15 else "ok"})
        profit = {"revenue": _som(revenue), "profit": _som(revenue - costed), "pct": round((revenue - costed) * 100 / revenue, 1),
                  "loss": revenue < costed} if revenue else None
        nocost = sum(1 for i in items if not i.unit_cost)
        pay = None
        if r.supplier_ref_id:
            from apps.stores.infrastructure import supplier_service
            from apps.stores.infrastructure.supplier_admin import RECEIPT_STATE

            info = supplier_service.balance(r.supplier_ref).receipt_status.get(r.pk) or {}
            color, label = RECEIPT_STATE.get(info.get("status"), ("#64748b", ""))
            pay = {"color": color, "label": label, "open": _som(info.get("open")) if info.get("open") else "", "url": reverse("admin:stores_supplier_card", args=[r.supplier_ref_id])}
        return render(request, "admin/wms/receipt_detail.html", {
            **self.admin_site.each_context(request), "opts": self.model._meta, "title": r.number, "r": r, "rows": rows,
            "created": _local(r.created_at), "value": _som(r.total_value) if r.total_value else "", "pay": pay, "profit": profit, "nocost": nocost,
        })

    @admin.display(description=_("Date"), ordering="created_at")
    def created_col(self, obj):
        return _local(obj.created_at)

    @admin.display(description=_("Location"))
    def store_col(self, obj):
        return format_html('<span class="wms-route"><i class="wms-dot" style="background:{}"></i>{}</span>', obj.store.color or "#94a3b8", obj.store.name)

    @admin.display(description=_("Products"))
    def lines_col(self, obj):
        return obj._lines

    @admin.display(description=_("Value"), ordering="total_value")
    def value_col(self, obj):
        return _som(obj.total_value) if obj.total_value else "—"

    @admin.display(description=_("Products"))
    def items_table(self, obj):
        from django.utils.html import format_html_join

        items = list(obj.items.all())
        infos = stock_service.variant_infos([i.product_variant_id for i in items])
        body = format_html_join("", "<tr><td>{}</td><td>{}</td><td>{}</td><td class='num'>{}</td><td class='num'>{}</td><td>{}</td><td>{}</td></tr>", (
            (n, infos[i.product_variant_id].name if i.product_variant_id in infos else "—", infos[i.product_variant_id].barcode if i.product_variant_id in infos else "",
             i.quantity, _som(i.unit_cost) if i.unit_cost is not None else "—", i.batch_number or "—", i.expiry_date or "—") for n, i in enumerate(items, 1)
        ))
        return format_html('<table class="wms-table"><thead><tr><th>#</th><th>{}</th><th>{}</th><th class="num">{}</th><th class="num">{}</th><th>{}</th><th>{}</th></tr></thead><tbody>{}</tbody></table>',
                           _t("Mahsulot"), _t("Shtrix-kod"), _t("Soni"), _t("Tannarx"), _t("Partiya"), _t("Yaroqlilik muddati"), body)

    def changelist_view(self, request, extra_context=None):
        month = timezone.localtime(timezone.now(), TZ).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        agg = StockReceipt.objects.filter(created_at__gte=month).aggregate(n=Count("id"), units=Sum("total_units"), value=Sum("total_value"))
        agg["value"] = _som(agg["value"]) if agg["value"] else "0"
        suppliers = StockReceipt.objects.filter(created_at__gte=month).exclude(supplier="").values("supplier").distinct().count()
        return super().changelist_view(request, {**(extra_context or {}), "wms_kpis": {**agg, "suppliers": suppliers}, "title": _t("Kirim (qabul)"),
                                                 "can_create": request.user.has_perm("stores.add_stockreceipt")})

    def get_urls(self):
        v = self.admin_site.admin_view
        return [
            path("new/", v(self.new_view), name="stores_stockreceipt_new"),
            path("new/post/", v(self.post_view), name="stores_stockreceipt_post"),
            path("new/excel/", v(self.excel_view), name="stores_stockreceipt_excel"),
            path("template.xlsx", v(self.template_view), name="stores_stockreceipt_template"),
            *super().get_urls(),
        ]

    def new_view(self, request):
        if not request.user.has_perm("stores.add_stockreceipt"):
            raise PermissionDenied
        locations = list(transfer_service.locations())
        return render(request, "admin/wms/receipt_new.html", {
            **self.admin_site.each_context(request), "opts": self.model._meta, "title": _t("Yangi kirim (qabul)"),
            "config": {
                "locations": [_location_payload(s) for s in locations],
                "default_location": next((str(s.id) for s in locations if s.kind == StoreKind.WAREHOUSE), str(locations[0].id) if locations else None),
                "suppliers": [{"name": n, "terms": d} for n, d in Supplier.objects.filter(is_active=True).values_list("name", "payment_terms_days")],
                "urls": {
                    "lookup": reverse("admin:stores_stocktransfer_lookup"), "search": reverse("admin:stores_stocktransfer_search"),
                    "attach": reverse("admin:stores_stocktransfer_attach_barcode"),
                    "post": reverse("admin:stores_stockreceipt_post"), "excel": reverse("admin:stores_stockreceipt_excel"),
                    "template": reverse("admin:stores_stockreceipt_template"), "list": reverse("admin:stores_stockreceipt_changelist"),
                    "transfer": reverse("admin:stores_stocktransfer_new") if request.user.has_perm("stores.add_stocktransfer") else "",
                },
            },
        })

    def post_view(self, request):
        if request.method != "POST" or not request.user.has_perm("stores.add_stockreceipt"):
            raise PermissionDenied
        b = _body(request)
        store = Store.objects.filter(id=b.get("location")).first()
        if store is None:
            return JsonResponse({"ok": False, "error": {"code": "LOCATION", "message": _t("Qabul qiladigan joyni tanlang.")}}, status=400)
        try:
            r = stock_service.post_receipt(store, b.get("lines") or [], _actor(request), supplier=str(b.get("supplier") or ""),
                                           invoice_number=str(b.get("invoice") or ""), note=str(b.get("note") or ""))
        except DomainException as exc:
            return _error(exc)
        return JsonResponse({"ok": True, "data": {"number": r.number, "units": r.total_units, "value": str(r.total_value),
                                                   "url": reverse("admin:stores_stockreceipt_change", args=[r.pk])}})

    def excel_view(self, request):
        """Upload → parsed lines for review on the page (nothing is posted until the user confirms)."""
        if request.method != "POST" or not request.user.has_perm("stores.add_stockreceipt"):
            raise PermissionDenied
        f = request.FILES.get("file")
        if f is None:
            return JsonResponse({"ok": False, "error": {"message": _t("Faylni tanlang.")}}, status=400)
        try:
            lines, problems = stock_service.parse_receipt_excel(f)
        except Exception:  # noqa: BLE001 - not an xlsx
            return JsonResponse({"ok": False, "error": {"message": _t("Faylni o'qib bo'lmadi. .xlsx formatidagi shablondan foydalaning.")}}, status=400)
        ids = [uuid.UUID(ln["variant_id"]) for ln in lines]
        infos, costs = stock_service.variant_infos(ids), _cost_fields(ids)
        for ln in lines:
            vid = uuid.UUID(ln["variant_id"])
            info = infos.get(vid)
            ln.update({"name": info.name, "sku": info.sku, "barcode": info.barcode, "size": info.size, "image_url": info.image_url,
                       "price": str(info.price) if info.price is not None else None} if info else {})
            ln.update({k: v for k, v in costs[vid].items() if k != "cost"})
            # a cost typed in the file wins; an empty cell takes the latest receipt / catalog cost
            ln["unit_cost"] = str(ln["unit_cost"]) if ln.get("unit_cost") is not None else costs[vid]["prefill_cost"]
        return JsonResponse({"ok": True, "data": {"lines": lines, "problems": problems}})

    def template_view(self, request):
        return excel_response(catalogue_workbook(), f"{_slug()}_kirim_shablon.xlsx")


def catalogue_workbook():
    """
    The goods-receipt template: every active product with its barcode, SKU and size, an empty
    "Soni" column to fill (or scan into), plus cost / batch / expiry -- ready to import back.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    from apps.products.infrastructure.models import ProductVariant

    variants = list(ProductVariant.objects.filter(status="active").order_by("sku"))
    infos = stock_service.variant_infos([v.id for v in variants])
    wb = Workbook()
    ws = wb.active
    ws.title = "Kirim"
    headers = ["Barcode", "Mahsulot nomi", "SKU", "Hajmi", "Soni", "Tannarx (so'm)", "Partiya (LOT)", "Yaroqlilik muddati (YYYY-MM-DD)"]
    ws.append(headers)
    for v in variants:
        i = infos.get(v.id)
        if i is None:
            continue
        ws.append([i.barcode or "", i.name, i.sku, i.size or "", None, None, None, None])
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    thin = Side(style="thin", color="CBD5E1")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    fill_in = PatternFill("solid", fgColor="FEF3C7")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
        row[0].number_format = "@"  # barcodes stay text (no 4.78E+12)
        row[4].fill = fill_in
        row[4].number_format = "0"
        row[5].number_format = "#,##0"
        row[7].number_format = "yyyy-mm-dd"
    for idx, width in enumerate([18, 48, 18, 10, 10, 16, 18, 22], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(ws.max_row, 2)}"
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100000", allow_blank=True,
                        error="Soni 0 dan 100 000 gacha butun son bo'lishi kerak.", errorTitle="Soni noto'g'ri")
    ws.add_data_validation(dv)
    dv.add(f"E2:E{max(ws.max_row, 2) + 500}")

    help_ws = wb.create_sheet("Yo'riqnoma")
    for line in [
        [f"{_brand()} — Sklad kirim shabloni"], [],
        ["1. «Soni» ustuniga kelgan miqdorni yozing (sariq ustun). Bo'sh yoki 0 qatorlar e'tiborga olinmaydi."],
        ["2. Yangi mahsulot bo'lsa, qator qo'shing: faqat Barcode va Soni majburiy."],
        ["3. Tannarx, Partiya va Yaroqlilik muddati ixtiyoriy (vitamin va qo'shimchalar uchun tavsiya etiladi)."],
        ["4. Dashboard → Ombor → Kirim → «Excel'dan yuklash» orqali faylni yuklang, tekshirib «Kirimni tasdiqlash» bosing."],
        ["5. Barcode ustuni matn formatida — Excel uzun shtrix-kodlarni buzmaydi."],
    ]:
        help_ws.append(line)
    help_ws["A1"].font = Font(bold=True, size=14, color="1E3A8A")
    help_ws.column_dimensions["A"].width = 110
    return wb


def excel_response(wb, filename: str) -> HttpResponse:
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ================================================================ ledger


@admin.register(StockMovement)
class StockMovementAdmin(admin.ModelAdmin):
    class Media:
        css = {"all": ("admin/css/nutrizone-wms.css",)}

    list_display = ("when_col", "product_col", "location_col", "type_col", "qty_col", "balance_after", "reserved_after", "reference_col", "actor", "batch_col", "note")
    list_filter = ("movement_type", "store", ("created_at", admin.DateFieldListFilter))
    search_fields = ("reference", "product_variant__sku", "product_variant__barcode", "product_variant__product__translations__name", "actor", "note")
    list_select_related = ("store", "product_variant")
    ordering = ("-created_at",)
    list_per_page = 50
    actions = ["export_xlsx"]
    change_list_template = "admin/wms/ledger_list.html"

    def get_urls(self):
        return [path("export.xlsx", self.admin_site.admin_view(self.export_view), name="stores_stockmovement_export"), *super().get_urls()]

    def export_view(self, request):
        """The current filter / search as Excel (no row selection needed)."""
        cl = self.get_changelist_instance(request)
        return self.export_xlsx(request, cl.queryset)

    def get_changelist_instance(self, request):
        cl = super().get_changelist_instance(request)
        rows = list(cl.result_list)
        infos = self._names(rows)
        for m in rows:
            info = infos.get(m.product_variant_id)
            reserve = m.movement_type in ("reservation", "release")
            link = None
            if m.reference.startswith("TR-") and m.document_id:
                link = reverse("admin:stores_stocktransfer_workspace", args=[m.document_id])
            elif m.reference.startswith("KR-") and m.document_id:
                link = reverse("admin:stores_stockreceipt_change", args=[m.document_id])
            elif m.reference.startswith("IC-") and m.document_id:
                link = reverse("admin:stores_stockcount_workspace", args=[m.document_id])
            elif m.reference.startswith("HC-") and m.document_id:
                link = reverse("admin:stores_writeoff_document", args=[m.document_id])
            elif m.reference.startswith("ORD-"):
                link = f"{reverse('admin:ordering_ordermodel_changelist')}?q={m.reference}"
            m.nz = {
                "name": info.name if info else m.product_variant.sku, "image": info.image_url if info else None, "when": _local(m.created_at),
                "color": MOVEMENT_TONE.get(m.movement_type, "#64748b"), "reserve": reserve,
                "qty": f"{m.quantity:+d}", "cls": "res" if reserve else "pos" if m.quantity > 0 else "neg", "link": link,
            }
        return cl

    def changelist_view(self, request, extra_context=None):
        today = datetime.combine(timezone.localtime(timezone.now(), TZ).date(), time.min, tzinfo=TZ)
        agg = StockMovement.objects.filter(created_at__gte=today).aggregate(
            n=Count("id"), inbound=Sum("quantity", filter=Q(quantity__gt=0) & ~Q(movement_type__in=["reservation", "release"])),
            outbound=Sum("quantity", filter=Q(quantity__lt=0) & ~Q(movement_type__in=["reservation", "release"])),
        )
        current_type = request.GET.get("movement_type__exact", "")
        params = request.GET.copy()
        params.pop("movement_type__exact", None)
        params.pop("p", None)
        base = params.urlencode()
        types = [{"key": "", "label": _t("All"), "color": "#64748b", "on": not current_type, "url": f"?{base}"}]
        for key, label in StockMovementType.choices:
            q = params.copy()
            q["movement_type__exact"] = key
            types.append({"key": key, "label": label, "color": MOVEMENT_TONE.get(key, "#64748b"), "on": current_type == key, "url": f"?{q.urlencode()}"})
        return super().changelist_view(request, {**(extra_context or {}), "title": _t("Ombor tarixi"), "nz_types": types,
                                                 "nz_today": {"n": agg["n"] or 0, "in": agg["inbound"] or 0, "out": -(agg["outbound"] or 0)},
                                                 "nz_locations": list(transfer_service.locations()), "nz_store": request.GET.get("store__id__exact", ""),
                                                 "nz_from": request.GET.get("created_at__gte", ""), "nz_to": request.GET.get("created_at__lt", ""),
                                                 "nz_export": f"{reverse('admin:stores_stockmovement_export')}?{request.GET.urlencode()}"})

    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def get_search_results(self, request, queryset, search_term):
        qs, dup = super().get_search_results(request, queryset, search_term)
        return qs.distinct(), dup

    def _names(self, objs):
        if not hasattr(self, "_cache"):
            self._cache = {}
        missing = [o.product_variant_id for o in objs if o.product_variant_id not in self._cache]
        if missing:
            self._cache.update(stock_service.variant_infos(missing))
        return self._cache

    @admin.display(description=_("When"), ordering="created_at")
    def when_col(self, obj):
        return _local(obj.created_at)

    @admin.display(description=_("Product"))
    def product_col(self, obj):
        info = self._names([obj]).get(obj.product_variant_id)
        return format_html("<strong>{}</strong><br><small class='wms-muted'>{}</small>", info.name if info else "—", obj.product_variant.sku)

    @admin.display(description=_("Location"))
    def location_col(self, obj):
        return format_html('<span class="wms-route"><i class="wms-dot" style="background:{}"></i>{}</span>', obj.store.color or "#94a3b8", obj.store.name)

    @admin.display(description=_("Type"), ordering="movement_type")
    def type_col(self, obj):
        return format_html('<span class="wms-pill" style="--t:{}">{}</span>', MOVEMENT_TONE.get(obj.movement_type, "#64748b"), obj.get_movement_type_display())

    @admin.display(description=_("Quantity"), ordering="quantity")
    def qty_col(self, obj):
        reserve = obj.movement_type in ("reservation", "release")
        color = "#0ea5e9" if reserve else "#16a34a" if obj.quantity > 0 else "#dc2626"
        return format_html('<strong style="color:{}">{}{}</strong>', color, f"{obj.quantity:+d}", " " + _t("band") if reserve else "")

    @admin.display(description=_("Reference"), ordering="reference")
    def reference_col(self, obj):
        if obj.reference.startswith("TR-") and obj.document_id:
            return format_html('<a href="{}">{}</a>', reverse("admin:stores_stocktransfer_workspace", args=[obj.document_id]), obj.reference)
        if obj.reference.startswith("KR-") and obj.document_id:
            return format_html('<a href="{}">{}</a>', reverse("admin:stores_stockreceipt_change", args=[obj.document_id]), obj.reference)
        return obj.reference or "—"

    @admin.display(description=_("Batch"))
    def batch_col(self, obj):
        if not obj.batch_number:
            return "—"
        return format_html("{}<br><small class='wms-muted'>{}</small>", obj.batch_number, obj.expiry_date or "")

    @admin.action(description=_("Tanlanganlarni Excel'ga eksport qilish"))
    def export_xlsx(self, request, queryset):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        rows = list(queryset.select_related("store", "product_variant")[:20000])
        infos = stock_service.variant_infos({r.product_variant_id for r in rows})
        wb = Workbook()
        ws = wb.active
        ws.title = "Ombor tarixi"
        ws.append(["Sana", "Joy", "Mahsulot", "SKU", "Turi", "Miqdor", "Qoldiq", "Band", "Hujjat", "Foydalanuvchi", "Partiya", "Muddati", "Izoh"])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A8A")
        for r in rows:
            ws.append([_local(r.created_at), r.store.name, infos[r.product_variant_id].name if r.product_variant_id in infos else "", r.product_variant.sku,
                       r.get_movement_type_display(), r.quantity, r.balance_after, r.reserved_after, r.reference, r.actor, r.batch_number,
                       r.expiry_date.isoformat() if r.expiry_date else "", r.note])
        return excel_response(wb, f"{_slug()}_ombor_tarixi.xlsx")


# ================================================================ stock by location


LOW_STOCK = 5  # "running low" threshold (units available)


def stock_view(request):
    """
    Ombor → Qoldiqlar: stock per location (on hand / reserved / available, value at cost and at selling price,
    last movement), filters (low / out / reserved), and a matrix of every product across every location.
    """
    from django.db.models import Max

    from apps.stores.infrastructure.count_service import cost_map

    if not request.user.is_staff:
        raise PermissionDenied
    locations = list(transfer_service.locations())
    view = request.GET.get("view", "location")
    current = next((s for s in locations if str(s.id) == request.GET.get("location")), None) or (locations[0] if locations else None)
    q = (request.GET.get("q") or "").strip()
    state = request.GET.get("state", "")

    # per-location cards: lines / units / value at cost
    all_rows = list(StoreProductStock.objects.filter(quantity__gt=0).values("store_id", "product_variant_id", "quantity"))
    costs_all = cost_map({r["product_variant_id"] for r in all_rows})
    loc_stats: dict = {}
    for r in all_rows:
        st = loc_stats.setdefault(str(r["store_id"]), {"lines": 0, "units": 0, "value": Decimal(0)})
        st["lines"] += 1
        st["units"] += r["quantity"]
        st["value"] += (costs_all.get(r["product_variant_id"]) or 0) * r["quantity"]
    cards = [{"s": s, "lines": loc_stats.get(str(s.id), {}).get("lines", 0), "units": loc_stats.get(str(s.id), {}).get("units", 0),
              "value": _som(loc_stats.get(str(s.id), {}).get("value", 0))} for s in locations]
    ctx = {**admin.site.each_context(request), "title": _t("Qoldiqlar"), "locations": locations, "cards": cards, "current": current,
           "q": q, "state": state, "view": view, "low_limit": LOW_STOCK}

    if view == "matrix":
        return _matrix_view(request, ctx, locations, q)

    if q and current and not request.GET.get("export"):
        exact = _exact_variant(q)
        if exact is not None and not StoreProductStock.objects.filter(store=current, product_variant=exact).exists():
            return HttpResponseRedirect(reverse("admin:nz_wms_stock_product", args=[exact.pk]))  # scanned a product this place doesn't hold -> where is it?
    if state == "missing" and current:
        missing = _missing_here(current, q)
        ctx["missing"] = missing
    rows = _stock_rows(current, q=q, limit=3000) if current else []
    costs = cost_map([uuid.UUID(r["variant_id"]) for r in rows])
    last_moves = dict(StockMovement.objects.filter(store=current, product_variant_id__in=[r["variant_id"] for r in rows]).values_list("product_variant_id").annotate(m=Max("created_at"))) if current else {}
    from apps.stores.infrastructure.models import ReorderRule

    minimums = dict(ReorderRule.objects.filter(store=current).values_list("product_variant_id", "min_quantity")) if current else {}
    from apps.stores.infrastructure.models import StockBatch

    batches: dict = {}
    today = timezone.localdate()
    for b in StockBatch.objects.filter(store=current, quantity__gt=0).exclude(batch_number="", expiry_date=None).order_by("expiry_date").values("product_variant_id", "batch_number", "expiry_date", "quantity") if current else []:
        days = (b["expiry_date"] - today).days if b["expiry_date"] else None
        batches.setdefault(b["product_variant_id"], []).append({**b, "tone": "#b91c1c" if days is not None and days < 0 else "#ef4444" if days is not None and days <= 30 else "#f59e0b" if days is not None and days <= 90 else "#6366f1"})
    counts = {"all": len(rows), "low": 0, "out": 0, "reserved": 0}
    for r in rows:
        vid = uuid.UUID(r["variant_id"])
        cost = costs.get(vid) or Decimal(0)
        price = Decimal(r["price"]) if r["price"] else None
        r["cost"] = _som(cost) if cost else ""
        r["cost_value"] = cost * r["quantity"] if cost else None
        r["retail_value"] = price * r["quantity"] if price is not None else None
        r["batches"] = batches.get(vid, [])
        r["min"] = minimums.get(vid)  # the location's reorder point, else the global "running low" threshold
        r["state"] = "out" if r["available"] <= 0 else "low" if r["available"] <= (r["min"] if r["min"] is not None else LOW_STOCK) else "ok"
        r["last"] = _local(last_moves.get(vid), "%d.%m.%Y") or "—"
        counts["low"] += r["state"] == "low"
        counts["out"] += r["state"] == "out"
        counts["reserved"] += r["reserved"] > 0
    counts["missing"] = len(ctx["missing"]) if "missing" in ctx else _missing_count(current)
    if state == "missing":
        rows = []
    elif state in ("low", "out"):
        rows = [r for r in rows if r["state"] == state]
    elif state == "reserved":
        rows = [r for r in rows if r["reserved"] > 0]
    sort = request.GET.get("sort", "")
    if sort == "value":
        rows.sort(key=lambda r: r["cost_value"] or 0, reverse=True)
    elif sort == "available":
        rows.sort(key=lambda r: r["available"])
    elif sort == "name":
        rows.sort(key=lambda r: r["name"])
    if request.GET.get("export") == "xlsx" and current:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = current.name[:30]
        ws.append(["Barcode", "Mahsulot", "SKU", "Hajmi", "Qoldiq", "Band", "Mavjud", "Tannarx", "Qiymat (tannarx)", "Sotuv narxi", "Qiymat (sotuv)", "Oxirgi harakat"])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A8A")
        for r in rows:
            ws.append([r["barcode"], r["name"], r["sku"], r["size"] or "", r["quantity"], r["reserved"], r["available"],
                       float(costs.get(uuid.UUID(r["variant_id"])) or 0) or None, float(r["cost_value"]) if r["cost_value"] is not None else None,
                       float(r["price"]) if r["price"] else None, float(r["retail_value"]) if r["retail_value"] is not None else None, r["last"]])
        return excel_response(wb, f"{_slug()}_qoldiq_{current.slug or 'joy'}.xlsx")
    totals = {
        "lines": counts["all"], "units": sum(r["quantity"] for r in rows), "reserved": sum(r["reserved"] for r in rows),
        "cost_value": _som(sum((r["cost_value"] or 0 for r in rows), Decimal(0))), "retail_value": _som(sum((r["retail_value"] or 0 for r in rows), Decimal(0))),
        "unpriced": sum(1 for r in rows if r["cost_value"] is None),
    }
    for r in rows:
        r["cost_value"] = _som(r["cost_value"]) if r["cost_value"] is not None else ""
        r["retail_value"] = _som(r["retail_value"]) if r["retail_value"] is not None else ""
    return render(request, "admin/wms/stock.html", {**ctx, "rows": rows, "totals": totals, "counts": counts, "sort": sort})


def _exact_variant(code: str):
    """A scanned barcode / typed SKU that names exactly one variant (None for free-text searches)."""
    try:
        return stock_service.find_variant(code)
    except Exception:  # noqa: BLE001 - a search box never errors
        return None


def _catalog_variant_ids(q: str = "") -> list:
    from apps.products.infrastructure.models import ProductVariant

    qs = ProductVariant.objects.filter(status="active", product__is_deleted=False)
    if q:
        from apps.products.infrastructure.barcodes import barcode_contains_q

        qs = qs.filter(Q(sku__icontains=q) | barcode_contains_q(q) | Q(product__translations__name__icontains=q)).distinct()
    return list(qs.values_list("id", flat=True))


def _missing_count(store) -> int:
    if store is None:
        return 0
    held = set(StoreProductStock.objects.filter(store=store, quantity__gt=0).values_list("product_variant_id", flat=True))
    return sum(1 for v in _catalog_variant_ids() if v not in held)


def _missing_here(store, q: str = "") -> list[dict]:
    """Active catalogue variants this location has none of -- with where they are available instead (transfer source)."""
    held = set(StoreProductStock.objects.filter(store=store, quantity__gt=0).values_list("product_variant_id", flat=True))
    ids = [v for v in _catalog_variant_ids(q) if v not in held]
    infos = stock_service.variant_infos(ids)
    elsewhere: dict = {}
    for r in StoreProductStock.objects.filter(product_variant_id__in=ids, quantity__gt=0).exclude(store=store).select_related("store"):
        if r.available > 0:
            elsewhere.setdefault(r.product_variant_id, []).append({"store": r.store, "available": r.available})
    out = []
    for vid in ids:
        info = infos.get(vid)
        if info is None:
            continue
        where = sorted(elsewhere.get(vid, []), key=lambda w: (w["store"].kind != "warehouse", -w["available"]))
        out.append({"info": info, "where": where, "best": where[0]["store"] if where else None, "total": sum(w["available"] for w in where)})
    out.sort(key=lambda m: (-m["total"], m["info"].name))
    return out


def _matrix_view(request, ctx: dict, locations: list, q: str):
    """Every product x every location: available units, coloured against the location's reorder minimum."""
    from apps.stores.infrastructure.models import ReorderRule

    scope = request.GET.get("scope", "")      # "" = held somewhere, "all" = the whole active catalogue
    gap = request.GET.get("gap", "")          # "some" = out in at least one location, "none" = nowhere in stock
    if q and not request.GET.get("export"):
        exact = _exact_variant(q)
        if exact is not None:
            return HttpResponseRedirect(reverse("admin:nz_wms_stock_product", args=[exact.pk]))
    stock = StoreProductStock.objects.filter(Q(quantity__gt=0) | Q(reserved_quantity__gt=0))
    if q:
        from apps.products.infrastructure.barcodes import barcode_contains_q

        stock = stock.filter(Q(product_variant__sku__icontains=q) | barcode_contains_q(q, "product_variant__") | Q(product_variant__product__translations__name__icontains=q)).distinct()
    grid: dict = {}
    for r in stock.values("product_variant_id", "store_id", "quantity", "reserved_quantity"):
        grid.setdefault(r["product_variant_id"], {})[str(r["store_id"])] = (r["quantity"] - r["reserved_quantity"], r["reserved_quantity"])
    if scope == "all" or gap == "none":
        for vid in _catalog_variant_ids(q):
            grid.setdefault(vid, {})
    mins = {(str(s), v): m for s, v, m in ReorderRule.objects.values_list("store_id", "product_variant_id", "min_quantity")}
    infos = stock_service.variant_infos(list(grid))
    matrix = []
    for vid, cells in grid.items():
        info = infos.get(vid)
        if info is None:
            continue
        row_cells = []
        for s in locations:
            avail, reserved = cells.get(str(s.id), (0, 0))
            low_at = mins.get((str(s.id), vid))
            tone = "zero" if avail <= 0 else "low" if avail <= (low_at if low_at is not None else LOW_STOCK) else "ok"
            row_cells.append({"v": avail, "res": reserved, "tone": tone, "min": low_at, "store": s})
        total = sum(c["v"] for c in row_cells)
        branches = [c for c in row_cells if c["store"].kind != "warehouse"]
        out_in = sum(1 for c in branches if c["v"] <= 0)
        matrix.append({"info": info, "cells": row_cells, "total": total, "out_in": out_in, "covered": sum(1 for c in row_cells if c["v"] > 0)})
    if gap == "some":
        matrix = [m for m in matrix if m["out_in"] and m["total"] > 0]
    elif gap == "none":
        matrix = [m for m in matrix if m["total"] <= 0]
    matrix.sort(key=lambda m: m["info"].name)
    if request.GET.get("export") == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Matritsa"
        ws.append(["Mahsulot", "SKU", "Shtrix-kod", *[s.name for s in locations], "Jami"])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A8A")
        fills = {"zero": PatternFill("solid", fgColor="FDE2E4"), "low": PatternFill("solid", fgColor="FEF3C7")}
        for m in matrix:
            ws.append([m["info"].name, m["info"].sku, m["info"].barcode, *[c["v"] for c in m["cells"]], m["total"]])
            for i, c in enumerate(m["cells"]):
                if c["tone"] in fills:
                    ws.cell(row=ws.max_row, column=4 + i).fill = fills[c["tone"]]
        return excel_response(wb, f"{_slug()}_qoldiq_matritsa.xlsx")
    foot = [{"units": sum(m["cells"][i]["v"] for m in matrix), "skus": sum(1 for m in matrix if m["cells"][i]["v"] > 0),
             "out": sum(1 for m in matrix if m["cells"][i]["v"] <= 0)} for i in range(len(locations))]
    summary = {
        "skus": len(matrix), "units": sum(m["total"] for m in matrix),
        "gaps": sum(1 for m in matrix if m["out_in"] and m["total"] > 0), "nowhere": sum(1 for m in matrix if m["total"] <= 0),
    }
    return render(request, "admin/wms/stock.html", {**ctx, "matrix": matrix, "foot": foot, "scope": scope, "gap": gap, "msum": summary})


def stock_product_view(request, variant_id):
    """
    Ombor → one product: where it is and how much -- every location's on hand / reserved / available against its
    reorder minimum, batches with expiry, last movement and 30-day sales there, the recent ledger, and one-click
    transfers from the best-stocked location.
    """
    from datetime import timedelta as _td

    from django.db.models import Max

    from apps.products.infrastructure.models import ProductVariant
    from apps.stores.infrastructure.count_service import cost_info
    from apps.stores.infrastructure.models import ReorderRule, StockBatch

    if not request.user.is_staff:
        raise PermissionDenied
    variant = get_object_or_404(ProductVariant._base_manager, pk=variant_id)  # a deleted product may still hold stock
    info = stock_service.variant_infos([variant.pk]).get(variant.pk)
    locations = list(transfer_service.locations())
    stock = {r.store_id: r for r in StoreProductStock.objects.filter(product_variant=variant)}
    mins = dict(ReorderRule.objects.filter(product_variant=variant).values_list("store_id", "min_quantity"))
    targets = dict(ReorderRule.objects.filter(product_variant=variant).values_list("store_id", "target_quantity"))
    last = dict(StockMovement.objects.filter(product_variant=variant).values_list("store_id").annotate(m=Max("created_at")))
    since = timezone.now() - _td(days=30)
    sold: dict = {}
    for store_id, qty in StockMovement.objects.filter(product_variant=variant, movement_type="sale", created_at__gte=since).values_list("store_id", "quantity"):
        sold[store_id] = sold.get(store_id, 0) + abs(qty)
    today = timezone.localdate()
    batches: dict = {}
    for b in StockBatch.objects.filter(product_variant=variant, quantity__gt=0).exclude(batch_number="", expiry_date=None).order_by("expiry_date").values("store_id", "batch_number", "expiry_date", "quantity"):
        days = (b["expiry_date"] - today).days if b["expiry_date"] else None
        b["tone"] = "#b91c1c" if days is not None and days < 0 else "#ef4444" if days is not None and days <= 30 else "#f59e0b" if days is not None and days <= 90 else "#6366f1"
        b["days"] = days
        batches.setdefault(b["store_id"], []).append(b)
    cost, cost_src = cost_info([variant.pk]).get(variant.pk, (None, None))
    if cost_src == "price":
        cost = None  # valued at the selling price would look like a real cost: show "no cost" and ask for a receipt cost instead
    rows = []
    for s in locations:
        r = stock.get(s.id)
        qty, res = (r.quantity, r.reserved_quantity) if r else (0, 0)
        avail = qty - res
        mn = mins.get(s.id)
        state = "out" if avail <= 0 else "low" if avail <= (mn if mn is not None else LOW_STOCK) else "ok"
        need = max((targets.get(s.id) or 0) - avail, 0) if mn is not None and avail <= mn else 0
        daily = sold.get(s.id, 0) / 30
        rows.append({
            "s": s, "quantity": qty, "reserved": res, "available": avail, "min": mn, "target": targets.get(s.id), "state": state, "need": need,
            "sold": sold.get(s.id, 0), "days_left": int(avail / daily) if daily and avail > 0 else None,
            "last": _local(last.get(s.id), "%d.%m.%Y") or "—", "batches": batches.get(s.id, []),
            "value": _som(cost * qty) if cost and qty else "",
        })
    peak = max([r["available"] for r in rows] + [1])
    for r in rows:
        r["pct"] = max(int(r["available"] * 100 / peak), 2) if r["available"] > 0 else 0
    donors = sorted([r for r in rows if r["available"] > 0], key=lambda r: (r["s"].kind != "warehouse", -r["available"]))
    for r in rows:
        donor = next((d for d in donors if d["s"].pk != r["s"].pk), None)
        r["donor"] = donor["s"] if donor and r["state"] != "ok" else None
    labels = dict(StockMovement._meta.get_field("movement_type").choices)
    moves = list(StockMovement.objects.filter(product_variant=variant).select_related("store").order_by("-created_at")[:12])
    for m in moves:
        m.label, m.when = labels.get(m.movement_type, m.movement_type), _local(m.created_at)
    siblings = list(ProductVariant.objects.filter(product_id=variant.product_id).exclude(pk=variant.pk).order_by("-is_default", "sku"))
    sib_stock = dict(StoreProductStock.objects.filter(product_variant__in=siblings).values_list("product_variant_id").annotate(t=Sum("quantity")))
    totals = {
        "quantity": sum(r["quantity"] for r in rows), "reserved": sum(r["reserved"] for r in rows), "available": sum(r["available"] for r in rows),
        "covered": sum(1 for r in rows if r["available"] > 0), "locations": len(rows), "sold": sum(sold.values()),
        "value": _som(cost * sum(r["quantity"] for r in rows)) if cost else "", "cost": _som(cost) if cost else "",
        "cost_src": {"receipt": _t("oxirgi kirim narxi"), "catalog": _t("katalogdagi tannarx")}.get(cost_src or "", ""),
        "margin": round((info.price - cost) * 100 / info.price, 1) if cost and info and info.price else None,
        "low": sum(1 for r in rows if r["state"] == "low"), "out": sum(1 for r in rows if r["state"] == "out" and r["s"].kind != "warehouse"),
    }
    daily_all = totals["sold"] / 30
    totals["days_left"] = int(totals["available"] / daily_all) if daily_all and totals["available"] > 0 else None
    ctx = {
        **admin.site.each_context(request), "title": info.name if info else variant.sku, "variant": variant, "info": info, "rows": rows,
        "totals": totals, "moves": moves, "siblings": [{"v": v, "units": sib_stock.get(v.pk, 0)} for v in siblings],
        "low_limit": LOW_STOCK,
    }
    return render(request, "admin/wms/stock_product.html", ctx)


def install_urls() -> None:
    """Adds the stock page and redirects the old transfer page to the new WMS list (idempotent)."""
    if getattr(admin.site, "_nz_wms_installed", False):
        return
    original = admin.site.get_urls

    def get_urls():
        from django.views.generic import RedirectView

        return [
            path("wms/stock/", admin.site.admin_view(stock_view), name="nz_wms_stock"),
            path("wms/stock/product/<uuid:variant_id>/", admin.site.admin_view(stock_product_view), name="nz_wms_stock_product"),
            path("inventory/transfers/", RedirectView.as_view(pattern_name="admin:stores_stocktransfer_changelist", permanent=False), name="nz_wms_old_transfers"),
            *original(),
        ]

    admin.site.get_urls = get_urls
    admin.site._nz_wms_installed = True
