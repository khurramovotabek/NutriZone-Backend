"""
Dashboard → Ombor → To'ldirish (reorder points / replenishment) and Yaroqlilik muddati (expiry, FEFO).
Business rules live in planning_service.py.

To'ldirish   the ReorderRule changelist is a planning page: location cards (how many products need stock),
             per product available / incoming / sales speed / days of cover / min / fill-up-to / suggestion,
             inline editing of the rules (JSON save), rule suggestions from sales, and "create transfer draft".
Expiry       estimated batches on the shelf by expiry bucket, FEFO advice, write-off shortcut, Excel.
"""
from __future__ import annotations

import uuid
from decimal import Decimal

from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.http import HttpResponseRedirect, JsonResponse
from django.shortcuts import render
from django.urls import path, reverse
from django.utils import timezone
from django.utils.translation import gettext as _t

from apps.stores.infrastructure import planning_service as plan
from apps.stores.infrastructure import stock_service, transfer_service
from apps.stores.infrastructure.models import ReorderRule, StoreKind
from apps.stores.infrastructure.warehouse_admin import _actor, _body, _error, _slug, _som, excel_response
from common.exceptions import DomainException


@admin.register(ReorderRule)
class ReorderRuleAdmin(admin.ModelAdmin):
    """Only the planning page -- rules are edited inline there, never through a raw form."""

    def has_add_permission(self, request):
        return False

    def change_view(self, request, object_id, form_url="", extra_context=None):
        return HttpResponseRedirect(reverse("admin:stores_reorderrule_changelist"))

    def get_urls(self):
        v = self.admin_site.admin_view
        return [
            path("save/", v(self.save_view), name="stores_reorderrule_save"),
            path("transfer/", v(self.transfer_view), name="stores_reorderrule_transfer"),
            *super().get_urls(),
        ]

    def changelist_view(self, request, extra_context=None):
        if not self.has_view_permission(request):
            raise PermissionDenied
        locations = list(transfer_service.locations())
        branches = [s for s in locations if s.kind != StoreKind.WAREHOUSE] or locations
        store = next((s for s in locations if str(s.id) == request.GET.get("location")), None) or (branches[0] if branches else None)
        sources = [s for s in locations if store is None or s.id != store.id]
        source = next((s for s in sources if str(s.id) == request.GET.get("source")), None) or plan.default_source(exclude=store) or (sources[0] if sources else None)
        overview = plan.overview()
        tab = request.GET.get("tab") or "need"
        ctx = {**self.admin_site.each_context(request), "opts": self.model._meta, "title": _t("To'ldirish (minimal qoldiq)"),
               "locations": [{"s": s, **overview.get(s.id, {"rules": 0, "need": 0, "out": 0})} for s in locations], "store": store, "source": source,
               "sources": sources, "tab": tab, "can_edit": request.user.has_perm("stores.change_reorderrule"),
               "can_transfer": request.user.has_perm("stores.add_stocktransfer"),
               "days": {"sales": plan.SALES_DAYS, "safety": plan.SAFETY_DAYS, "cover": plan.COVER_DAYS}}
        if store is None:
            return render(request, "admin/wms/replenish.html", ctx)
        p = plan.plan(store, source)
        rows = []
        for ln in p["lines"]:
            low, target = plan.suggested_rule(ln.sold) if ln.sold else (None, None)
            rows.append({"ln": ln, "sug_min": low, "sug_target": target, "per_day": f"{ln.per_day:.1f}".rstrip("0").rstrip(".") if ln.sold else ""})
        shown = {"need": [r for r in rows if r["ln"].needs], "rules": [r for r in rows if r["ln"].rule], "norule": [r for r in rows if not r["ln"].rule]}.get(tab, rows)
        ctx.update({
            "rows": shown, "counts": {"need": len(p["need"]), "rules": sum(1 for r in rows if r["ln"].rule), "norule": len(p["norule"])},
            "send_units": sum(r["ln"].can_send for r in rows if r["ln"].needs), "short_units": sum(r["ln"].suggest - r["ln"].can_send for r in rows if r["ln"].needs),
            "search_url": reverse("admin:stores_stocktransfer_search"),
        })
        return render(request, "admin/wms/replenish.html", ctx)

    def save_view(self, request):
        if request.method != "POST" or not request.user.has_perm("stores.change_reorderrule"):
            raise PermissionDenied
        b = _body(request)
        try:
            n = plan.save_rules(plan.location(b.get("location")), _actor(request), b.get("rules") or {})
        except DomainException as exc:
            return _error(exc)
        return JsonResponse({"ok": True, "data": {"saved": n}})

    def transfer_view(self, request):
        if request.method != "POST" or not request.user.has_perm("stores.add_stocktransfer"):
            raise PermissionDenied
        back = reverse("admin:stores_reorderrule_changelist") + f"?location={request.POST.get('location', '')}&source={request.POST.get('source', '')}"
        qty = {}
        for key, value in request.POST.items():
            if key.startswith("qty_") and request.POST.get("pick_" + key[4:]):
                try:
                    qty[uuid.UUID(key[4:])] = int(value or 0)
                except ValueError:
                    continue
        try:
            t = plan.create_transfer(plan.location(request.POST.get("location")), plan.location(request.POST.get("source")), _actor(request), qty)
        except DomainException as exc:
            messages.error(request, str(exc))
            return HttpResponseRedirect(back)
        messages.success(request, _t("Transfer draft %s created -- check it and mark it ready.") % t.number)
        return HttpResponseRedirect(reverse("admin:stores_stocktransfer_workspace", args=[t.pk]))


# ================================================================ expiry


BUCKETS = (
    ("expired", "#b91c1c", "fa-skull-crossbones"), ("d30", "#ef4444", "fa-hourglass-end"),
    ("d60", "#f59e0b", "fa-hourglass-half"), ("d90", "#eab308", "fa-hourglass-start"), ("ok", "#10b981", "fa-circle-check"),
)


def expiry_view(request):
    """Ombor → Yaroqlilik muddati: what expires soon where, the batch to sell first, and expired stock to write off."""
    if not request.user.has_perm("stores.view_stockmovement"):
        raise PermissionDenied
    from apps.stores.infrastructure.count_service import cost_map

    locations = list(transfer_service.locations())
    current = next((s for s in locations if str(s.id) == request.GET.get("location")), None)
    report = plan.expiry_report([current.id] if current else None)
    today = report["today"]
    names = {s.id: s for s in locations}
    infos = stock_service.variant_infos({ly.variant_id for ly in report["layers"]})
    costs = cost_map({ly.variant_id for ly in report["layers"]})
    q = (request.GET.get("q") or "").strip().lower()
    labels = {"expired": _t("Expired"), "d30": _t("≤ 30 days"), "d60": _t("31–60 days"), "d90": _t("61–90 days"), "ok": _t("More than 90 days")}
    stats = {key: {"key": key, "label": labels[key], "color": color, "icon": icon, "units": 0, "value": Decimal(0), "lines": 0} for key, color, icon in BUCKETS}
    rows = []
    for ly in report["layers"]:
        info = infos.get(ly.variant_id)
        store = names.get(ly.store_id)
        if info is None or store is None:
            continue
        days = ly.days_left(today)
        key = plan.bucket(days)
        value = (costs.get(ly.variant_id) or 0) * ly.quantity
        st = stats[key]
        st["units"] += ly.quantity
        st["value"] += value
        st["lines"] += 1
        if q and q not in info.name.lower() and q not in info.sku.lower() and q not in (info.barcode or "") and q not in ly.batch.lower():
            continue
        first = report["fefo"].get((ly.store_id, ly.variant_id))
        rows.append({"info": info, "store": store, "ly": ly, "days": days, "ago": abs(days), "key": key, "color": stats[key]["color"], "value": _som(value) if value else "",
                     "fefo_first": bool(first) and first.expiry == ly.expiry, "fefo_later": bool(first) and first.expiry != ly.expiry,
                     "writeoff": reverse("admin:stores_writeoff_new") + f"?store={ly.store_id}&variant={ly.variant_id}&qty={ly.quantity}&reason=expired&batch={ly.batch}&expiry={ly.expiry.isoformat()}"})
    tab = request.GET.get("tab") or "soon"
    if tab == "soon":
        rows = [r for r in rows if r["key"] in ("expired", "d30", "d60", "d90")]
    elif tab in stats:
        rows = [r for r in rows if r["key"] == tab]
    rows.sort(key=lambda r: (r["ly"].expiry, r["info"].name))
    fefo = []
    for (sid, vid), first in report["fefo"].items():
        info, store = infos.get(vid), names.get(sid)
        if info is None or store is None:
            continue
        group = sorted((ly for ly in report["layers"] if ly.store_id == sid and ly.variant_id == vid), key=lambda ly: ly.expiry)
        fefo.append({"info": info, "store": store, "first": first, "group": group, "days": first.days_left(today)})
    fefo.sort(key=lambda f: f["first"].expiry)

    if request.GET.get("export") == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Yaroqlilik"
        ws.append(["Mahsulot", "SKU", "Shtrix-kod", "Joy", "Partiya", "Yaroqlilik", "Qolgan kun", "Taxminiy soni", "Qiymat (tannarx)"])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A8A")
        for r in rows:
            ws.append([r["info"].name, r["info"].sku, r["info"].barcode, r["store"].name, r["ly"].batch, r["ly"].expiry, r["days"], r["ly"].quantity,
                       float((costs.get(r["ly"].variant_id) or 0) * r["ly"].quantity) or None])
        for col, width in zip("ABCDEFGHI", (44, 16, 16, 22, 14, 12, 11, 13, 16)):
            ws.column_dimensions[col].width = width
        return excel_response(wb, f"{_slug()}_yaroqlilik_muddati.xlsx")

    for st in stats.values():
        st["value"] = _som(st["value"])
    return render(request, "admin/wms/expiry.html", {
        **admin.site.each_context(request), "title": _t("Yaroqlilik muddati"), "locations": locations, "current": current, "rows": rows[:500], "total_rows": len(rows),
        "stats": list(stats.values()), "tab": tab, "q": request.GET.get("q", ""), "fefo": fefo[:40], "unknown_units": report["unknown_units"],
        "unknown_products": report["unknown_products"], "today": _local_date(today), "can_writeoff": request.user.has_perm("stores.add_writeoff"),
    })


def _local_date(d) -> str:
    return d.strftime("%d.%m.%Y")



# ================================================================ reports


def reports_view(request):
    """Ombor → Hisobotlar: turnover, days of inventory, dead stock, ABC, flow by movement type (report_service.py)."""
    if not request.user.has_perm("stores.view_stockmovement"):
        raise PermissionDenied
    from apps.stores.infrastructure import report_service
    from apps.stores.infrastructure.models import StockMovementType

    locations = list(transfer_service.locations())
    current = next((s for s in locations if str(s.id) == request.GET.get("location")), None)
    try:
        days = int(request.GET.get("days") or 30)
    except ValueError:
        days = 30
    days = days if days in (7, 30, 90, 180, 365) else 30
    rep = report_service.build([current.id] if current else None, days)
    tab = request.GET.get("tab") or "all"
    rows = rep["rows"]
    if tab == "dead":
        rows = sorted(rep["dead"], key=lambda r: r.value, reverse=True)
    elif tab in ("A", "B", "C"):
        rows = sorted((r for r in rows if r.sold and r.abc == tab), key=lambda r: r.sales_value or r.sold, reverse=True)
    elif tab == "fast":
        rows = sorted((r for r in rows if r.sold), key=lambda r: r.sold, reverse=True)[:50]
    else:
        rows = sorted(rows, key=lambda r: (r.sales_value, r.sold, r.value), reverse=True)
    if request.GET.get("export") == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Hisobot"
        ws.append(["Mahsulot", "SKU", "Qoldiq", f"Sotildi ({days} kun)", "Kunlik sotuv", "Necha kunga yetadi", "Oxirgi sotuv", "Tannarx", "Qoldiq qiymati", "Sotuv (tannarxda)", "ABC"])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A8A")
        for r in rows:
            ws.append([r.info.name, r.info.sku, r.on_hand, r.sold, round(r.per_day, 2), r.cover, timezone.localtime(r.last_sale).strftime("%d.%m.%Y") if r.last_sale else "",
                       float(r.cost) or None, float(r.value) or None, float(r.sales_value) or None, r.abc if r.sold else ""])
        for col, width in zip("ABCDEFGHIJK", (44, 16, 9, 12, 12, 14, 12, 12, 15, 15, 6)):
            ws.column_dimensions[col].width = width
        return excel_response(wb, f"{_slug()}_ombor_hisoboti_{days}kun.xlsx")
    labels = dict(StockMovementType.choices)
    tones = {"receipt": "#10b981", "sale": "#3b82f6", "return": "#a855f7", "transfer_out": "#f97316", "transfer_in": "#14b8a6", "adjustment": "#64748b",
             "damage": "#ef4444", "expired": "#b91c1c", "write_off": "#be123c"}

    def flow(d):
        total = sum(v["units"] for v in d.values()) or 1
        return [{"label": labels.get(k, k), "color": tones.get(k, "#64748b"), "units": v["units"], "value": _som(v["value"]), "share": round(v["units"] * 100 / total)}
                for k, v in sorted(d.items(), key=lambda kv: -kv[1]["units"])]
    today = timezone.localdate()
    out_rows = [{"r": r, "value": _som(r.value) if r.value else "", "sales_value": _som(r.sales_value) if r.sales_value else "",
                 "per_day": f"{r.per_day:.1f}".rstrip("0").rstrip(".") if r.sold else "", "last": timezone.localtime(r.last_sale).strftime("%d.%m.%Y") if r.last_sale else "",
                 "idle": (today - timezone.localtime(r.last_sale).date()).days if r.last_sale else None} for r in rows[:400]]
    return render(request, "admin/wms/reports.html", {
        **admin.site.each_context(request), "title": _t("Ombor hisobotlari"), "locations": locations, "current": current, "days": days, "tab": tab,
        "rep": rep, "rows": out_rows, "total_rows": len(rows), "flow_in": flow(rep["flow_in"]), "flow_out": flow(rep["flow_out"]),
        "k": {"stock_value": _som(rep["stock_value"]), "cogs": _som(rep["cogs"]), "dead_value": _som(rep["dead_value"]), "loss_value": _som(rep["loss_value"]),
              "abc": {k: {"n": v["n"], "value": _som(v["value"])} for k, v in rep["abc"].items()}},
        "periods": (7, 30, 90, 180, 365),
    })
