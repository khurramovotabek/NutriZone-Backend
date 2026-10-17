"""
Location stock -- the ONE way stock at a branch or warehouse changes.

Every change (goods receipt, sale, return, transfer out/in, adjustment, reservation)
goes through `move()`, which in a single transaction:

  * locks the location's stock row (select_for_update -- concurrent tills / transfers
    queue up instead of overselling),
  * refuses anything that would make stock negative or reserve more than is on hand,
  * writes an immutable StockMovement ledger row with the balance after the change,
  * mirrors on-hand changes into the inventory module's per-location rows, which the
    storefront reads for "available in stores" (best effort, logged, never blocks).

Sellable / transferable stock is `quantity − reserved_quantity`.
"""
from __future__ import annotations

import logging
import uuid
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal

from django.db import connection, transaction
from django.db.models import F, Q
from django.utils import timezone

from apps.stores.infrastructure.models import (
    StockMovement,
    StockMovementType,
    StockReceipt,
    StockReceiptItem,
    Store,
    StoreProductStock,
)
from common.exceptions import ConflictError, ResourceNotFound, ValidationFailed

logger = logging.getLogger("nutrizone.app")
MAX_UNITS_PER_LINE = 100_000


def conflict(code: str, message: str, **details) -> ConflictError:
    err = ConflictError(message, details=details)
    err.code = code
    return err


def next_number(prefix: str, sequence: str) -> str:
    with connection.cursor() as cursor:
        cursor.execute(f"SELECT nextval('{sequence}')")
        return f"{prefix}-{cursor.fetchone()[0]:06d}"


# ---------------------------------------------------------------- products


@dataclass(frozen=True)
class VariantInfo:
    id: uuid.UUID
    product_id: uuid.UUID
    name: str
    sku: str
    barcode: str
    size: str | None
    image_url: str | None
    price: Decimal | None
    active: bool


def variant_infos(variant_ids, locale: str = "uz") -> dict:
    from apps.products.infrastructure.models import ProductMedia, ProductTranslation, ProductVariant
    from apps.products.infrastructure.selectors import bulk_variant_prices

    from apps.products.infrastructure.variant_labels import variant_label

    variants = list(ProductVariant.objects.filter(id__in=list(variant_ids)).select_related("package", "flavour").prefetch_related("flavour__translations"))
    if not variants:
        return {}
    pids = {v.product_id for v in variants}
    names: dict = {}
    for tr in ProductTranslation.objects.filter(product_id__in=pids).values("product_id", "locale_code", "name"):
        names.setdefault(tr["product_id"], {})[tr["locale_code"]] = tr["name"]
    images: dict = {}
    for m in ProductMedia.objects.filter(product_id__in=pids).order_by("product_id", "sort_order").values("product_id", "storage_path", "thumbnail_paths"):
        images.setdefault(m["product_id"], f"/media/{(m['thumbnail_paths'] or {}).get('small') or m['storage_path']}")
    prices = bulk_variant_prices([v.id for v in variants])
    out = {}
    for v in variants:
        by = names.get(v.product_id, {})
        size = variant_label(v, locale, package=False)  # "2.27 kg · Chocolate"
        price = prices.get(v.id)
        out[v.id] = VariantInfo(
            id=v.id, product_id=v.product_id, name=by.get(locale) or by.get("uz") or by.get("en") or next(iter(by.values()), v.sku),
            sku=v.sku, barcode=v.barcode or "", size=size, image_url=images.get(v.product_id),
            price=Decimal(price["amount"]) if price else None, active=v.status == "active",
        )
    return out


def _barcodes():
    from apps.products.infrastructure import barcodes

    return barcodes


def find_variant(code: str):
    """A scanned barcode (main or an extra batch barcode) or a typed SKU -> the product variant (None if nothing matches)."""
    match = find_match(code)
    return match.variant if match else None


def find_match(code: str):
    """Like find_variant, but also says which barcode matched and the batch/expiry tied to it (see products.barcodes)."""
    from apps.products.infrastructure import barcodes

    return barcodes.resolve(code)


def search_variants(q: str, limit: int = 20):
    from apps.products.infrastructure.models import ProductVariant

    q = (q or "").strip()
    if not q:
        return []
    exact = find_variant(q)
    qs = ProductVariant.objects.filter(
        Q(sku__icontains=q) | _barcodes().barcode_contains_q(q) | Q(product__translations__name__icontains=q)
    ).distinct().order_by("sku")[:limit]
    rows = list(qs)
    if exact and exact not in rows:
        rows.insert(0, exact)
    return rows[:limit]


# ---------------------------------------------------------------- the one stock mutation


def _mirror_inventory(store_id, variant_id, delta: int, reference: str) -> None:
    """Keep the inventory module's per-location row (read by the storefront) in step with this location's stock."""
    if not delta:
        return
    try:
        from django.db.models.functions import Greatest

        from apps.inventory.infrastructure.models import InventoryModel, WarehouseInventoryModel

        inventory = InventoryModel.objects.filter(product_variant_id=variant_id).first()
        if inventory is None:
            inventory = InventoryModel.objects.create(code=f"INV-{uuid.uuid4().hex[:8].upper()}", product_variant_id=variant_id)
        row = WarehouseInventoryModel.objects.filter(inventory=inventory, warehouse_id=store_id).order_by("id").first()
        if row is None:
            WarehouseInventoryModel.objects.create(inventory=inventory, warehouse_id=store_id, available_quantity=max(0, delta))
        else:
            WarehouseInventoryModel.objects.filter(pk=row.pk).update(available_quantity=Greatest(F("available_quantity") + delta, 0))
    except Exception:  # noqa: BLE001 - the storefront mirror must never block a stock operation
        logger.exception("stock mirror failed (%s, variant %s, %+d)", reference, variant_id, delta)


# ---------------------------------------------------------------- batches (partiya)


def _batches_locked(store, variant_id):
    from apps.stores.infrastructure.models import StockBatch

    return list(StockBatch.objects.select_for_update().filter(store=store, product_variant_id=variant_id))


def _fefo(rows):
    """Outgoing order: the batch that expires first, then batches without an expiry, oldest arrival first."""
    return sorted((b for b in rows if b.quantity > 0), key=lambda b: (b.expiry_date is None, b.expiry_date or date.max, b.received_at))


def _batch(store, variant_id, batch_number: str, expiry_date, rows: list):
    from apps.stores.infrastructure.models import StockBatch

    found = next((b for b in rows if b.batch_number == batch_number and b.expiry_date == expiry_date), None)
    if found is None:
        found, _ = StockBatch.objects.get_or_create(store=store, product_variant_id=variant_id, batch_number=batch_number, expiry_date=expiry_date)
        found = StockBatch.objects.select_for_update().get(pk=found.pk)
        rows.append(found)
    return found


SELLABLE_FIRST = (StockMovementType.SALE, StockMovementType.TRANSFER_OUT)


def _take(rows, units: int, prefer: tuple | None = None, sellable_first: bool = False) -> list:
    """
    Takes `units` out of the batches (the preferred batch first, then FEFO); returns [(batch, n)].
    sellable_first: a sale / transfer takes unexpired batches before expired ones (those are to be written off).
    """
    order = _fefo(rows)
    if sellable_first:
        today = date.today()
        order.sort(key=lambda b: bool(b.expiry_date and b.expiry_date < today))
    if prefer is not None:
        order.sort(key=lambda b: (b.batch_number, b.expiry_date) != prefer)
    out = []
    for b in order:
        if units <= 0:
            break
        n = min(units, b.quantity)
        b.quantity -= n
        units -= n
        out.append((b, n))
    return out


def _apply_batches(store, variant_id, before: int, quantity: int, movement_type: str, batch_number: str, expiry_date, placement: list | None) -> list:
    """
    Keeps the product's batches at this location in step with an on-hand change of `quantity` and returns the
    allocations [(batch, signed units)]. First heals any drift (batches must add up to the on-hand `before`).
    """
    from apps.stores.infrastructure.models import StockBatch

    rows = _batches_locked(store, variant_id)
    total = sum(b.quantity for b in rows)
    if total < before:  # stock that predates batch tracking (or a change made around move()): "no batch"
        nb = _batch(store, variant_id, "", None, rows)
        nb.quantity += before - total
    elif total > before:
        _take(rows, total - before)
    alloc: list = []
    if quantity > 0:
        parts = [(str(p.get("batch") or "")[:48], p.get("expiry"), int(p.get("qty") or 0)) for p in (placement or []) if int(p.get("qty") or 0) > 0]
        left = quantity
        for batch, expiry, qty in parts:
            qty = min(qty, left)
            if qty <= 0:
                continue
            b = _batch(store, variant_id, batch, _parse_date(expiry), rows)
            b.quantity += qty
            alloc.append((b, qty))
            left -= qty
        if left > 0:
            b = _batch(store, variant_id, (batch_number or "")[:48], expiry_date, rows)
            b.quantity += left
            alloc.append((b, left))
    elif quantity < 0:
        prefer = ((batch_number or "")[:48], expiry_date) if (batch_number or expiry_date) else None
        alloc = [(b, -n) for b, n in _take(rows, -quantity, prefer, sellable_first=movement_type in SELLABLE_FIRST)]
    for b in rows:
        StockBatch.objects.filter(pk=b.pk).update(quantity=b.quantity)
    return alloc


def return_placement(store, variant_id, *, reference: str = "", document_id=None) -> list:
    """
    A customer return goes back into the batches the sale took the units from (minus what was already returned):
    [{batch, expiry, qty}] for move(..., placement=...).
    """
    from apps.stores.infrastructure.models import StockMovementBatch

    q = StockMovementBatch.objects.filter(movement__store=store, movement__product_variant_id=variant_id)
    if document_id:
        q = q.filter(movement__document_id=document_id)
    elif reference:
        q = q.filter(movement__reference=reference)
    else:
        return []
    sold: dict = {}
    for mb in q.filter(movement__movement_type__in=(StockMovementType.SALE, StockMovementType.RETURN)).select_related("batch"):
        key = (mb.batch.batch_number, mb.batch.expiry_date)
        sold[key] = sold.get(key, 0) - mb.quantity  # sales are negative, returns positive
    return [{"batch": k[0], "expiry": k[1], "qty": n} for k, n in sorted(sold.items(), key=lambda kv: (kv[0][1] is None, kv[0][1] or date.max)) if n > 0]


def move(
    store: Store, variant_id, quantity: int, movement_type: str, *, reference: str = "", actor: str = "", note: str = "",
    reserve: int = 0, batch_number: str = "", expiry_date: date | None = None, document_id=None, mirror: bool = True,
    placement: list | None = None,
) -> StockMovement:
    """
    Changes a location's on-hand `quantity` by `quantity` and/or its reserved units by `reserve`,
    atomically, never below zero, and records it in the ledger. Must run inside a transaction.
    Batches: incoming units go to `batch_number`/`expiry_date` (or split by `placement` = [{batch, expiry, qty}]);
    outgoing units leave the given batch first, then FEFO. The allocations are on `movement.allocations`.
    """
    entry, _ = StoreProductStock.objects.get_or_create(store=store, product_variant_id=variant_id, defaults={"quantity": 0})
    entry = StoreProductStock.objects.select_for_update().get(pk=entry.pk)
    new_qty = entry.quantity + quantity
    new_reserved = entry.reserved_quantity + reserve
    if new_qty < 0:
        raise conflict("INSUFFICIENT_STOCK", f"Only {entry.quantity - entry.reserved_quantity} available.", variant_id=str(variant_id), available=max(0, entry.quantity - entry.reserved_quantity))
    if new_reserved < 0 or new_reserved > new_qty:
        raise conflict("INSUFFICIENT_STOCK", f"Only {entry.quantity - entry.reserved_quantity} available.", variant_id=str(variant_id), available=max(0, entry.quantity - entry.reserved_quantity))
    StoreProductStock.objects.filter(pk=entry.pk).update(quantity=new_qty, reserved_quantity=new_reserved)
    if movement_type == StockMovementType.RETURN and quantity > 0 and not placement and not batch_number:
        placement = return_placement(store, variant_id, reference=reference, document_id=document_id)
    alloc = _apply_batches(store, variant_id, entry.quantity, quantity, movement_type, batch_number, expiry_date, placement) if quantity else []
    if quantity < 0 and len(alloc) == 1 and not batch_number:
        batch_number, expiry_date = alloc[0][0].batch_number, alloc[0][0].expiry_date  # one batch: show it on the ledger row
    movement = StockMovement.objects.create(
        store=store, product_variant_id=variant_id, movement_type=movement_type, quantity=quantity if quantity else reserve,
        balance_after=new_qty, reserved_after=new_reserved, reference=reference[:40], actor=actor[:150], note=note[:255],
        batch_number=(batch_number or "")[:48], expiry_date=expiry_date, document_id=document_id,
    )
    if alloc:
        from apps.stores.infrastructure.models import StockMovementBatch

        StockMovementBatch.objects.bulk_create([StockMovementBatch(movement=movement, batch=b, quantity=n) for b, n in alloc])
    movement.allocations = [{"batch": b.batch_number, "expiry": b.expiry_date.isoformat() if b.expiry_date else None, "qty": abs(n)} for b, n in alloc]
    if mirror and quantity:
        _mirror_inventory(store.id, variant_id, quantity, reference)
    return movement


def available_map(store: Store, variant_ids) -> dict:
    return {
        row["product_variant_id"]: max(0, row["quantity"] - row["reserved_quantity"])
        for row in StoreProductStock.objects.filter(store=store, product_variant_id__in=list(variant_ids)).values("product_variant_id", "quantity", "reserved_quantity")
    }


# ---------------------------------------------------------------- goods receipt (Kirim)


def _parse_date(value) -> date | None:
    if not value:
        return None
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError as exc:
        raise ValidationFailed(f"Invalid date: {value}") from exc


EXPIRY_PAST_DAYS = 365 * 2  # goods are not received with an expiry years ago -- that is a typo (e.g. 0002-05-01 for 2027-05-01)
EXPIRY_MAX_YEARS = 15


def expiry_on_receipt(value) -> date | None:
    """A goods-receipt expiry date, refusing typos that would poison FEFO and the expiry alerts (year 0002, 2207, …)."""
    d = _parse_date(value)
    if d is None:
        return None
    today = timezone.localdate()
    if d < today - timedelta(days=EXPIRY_PAST_DAYS) or d.year > today.year + EXPIRY_MAX_YEARS:
        raise ValidationFailed(f"Yaroqlilik sanasi noto'g'ri: {d.strftime('%d.%m.%Y')} — yilni tekshiring.")
    return d


def post_receipt(store: Store, lines: list[dict], actor: str, supplier: str = "", invoice_number: str = "", note: str = "") -> StockReceipt:
    """
    Goods arriving at a location (usually the main warehouse): lines = [{variant_id, quantity, unit_cost?, batch?, expiry?}].
    Same product + batch lines are merged. All or nothing.
    """
    from apps.products.infrastructure.models import ProductVariant

    merged: dict = {}
    for raw in lines:
        try:
            vid = uuid.UUID(str(raw.get("variant_id")))
            qty = int(raw.get("quantity"))
        except (TypeError, ValueError) as exc:
            raise ValidationFailed("Each line needs a product and a whole quantity.") from exc
        if qty <= 0 or qty > MAX_UNITS_PER_LINE:
            raise ValidationFailed(f"Quantity must be between 1 and {MAX_UNITS_PER_LINE}.")
        batch = str(raw.get("batch") or "").strip()[:48]
        key = (vid, batch)
        cost = raw.get("unit_cost")
        try:
            cost = Decimal(str(cost).replace(" ", "").replace(",", ".")) if cost not in (None, "") else None
        except ArithmeticError as exc:
            raise ValidationFailed("The unit cost must be a number.") from exc
        if cost is not None and (cost < 0 or cost >= Decimal("1e12")):
            raise ValidationFailed("The unit cost must be zero or positive.")
        line = merged.setdefault(key, {"variant_id": vid, "quantity": 0, "batch": batch, "expiry": expiry_on_receipt(raw.get("expiry")), "unit_cost": cost})
        if line["unit_cost"] is None:
            line["unit_cost"] = cost
        line["quantity"] += qty
    if not merged:
        raise conflict("EMPTY_RECEIPT", "Add at least one product.")
    found = set(ProductVariant.objects.filter(id__in={k[0] for k in merged}).values_list("id", flat=True))
    missing = [str(k[0]) for k in merged if k[0] not in found]
    if missing:
        raise ResourceNotFound(f"Unknown product(s): {', '.join(missing)}")

    from apps.stores.infrastructure import supplier_service

    with transaction.atomic():
        vendor = supplier_service.resolve(supplier)  # the typed name becomes / matches a Supplier record (accounts payable)
        receipt = StockReceipt.objects.create(
            number=next_number("KR", "stores_receipt_number_seq"), store=store, supplier=vendor.name if vendor else "", supplier_ref=vendor,
            due_date=supplier_service.due_date(vendor, timezone.localdate()), invoice_number=invoice_number[:60], note=note[:255], created_by=actor[:150],
        )
        units, value = 0, Decimal("0")
        for line in merged.values():
            StockReceiptItem.objects.create(receipt=receipt, product_variant_id=line["variant_id"], quantity=line["quantity"], unit_cost=line["unit_cost"],
                                            batch_number=line["batch"], expiry_date=line["expiry"])
            move(store, line["variant_id"], line["quantity"], StockMovementType.RECEIPT, reference=receipt.number, actor=actor,
                 batch_number=line["batch"], expiry_date=line["expiry"], document_id=receipt.id, note=supplier)
            units += line["quantity"]
            value += (line["unit_cost"] or 0) * line["quantity"]
        StockReceipt.objects.filter(pk=receipt.pk).update(total_units=units, total_value=value)
        receipt.refresh_from_db()
    return receipt


def parse_receipt_excel(file) -> tuple[list[dict], list[str]]:
    """
    Reads an .xlsx with columns Barcode | (Name) | Quantity [| Unit cost | Batch | Expiry] (header row, any order,
    matched by name). Returns (lines ready for post_receipt, human-readable problems per row).
    """
    from openpyxl import load_workbook

    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ["Fayl bo'sh."]
    header = [str(c or "").strip().lower() for c in rows[0]]

    import re

    def col(*names):
        # a name must start a word ("lot" matches "partiya (lot)" but not "mahsulot")
        for n in names:
            for i, h in enumerate(header):
                if re.search(r"(?<![\w'])" + re.escape(n), h):
                    return i
        return None

    c_code = col("barcode", "shtrix", "sku", "штрих")
    c_qty = col("soni", "quantity", "qty", "miqdor", "кол")
    c_cost = col("tannarx", "cost", "narx", "цена", "себест")
    c_batch = col("partiya", "batch", "lot", "парт")
    c_exp = col("muddat", "expiry", "срок")
    if c_code is None or c_qty is None:
        return [], ["Sarlavhada 'Barcode' va 'Soni' ustunlari bo'lishi kerak."]
    lines, problems = [], []
    for n, row in enumerate(rows[1:], start=2):
        if not row or all(v in (None, "") for v in row):
            continue
        code = str(row[c_code] or "").strip()
        if code.endswith(".0"):
            code = code[:-2]  # Excel turns long barcodes into floats
        try:
            qty = int(float(row[c_qty] or 0))
        except (TypeError, ValueError):
            problems.append(f"{n}-qator: soni noto'g'ri ({row[c_qty]!r}).")
            continue
        if qty <= 0:
            continue
        variant = find_variant(code)
        if variant is None:
            problems.append(f"{n}-qator: '{code}' shtrix-kodli mahsulot topilmadi.")
            continue
        expiry = row[c_exp] if c_exp is not None else None
        lines.append({
            "variant_id": str(variant.id), "quantity": qty,
            "unit_cost": row[c_cost] if c_cost is not None and row[c_cost] not in (None, "") else None,
            "batch": str(row[c_batch] or "").strip() if c_batch is not None else "",
            "expiry": expiry.isoformat()[:10] if hasattr(expiry, "isoformat") else (str(expiry)[:10] if expiry else None),
        })
    return lines, problems
