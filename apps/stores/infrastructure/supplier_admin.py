"""
Dashboard → Ombor → Yetkazib beruvchilar: suppliers, what we bought from them and what we owe (see supplier_service.py).

Pages: the list (debt / overdue KPIs and filters), the supplier card (statement with running balance,
receipts with payment status, products and cost trend, payments), Excel statement (akt sverka).
The add / edit form is the standard admin form; saving returns to the card.
Permissions: Supplier / SupplierPayment model permissions.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal

from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.db.models import Sum
from django.http import HttpResponseRedirect
from django.shortcuts import get_object_or_404, render
from django.urls import path, reverse
from django.utils import timezone
from django.utils.text import slugify
from django.utils.translation import gettext as _t
from django.utils.translation import gettext_lazy as _

from apps.platform.infrastructure import excel_export as xl
from apps.platform.infrastructure.excel_export import Col, ExcelExportMixin
from apps.stores.infrastructure import stock_service, supplier_service
from apps.stores.infrastructure.models import PaymentMethod, StockReceipt, Supplier, SupplierPayment
from apps.stores.infrastructure.warehouse_admin import _actor, _local, _slug, _som, excel_response
from common.exceptions import DomainException

RECEIPT_STATE = {
    "paid": ("#10b981", _("To'langan")), "partial": ("#0ea5e9", _("Qisman to'langan")), "open": ("#64748b", _("To'lanmagan")),
    "overdue": ("#ef4444", _("Muddati o'tgan")), "nocost": ("#94a3b8", _("Tannarxsiz")),
}


class DebtFilter(admin.SimpleListFilter):
    title = _("Debt")
    parameter_name = "debt"

    def lookups(self, request, model_admin):
        return (("yes", _("With debt")), ("overdue", _("Overdue")))

    def queryset(self, request, queryset):
        if self.value() not in ("yes", "overdue"):
            return queryset
        bal = supplier_service.balances()
        ids = [sid for sid, b in bal.items() if (b.overdue > 0 if self.value() == "overdue" else b.debt > 0)]
        return queryset.filter(id__in=ids)


@admin.register(Supplier)
class SupplierAdmin(ExcelExportMixin, admin.ModelAdmin):
    excel_title = "Yetkazib beruvchilar"
    excel_columns = [
        Col("Nomi", "name", 26), Col("Yuridik nomi", "legal_name", 26), Col("STIR", "tin", 12), Col("Mas'ul shaxs", "contact_person", 20),
        Col("Telefon", "phone", 15), Col("Email", "email", 20), Col("Manzil", "address", 30), Col("To'lov muddati (kun)", "payment_terms_days", 12, number=True),
        Col("Faol", "is_active", 7), Col("Izoh", "note", 26),
    ]

    list_display = ("name", "contact_person", "phone", "is_active")
    list_filter = (DebtFilter, "is_active")
    search_fields = ("name", "legal_name", "tin", "contact_person", "phone")
    list_per_page = 40
    change_list_template = "admin/wms/supplier_list.html"
    fieldsets = (
        (None, {"fields": ("name", "is_active")}),
        (_("Requisites"), {"fields": ("legal_name", "tin", "bank_details", "address")}),
        (_("Contact"), {"fields": ("contact_person", "phone", "email")}),
        (_("Terms"), {"fields": ("payment_terms_days", "note")}),
    )

    def has_delete_permission(self, request, obj=None):
        return False  # receipts point at it; deactivate instead

    def get_changelist_instance(self, request):
        cl = super().get_changelist_instance(request)
        bal = supplier_service.balances([s.pk for s in cl.result_list])
        today = timezone.localdate()
        for s in cl.result_list:
            b = bal.get(s.pk) or supplier_service.Balance()
            s.nz = {"purchased": _som(b.purchased), "paid": _som(b.paid), "debt": _som(b.debt), "has_debt": b.debt > 0, "advance": b.debt < 0,
                    "advance_amount": _som(-b.debt), "overdue": _som(b.overdue) if b.overdue else "", "receipts": b.receipts,
                    "last": _local(b.last_receipt, "%d.%m.%Y"), "late_days": (today - b.oldest_overdue).days if b.oldest_overdue else 0}
        return cl

    def changelist_view(self, request, extra_context=None):
        bal = supplier_service.balances()
        month = timezone.localtime(timezone.now()).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        bought = StockReceipt.objects.filter(supplier_ref__isnull=False, created_at__gte=month).aggregate(v=Sum("total_value"))["v"]
        paid = SupplierPayment.objects.filter(paid_on__gte=month.date()).aggregate(v=Sum("amount"))["v"]
        current = request.GET.get("debt", "")
        active = request.GET.get("is_active__exact", "")
        tabs = [
            {"url": "?", "label": _t("All"), "n": Supplier.objects.count(), "color": "#64748b", "icon": "fa-layer-group", "on": not current and not active},
            {"url": "?debt=yes", "label": _t("With debt"), "n": sum(1 for b in bal.values() if b.debt > 0), "color": "#f59e0b", "icon": "fa-scale-unbalanced", "on": current == "yes"},
            {"url": "?debt=overdue", "label": _t("Overdue"), "n": sum(1 for b in bal.values() if b.overdue > 0), "color": "#ef4444", "icon": "fa-clock", "on": current == "overdue"},
            {"url": "?is_active__exact=0", "label": _t("Inactive"), "n": Supplier.objects.filter(is_active=False).count(), "color": "#94a3b8", "icon": "fa-box-archive", "on": active == "0"},
        ]
        extra = {
            "title": _t("Yetkazib beruvchilar"), "nz_tabs": tabs,
            "nz_kpis": {"debt": _som(sum((b.debt for b in bal.values() if b.debt > 0), Decimal(0))), "overdue": _som(sum((b.overdue for b in bal.values()), Decimal(0))),
                        "overdue_n": sum(1 for b in bal.values() if b.overdue > 0), "bought": _som(bought), "paid": _som(paid)},
            "unlinked": StockReceipt.objects.filter(supplier_ref__isnull=True).count(),
            "can_add": request.user.has_perm("stores.add_supplier"),
        }
        return super().changelist_view(request, {**extra, **(extra_context or {})})

    def get_urls(self):
        v = self.admin_site.admin_view
        return [
            path("<uuid:pk>/card/", v(self.card_view), name="stores_supplier_card"),
            path("<uuid:pk>/pay/", v(self.pay_view), name="stores_supplier_pay"),
            path("<uuid:pk>/pay/<uuid:pid>/delete/", v(self.delete_payment_view), name="stores_supplier_pay_delete"),
            path("<uuid:pk>/statement.xlsx", v(self.statement_xlsx), name="stores_supplier_statement"),
            *super().get_urls(),
        ]

    def response_add(self, request, obj, post_url_continue=None):
        if "_addanother" in request.POST or "_continue" in request.POST:
            return super().response_add(request, obj, post_url_continue)
        return HttpResponseRedirect(reverse("admin:stores_supplier_card", args=[obj.pk]))

    def response_change(self, request, obj):
        if "_addanother" in request.POST or "_continue" in request.POST:
            return super().response_change(request, obj)
        return HttpResponseRedirect(reverse("admin:stores_supplier_card", args=[obj.pk]))

    # ------------------------------------------------------------------ card
    @staticmethod
    def _period(request) -> tuple[date | None, date | None]:
        out = []
        for key in ("from", "to"):
            try:
                out.append(date.fromisoformat(request.GET.get(key, "")[:10]))
            except ValueError:
                out.append(None)
        return out[0], out[1]

    def card_view(self, request, pk):
        if not self.has_view_permission(request):
            raise PermissionDenied
        s = get_object_or_404(Supplier, pk=pk)
        b = supplier_service.balance(s)
        start, end = self._period(request)
        st = supplier_service.statement(s, start, end)
        for row in st["rows"]:
            if row["kind"] == "receipt":
                info = b.receipt_status.get(row["obj"].pk) or {}
                color, label = RECEIPT_STATE.get(info.get("status"), ("#64748b", ""))
                row.update({"state_color": color, "state_label": label, "open": _som(info.get("open")) if info.get("open") else "",
                            "due": row["obj"].due_date, "url": reverse("admin:stores_stockreceipt_change", args=[row["obj"].pk])})
            for k in ("debit", "credit", "balance"):
                row[k + "_s"] = _som(row[k]) if row[k] else ""
        top = supplier_service.top_products(s)
        infos = stock_service.variant_infos([t["product_variant_id"] for t in top])
        for t in top:
            t["info"] = infos.get(t["product_variant_id"])
            t["trend"] = None if t["last_cost"] is None or t["prev_cost"] is None or t["last_cost"] == t["prev_cost"] else ("up" if t["last_cost"] > t["prev_cost"] else "down")
            t["last_cost_s"] = _som(t["last_cost"]) if t["last_cost"] is not None else ""
            t["change"] = round((t["last_cost"] - t["prev_cost"]) * 100 / t["prev_cost"]) if t["trend"] and t["prev_cost"] else None
        today = timezone.localdate()
        return render(request, "admin/wms/supplier_card.html", {
            **self.admin_site.each_context(request), "opts": self.model._meta, "title": s.name, "s": s, "b": b,
            "kpi": {"purchased": _som(b.purchased), "paid": _som(b.paid), "debt": _som(abs(b.debt)), "advance": b.debt < 0, "overdue": _som(b.overdue) if b.overdue else "",
                    "late_days": (today - b.oldest_overdue).days if b.oldest_overdue else 0},
            "st": {**st, "opening_s": _som(st["opening"]), "debit_s": _som(st["debit"]), "credit_s": _som(st["credit"]), "closing_s": _som(st["closing"])},
            "start": start.isoformat() if start else "", "end": end.isoformat() if end else "", "top": top,
            "methods": PaymentMethod.choices, "today": today.isoformat(), "suggest": int(b.overdue or max(b.debt, 0)),
            "can_pay": request.user.has_perm("stores.add_supplierpayment"), "can_edit": request.user.has_perm("stores.change_supplier"),
            "can_delete_pay": request.user.has_perm("stores.delete_supplierpayment"),
        })

    def pay_view(self, request, pk):
        if request.method != "POST" or not request.user.has_perm("stores.add_supplierpayment"):
            raise PermissionDenied
        s = get_object_or_404(Supplier, pk=pk)
        p = request.POST
        try:
            pay = supplier_service.add_payment(s, _actor(request), p.get("amount"), p.get("paid_on"), p.get("method") or PaymentMethod.BANK, p.get("reference", ""), p.get("note", ""))
        except DomainException as exc:
            messages.error(request, str(exc))
        else:
            messages.success(request, _t("Payment of %s so'm recorded.") % _som(pay.amount))
        return HttpResponseRedirect(reverse("admin:stores_supplier_card", args=[s.pk]))

    def delete_payment_view(self, request, pk, pid):
        if request.method != "POST" or not request.user.has_perm("stores.delete_supplierpayment"):
            raise PermissionDenied
        pay = get_object_or_404(SupplierPayment, pk=pid, supplier_id=pk)
        pay.delete()
        messages.success(request, _t("Payment deleted."))
        return HttpResponseRedirect(reverse("admin:stores_supplier_card", args=[pk]))

    def statement_xlsx(self, request, pk):
        if not self.has_view_permission(request):
            raise PermissionDenied
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        s = get_object_or_404(Supplier, pk=pk)
        start, end = self._period(request)
        st = supplier_service.statement(s, start, end)
        wb = Workbook()
        ws = wb.active
        ws.title = "Akt sverka"
        ws.append([f"Solishtirma dalolatnoma: {s.name}" + (f" (STIR {s.tin})" if s.tin else "")])
        ws.append([f"Davr: {start or '…'} — {end or timezone.localdate()}"])
        ws.append([])
        ws.append(["Sana", "Hujjat", "Mazmun", "Kirim (qarz)", "To'lov", "Qoldiq"])
        for c in ws[4]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A8A")
        ws.append(["", "", "Boshlang'ich qoldiq", None, None, float(st["opening"])])
        for r in st["rows"]:
            ws.append([r["date"], r["doc"], ("Kirim · " if r["kind"] == "receipt" else "To'lov · ") + r["text"], float(r["debit"]) or None, float(r["credit"]) or None, float(r["balance"])])
        ws.append(["", "", "Jami / yakuniy qoldiq", float(st["debit"]), float(st["credit"]), float(st["closing"])])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        ws["A1"].font = Font(bold=True, size=13)
        for col, width in zip("ABCDEF", (12, 14, 48, 16, 16, 16)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=5, min_col=4, max_col=6):
            for c in row:
                c.number_format = "#,##0"
        return excel_response(wb, f"{_slug()}_akt_sverka_{slugify(s.name)[:30] or 'supplier'}.xlsx")
