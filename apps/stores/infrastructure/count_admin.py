"""
Dashboard → Ombor → Inventarizatsiya: stock counts (see count_service.py for the workflow and stock rules).

Pages: the list, "new count", the counting workspace (a JSON app: scan / type counts, review differences,
post), the printable count act (dalolatnoma) and an Excel report. Permissions: the StockCount model
permissions plus `post_stockcount` (only managers change stock).
"""
from __future__ import annotations

import uuid
from decimal import Decimal

from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.db.models import Count, Max, Q, Sum
from django.http import HttpResponseRedirect, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.urls import path, reverse
from django.utils import timezone
from django.utils.translation import gettext as _t
from django.utils.translation import gettext_lazy as _

from apps.stores.infrastructure import count_service, stock_service, transfer_service
from apps.stores.infrastructure.models import CountReason, StockCount, StockCountScope, StockCountStatus, Store, StoreKind
from apps.stores.infrastructure.warehouse_admin import _actor, _body, _error, _local, _location_payload, _som, excel_response
from common.exceptions import DomainException

STATUS_STYLE = {
    "counting": ("#0ea5e9", "fa-barcode"), "review": ("#f59e0b", "fa-magnifying-glass-chart"),
    "posted": ("#10b981", "fa-circle-check"), "cancelled": ("#ef4444", "fa-ban"),
}


@admin.register(StockCount)
class StockCountAdmin(admin.ModelAdmin):
    list_display = ("number", "store", "status", "created_at")
    list_filter = ("status", "store")
    search_fields = ("number", "note", "created_by")
    list_per_page = 30
    ordering = ("-created_at",)
    change_list_template = "admin/wms/count_list.html"

    def has_delete_permission(self, request, obj=None):
        return False

    def get_queryset(self, request):
        return super().get_queryset(request).select_related("store").annotate(
            _lines=Count("items"), _counted=Count("items", filter=Q(items__counted_quantity__isnull=False)),
        )

    def get_changelist_instance(self, request):
        cl = super().get_changelist_instance(request)
        for c in cl.result_list:
            color, icon = STATUS_STYLE.get(c.status, ("#64748b", "fa-circle"))
            c.nz = {"color": color, "icon": icon, "progress": round(c._counted * 100 / c._lines) if c._lines else 0, "counted": c._counted, "lines": c._lines,
                    "created": _local(c.created_at), "posted": _local(c.posted_at)}
        return cl

    def changelist_view(self, request, extra_context=None):
        month = timezone.localtime(timezone.now()).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        posted = StockCount.objects.filter(status=StockCountStatus.POSTED, posted_at__gte=month)
        agg = posted.aggregate(n=Count("id"), short=Sum("shortage_value"), excess=Sum("excess_value"), lines=Sum("lines_total"), matched=Sum("lines_matched"))
        current = request.GET.get("status__exact", "")
        counts = dict(StockCount.objects.values_list("status").annotate(n=Count("id")))
        tabs = [{"key": "", "label": _t("All"), "n": sum(counts.values()), "color": "#64748b", "icon": "fa-layer-group", "on": not current}]
        for key, label in StockCountStatus.choices:
            color, icon = STATUS_STYLE[key]
            tabs.append({"key": key, "label": label, "n": counts.get(key, 0), "color": color, "icon": icon, "on": current == key})
        open_counts = StockCount.objects.filter(status__in=(StockCountStatus.COUNTING, StockCountStatus.REVIEW)).select_related("store")
        extra = {
            "title": _t("Inventarizatsiya"), "nz_tabs": tabs, "nz_open": open_counts,
            "nz_kpis": {"month": agg["n"] or 0, "short": _som(agg["short"]), "excess": _som(agg["excess"]),
                        "accuracy": round((agg["matched"] or 0) * 100 / agg["lines"]) if agg["lines"] else None},
            "can_create": request.user.has_perm("stores.add_stockcount"),
        }
        return super().changelist_view(request, {**extra, **(extra_context or {})})

    def get_urls(self):
        v = self.admin_site.admin_view
        return [
            path("new/", v(self.new_view), name="stores_stockcount_new"),
            path("<uuid:pk>/work/", v(self.workspace_view), name="stores_stockcount_workspace"),
            path("<uuid:pk>/data/", v(self.api_data), name="stores_stockcount_data"),
            path("<uuid:pk>/do/<str:action>/", v(self.api_action), name="stores_stockcount_action"),
            path("<uuid:pk>/act/", v(self.act_view), name="stores_stockcount_act"),
            path("<uuid:pk>/report.xlsx", v(self.export_view), name="stores_stockcount_export"),
            *super().get_urls(),
        ]

    def change_view(self, request, object_id, form_url="", extra_context=None):
        return HttpResponseRedirect(reverse("admin:stores_stockcount_workspace", args=[object_id]))

    def add_view(self, request, form_url="", extra_context=None):
        return HttpResponseRedirect(reverse("admin:stores_stockcount_new"))

    def _ctx(self, request, title):
        return {**self.admin_site.each_context(request), "opts": self.model._meta, "title": title}

    # ------------------------------------------------------------------ new
    def new_view(self, request):
        if not request.user.has_perm("stores.add_stockcount"):
            raise PermissionDenied
        locations = list(transfer_service.locations())
        if request.method == "POST":
            ids = [v for v in request.POST.getlist("variant") if v]
            try:
                c = count_service.start(request.POST.get("store"), _actor(request), scope=request.POST.get("scope") or StockCountScope.FULL,
                                        blind=bool(request.POST.get("blind")), note=request.POST.get("note", ""), variant_ids=ids)
            except DomainException as exc:
                messages.error(request, str(exc))
            else:
                return HttpResponseRedirect(reverse("admin:stores_stockcount_workspace", args=[c.pk]))
        from apps.stores.infrastructure.models import StoreProductStock

        per_store = {str(r["store"]): r for r in StoreProductStock.objects.filter(quantity__gt=0).values("store").annotate(lines=Count("id"), units=Sum("quantity"))}
        last = {str(r["store"]): r["last"] for r in StockCount.objects.filter(status=StockCountStatus.POSTED).values("store").annotate(last=Max("posted_at"))}
        open_ = {str(c.store_id): c for c in StockCount.objects.filter(status__in=(StockCountStatus.COUNTING, StockCountStatus.REVIEW))}
        rows = [{"s": s, "stats": per_store.get(str(s.id)), "open": open_.get(str(s.id)), "last": _local(last.get(str(s.id)), "%d.%m.%Y")} for s in locations]
        first_free = next((r["s"].id for r in rows if not r["open"]), None)
        return render(request, "admin/wms/count_new.html", {
            **self._ctx(request, _t("Yangi inventarizatsiya")), "rows": rows, "first_free": first_free,
            "search_url": reverse("admin:stores_stocktransfer_search"),
        })

    # ------------------------------------------------------------------ workspace + API
    def workspace_view(self, request, pk):
        c = get_object_or_404(StockCount.objects.select_related("store"), pk=pk)
        return render(request, "admin/wms/count_workspace.html", {
            **self._ctx(request, f"{_t('Inventarizatsiya')} {c.number}"), "count": c,
            "config": {
                "id": str(c.pk), "number": c.number,
                "urls": {
                    "data": reverse("admin:stores_stockcount_data", args=[c.pk]),
                    "action": reverse("admin:stores_stockcount_action", args=[c.pk, "__"]),
                    "lookup": reverse("admin:stores_stocktransfer_lookup"),
                    "search": reverse("admin:stores_stocktransfer_search"),
                    "act": reverse("admin:stores_stockcount_act", args=[c.pk]),
                    "export": reverse("admin:stores_stockcount_export", args=[c.pk]),
                    "list": reverse("admin:stores_stockcount_changelist"),
                },
                "perms": {"count": request.user.has_perm("stores.change_stockcount"), "post": request.user.has_perm("stores.post_stockcount")},
                "reasons": [(k, str(v)) for k, v in CountReason.choices],
            },
            "i18n": self._i18n(),
        })

    @staticmethod
    def _i18n() -> dict:
        t = _t
        return {
            "error": t("Could not save. Please try again."), "saving": t("saving…"), "saved": t("✓ saved"), "cancel": t("Cancel count"),
            "submit": t("Finish counting"), "recount": t("Back to counting"), "post": t("Post the count"), "act": t("Count act"),
            "postedAt": t("Posted"), "back": t("Back"), "blind": t("Blind count"), "cancelledAt": t("Cancelled"),
            "blindNote": t("Blind count: count what is on the shelf -- the system quantity is shown to the manager at review."),
            "counted": t("Counted"), "accuracy": t("Accuracy"), "matched": t("matched"), "units": t("pcs"), "shortage": t("Shortage"),
            "excess": t("Excess"), "som": t("so'm"), "unitsCounted": t("Units counted"), "done": t("counted"),
            "tabAll": t("All"), "tabUncounted": t("Not counted"), "tabDiff": t("Differences"), "tabMatch": t("Matched"),
            "scanPh": t("Scan a barcode or type SKU and press Enter…"), "scanHint": t("1 scan = 1 unit"), "product": t("Product"),
            "system": t("System"), "diff": t("Difference"), "reason": t("Reason"), "searchPh": t("Find a product…"), "nothing": t("Nothing here"),
            "chooseReason": t("Choose a reason…"), "notePh": t("Comment"), "clamped": t("posted as"), "added": t("found during the count"),
            "lastEmpty": t("Scan the first product -- it shows up here big enough to see from a distance."), "lastScanned": t("Last scanned"),
            "addedToast": t("Not on the list -- added"), "submitTitle": t("Finish counting?"),
            "submitText": t("The count goes to review: differences are shown and need a reason before posting."),
            "zeroTitle": t("Set the %s uncounted product(s) to 0"), "zeroText": t("Products that weren't found on the shelf are counted as missing."),
            "submittedToast": t("Counting finished -- review the differences"), "recountToast": t("Back to counting"),
            "postTitle": t("Post the count?"), "postText": t("Every difference is written to the stock ledger as a correction. This can't be undone."),
            "missingReasons": t("%s difference(s) still need a reason."), "postedToast": t("Count posted -- stock corrected"),
            "cancelTitle": t("Cancel this count?"), "cancelText": t("Nothing changes in stock. The counted numbers are kept for reference."),
            "cancelledToast": t("Count cancelled"),
        }

    def _payload(self, request, c: StockCount) -> dict:
        items = list(c.items.all())
        infos = stock_service.variant_infos([i.product_variant_id for i in items])
        hide = c.blind and c.status == StockCountStatus.COUNTING and not request.user.has_perm("stores.post_stockcount")
        color, icon = STATUS_STYLE.get(c.status, ("#64748b", "fa-circle"))
        summary = count_service.summary(c)
        return {
            "id": str(c.pk), "number": c.number, "status": c.status, "status_label": c.get_status_display(), "color": color, "icon": icon,
            "store": _location_payload(c.store), "scope": c.scope, "scope_label": c.get_scope_display(), "blind": c.blind, "hidden": hide,
            "note": c.note, "created_by": c.created_by, "created_at": _local(c.created_at), "submitted_by": c.submitted_by, "submitted_at": _local(c.submitted_at),
            "posted_by": c.posted_by, "posted_at": _local(c.posted_at), "cancelled_at": _local(c.cancelled_at), "cancel_reason": c.cancel_reason,
            "summary": {k: (_som(v) if isinstance(v, Decimal) else v) for k, v in summary.items()},
            "items": [{
                "id": str(i.id), "variant_id": str(i.product_variant_id), "name": i.name, "sku": i.sku, "barcode": i.barcode,
                "size": infos[i.product_variant_id].size if i.product_variant_id in infos else None,
                "image_url": infos[i.product_variant_id].image_url if i.product_variant_id in infos else None,
                "expected": None if hide else i.expected_quantity, "counted": i.counted_quantity,
                "difference": None if hide else (i.difference if i.counted_quantity is not None else None),
                "value": None if hide else _som(abs(i.difference) * i.unit_cost) if i.counted_quantity is not None and i.difference else "",
                "reason": i.reason, "note": i.note, "counted_by": i.counted_by, "added": i.added_during_count, "adjusted": i.adjusted_quantity,
            } for i in items],
        }

    def api_data(self, request, pk):
        return JsonResponse({"ok": True, "data": self._payload(request, get_object_or_404(StockCount.objects.select_related("store"), pk=pk))})

    ACTION_PERMS = {
        "save_counts": "change_stockcount", "add": "change_stockcount", "scan": "change_stockcount", "submit": "change_stockcount",
        "recount": "change_stockcount", "save_reasons": "change_stockcount", "post": "post_stockcount", "cancel": "change_stockcount",
    }

    def api_action(self, request, pk, action):
        if request.method != "POST":
            return JsonResponse({"ok": False}, status=405)
        perm = self.ACTION_PERMS.get(action)
        if perm is None:
            return JsonResponse({"ok": False, "error": {"code": "UNKNOWN_ACTION", "message": action}}, status=404)
        if not request.user.has_perm(f"stores.{perm}"):
            return JsonResponse({"ok": False, "error": {"code": "FORBIDDEN", "message": _t("Bu amal uchun sizda ruxsat yo'q.")}}, status=403)
        b, a, s = _body(request), _actor(request), count_service
        extra = {}
        try:
            if action == "save_counts":
                s.save_counts(pk, a, b.get("counts") or {})
            elif action == "add":
                s.add_products(pk, a, b.get("variant_ids") or [])
            elif action == "scan":
                # a code that isn't on the list yet: the product is added (found on the shelf), the client then counts it
                match = stock_service.find_match(str(b.get("code") or ""))
                if match is None:
                    return JsonResponse({"ok": False, "error": {"code": "NOT_FOUND", "message": _t("Bu shtrix-kod bo'yicha mahsulot topilmadi.")}}, status=404)
                made = s.add_products(pk, a, [match.variant.id])
                extra["variant_id"] = str(match.variant.id)
                extra["added"] = bool(made)
            elif action == "submit":
                s.submit(pk, a, uncounted_as_zero=bool(b.get("uncounted_as_zero")))
            elif action == "recount":
                s.recount(pk, a)
            elif action == "save_reasons":
                s.save_reasons(pk, a, b.get("reasons") or {})
            elif action == "post":
                if b.get("reasons"):
                    s.save_reasons(pk, a, b.get("reasons"))
                s.post(pk, a)
            elif action == "cancel":
                s.cancel(pk, a, str(b.get("reason") or ""))
        except DomainException as exc:
            return _error(exc)
        c = StockCount.objects.select_related("store").get(pk=pk)
        return JsonResponse({"ok": True, "data": self._payload(request, c), **extra})

    # ------------------------------------------------------------------ act + Excel
    def _rows(self, c):
        rows = []
        for n, i in enumerate(sorted(c.items.all(), key=lambda x: (x.difference == 0, x.name)), 1):
            rows.append({"n": n, "i": i, "diff": i.difference, "value": abs(i.difference) * i.unit_cost})
        return rows

    def act_view(self, request, pk):
        c = get_object_or_404(StockCount.objects.select_related("store"), pk=pk)
        return render(request, "admin/wms/count_act.html", {
            **self._ctx(request, c.number), "c": c, "rows": self._rows(c), "summary": count_service.summary(c),
            "created": _local(c.created_at), "posted": _local(c.posted_at),
        })

    def export_view(self, request, pk):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        c = get_object_or_404(StockCount.objects.select_related("store"), pk=pk)
        wb = Workbook()
        ws = wb.active
        ws.title = c.number
        ws.append(["#", "Mahsulot", "SKU", "Shtrix-kod", "Tizimda", "Sanaldi", "Farq", "Tannarx", "Farq qiymati", "Sabab", "Izoh", "Sanagan"])
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A8A")
        for r in self._rows(c):
            i = r["i"]
            ws.append([r["n"], i.name, i.sku, i.barcode, i.expected_quantity, i.counted_quantity, r["diff"] if i.counted_quantity is not None else None,
                       float(i.unit_cost), float(r["value"]) * (-1 if r["diff"] < 0 else 1), i.get_reason_display() if i.reason else "", i.note, i.counted_by])
        for col, width in zip("ABCDEFGHIJKL", (5, 46, 16, 16, 10, 10, 9, 12, 14, 24, 30, 18)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        from apps.stores.infrastructure.warehouse_admin import _slug

        return excel_response(wb, f"{_slug()}_{c.number}.xlsx")
