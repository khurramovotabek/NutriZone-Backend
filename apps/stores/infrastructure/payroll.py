"""
Cashier payroll: a fixed monthly base salary + KPI.

    KPI   = kpi_percent × the cashier's own NET sales that month (what customers paid for the goods, coins
            included, minus everything returned) at their kassa -- the same revenue figure every POS report uses.
    Total = base salary + KPI + adjustment (a bonus or a deduction, with its reason).

Months follow Tashkent wall-clock time. An unpaid month is always computed live from the sales; paying it freezes
the figures into a CashierPayout (later returns or a changed rate never rewrite a paid month) and tells the cashier.
"""
from __future__ import annotations

from datetime import date, datetime, time
from decimal import ROUND_HALF_UP, Decimal
from zoneinfo import ZoneInfo

from django.db import IntegrityError, transaction

from common.exceptions import ConflictError, ValidationFailed

TZ = ZoneInfo("Asia/Tashkent")
MAX_PERCENT = Decimal("50")
MAX_SALARY = Decimal("1000000000")


def _err(code: str, message: str) -> ConflictError:
    e = ConflictError(message)
    e.code = code
    return e


def month_start(value: str | date | None = None) -> date:
    """The 1st of a month from "2026-10", a date, or None (= this month in Tashkent)."""
    if isinstance(value, date):
        return value.replace(day=1)
    if value:
        try:
            y, m = (int(x) for x in str(value)[:7].split("-"))
            return date(y, m, 1)
        except (ValueError, TypeError):
            pass
    return datetime.now(TZ).date().replace(day=1)


def next_month(m: date) -> date:
    return date(m.year + (m.month == 12), m.month % 12 + 1, 1)


def prev_month(m: date) -> date:
    return date(m.year - (m.month == 1), (m.month - 2) % 12 + 1, 1)


def month_bounds(m: date) -> tuple[datetime, datetime]:
    return datetime.combine(m, time.min, TZ), datetime.combine(next_month(m), time.min, TZ)


def is_current(m: date) -> bool:
    return m == month_start()


def kpi(revenue: Decimal, percent: Decimal) -> Decimal:
    return (Decimal(revenue) * Decimal(percent) / 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP)


def sales_of(month: date, user_ids=None, store_ids=None) -> dict:
    """{(user_id, store_id): {"sales": n, "revenue": Decimal}} -- completed till sales of the month, net of returns."""
    from apps.pos.infrastructure import sales_reports as rep
    from apps.pos.infrastructure.models import SaleModel

    start, end = month_bounds(month)
    qs = SaleModel.objects.filter(status__in=("completed", "refunded"), completed_at__gte=start, completed_at__lt=end)
    if user_ids is not None:
        qs = qs.filter(cashier_id__in=list(user_ids))
    if store_ids is not None:
        qs = qs.filter(branch_id__in=list(store_ids))
    out: dict = {}
    for r in rep.rows(qs):
        a = out.setdefault((r["cashier_id"], r["branch_id"]), {"sales": 0, "revenue": Decimal(0)})
        a["sales"] += r["status"] == "completed"
        a["revenue"] += r["revenue"]
    return out


def month_report(month: date) -> dict:
    """Every cashier's pay for the month: paid ones frozen, the rest live from their sales and current rates."""
    from apps.stores.infrastructure.models import CashierPayout, StoreCashier

    links = list(StoreCashier.objects.filter(is_active=True).select_related("user", "store").order_by("store__name", "user__full_name"))
    paid = {p.user_id: p for p in CashierPayout.objects.filter(month=month).select_related("user", "store")}
    sales = sales_of(month, user_ids={lnk.user_id for lnk in links} | set(paid))
    rows = []
    for lnk in links:
        if lnk.user_id in paid:
            continue
        s = sales.get((lnk.user_id, lnk.store_id), {"sales": 0, "revenue": Decimal(0)})
        k = kpi(s["revenue"], lnk.kpi_percent)
        rows.append({"user": lnk.user, "store": lnk.store, "link": lnk, "payout": None, "sales": s["sales"], "revenue": s["revenue"],
                     "base": lnk.base_salary, "percent": lnk.kpi_percent, "kpi": k, "adjustment": Decimal(0), "note": "", "total": lnk.base_salary + k})
    for p in paid.values():
        rows.append({"user": p.user, "store": p.store, "link": next((x for x in links if x.user_id == p.user_id), None), "payout": p, "sales": p.sales_count,
                     "revenue": p.revenue, "base": p.base_salary, "percent": p.kpi_percent, "kpi": p.kpi_amount, "adjustment": p.adjustment, "note": p.adjustment_note, "total": p.total})
    rows.sort(key=lambda r: (r["store"].name, r["user"].full_name or ""))
    totals = {k: sum((r[k] for r in rows), Decimal(0)) for k in ("revenue", "base", "kpi", "adjustment", "total")}
    totals["sales"] = sum(r["sales"] for r in rows)
    totals["paid"] = sum((r["total"] for r in rows if r["payout"]), Decimal(0))
    totals["unpaid"] = totals["total"] - totals["paid"]
    totals["paid_n"] = sum(1 for r in rows if r["payout"])
    return {"month": month, "rows": rows, "totals": totals, "current": is_current(month)}


def set_terms(link, base_salary, kpi_percent) -> None:
    try:
        base, pct = Decimal(str(base_salary).replace(" ", "")), Decimal(str(kpi_percent).replace(",", "."))
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailed("Enter the salary and the KPI percent as numbers.") from exc
    if not (0 <= base <= MAX_SALARY) or not (0 <= pct <= MAX_PERCENT):
        raise ValidationFailed("The salary must be 0 or more and the KPI between 0 and 50%.")
    link.base_salary, link.kpi_percent = base.quantize(Decimal("1")), pct.quantize(Decimal("0.01"))
    link.save(update_fields=["base_salary", "kpi_percent"])


def pay(link, month: date, adjustment=0, note: str = "", paid_by: str = ""):
    """Freezes the month's pay for this cashier and notifies them. A month can be paid once (undo first to redo)."""
    from apps.stores.infrastructure.models import CashierPayout

    if month > month_start():
        raise ValidationFailed("A future month can't be paid.")
    try:
        adj = Decimal(str(adjustment or 0).replace(" ", ""))
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailed("The adjustment must be a number.") from exc
    if adj and not note.strip():
        raise ValidationFailed("Write the reason for the bonus or deduction.")
    s = sales_of(month, user_ids=[link.user_id], store_ids=[link.store_id]).get((link.user_id, link.store_id), {"sales": 0, "revenue": Decimal(0)})
    k = kpi(s["revenue"], link.kpi_percent)
    total = link.base_salary + k + adj
    if total < 0:
        raise ValidationFailed("The deduction is larger than the pay.")
    try:
        with transaction.atomic():
            payout = CashierPayout.objects.create(
                user_id=link.user_id, store_id=link.store_id, month=month, sales_count=s["sales"], revenue=s["revenue"], base_salary=link.base_salary,
                kpi_percent=link.kpi_percent, kpi_amount=k, adjustment=adj, adjustment_note=note.strip()[:255], total=total, paid_by=paid_by[:150],
            )
    except IntegrityError as exc:
        raise _err("ALREADY_PAID", "This month is already paid for this cashier.") from exc
    from apps.notifications.infrastructure.notify import notify

    notify(link.user_id, "salary_paid", {"month": f"{month:%m.%Y}", "total": _som(total), "kpi": _som(k), "revenue": _som(s["revenue"])},
           link="/profile/store/pos", dedupe_key=f"salary_paid:{payout.id}")
    return payout


def undo(payout) -> None:
    payout.delete()


def earnings(user_id, store, month: date | None = None) -> dict | None:
    """What a cashier has earned so far this month (shown to them on the POS page)."""
    from apps.stores.infrastructure.models import CashierPayout, StoreCashier

    link = StoreCashier.objects.filter(user_id=user_id, store=store, is_active=True).first()
    if link is None:
        return None
    month = month or month_start()
    paid = CashierPayout.objects.filter(user_id=user_id, month=month).first()
    s = sales_of(month, user_ids=[user_id], store_ids=[store.id]).get((user_id, store.id), {"sales": 0, "revenue": Decimal(0)})
    k = kpi(s["revenue"], link.kpi_percent)
    return {"month": f"{month:%Y-%m}", "sales": s["sales"], "revenue": str(s["revenue"]), "base_salary": str(link.base_salary), "kpi_percent": str(link.kpi_percent.normalize()),
            "kpi": str(k), "total": str(link.base_salary + k), "paid": str(paid.total) if paid else None}


def _som(v) -> str:
    return f"{int(Decimal(v or 0)):,}".replace(",", " ")


def excel(report: dict) -> bytes:
    """The month's payroll sheet (vedomost) for the accountant."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = f"{report['month']:%Y-%m}"
    ws.append([f"Kassirlar oyligi — {report['month']:%m.%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    head = ["№", "Kassir", "Telefon", "Kassa", "Savdolar", "Sof savdo", "KPI %", "KPI", "Asosiy oylik", "Tuzatish", "Izoh", "Jami", "Holat"]
    ws.append(head)
    fill, thin = PatternFill("solid", fgColor="E2E8F0"), Side(style="thin", color="CBD5E1")
    for c in ws[3]:
        c.font, c.fill, c.alignment, c.border = Font(bold=True), fill, Alignment(horizontal="center", vertical="center", wrap_text=True), Border(bottom=thin)
    for i, r in enumerate(report["rows"], 1):
        ws.append([i, r["user"].full_name or "", r["user"].phone_number or "", r["store"].name, r["sales"], float(r["revenue"]), float(r["percent"]),
                   float(r["kpi"]), float(r["base"]), float(r["adjustment"]), r["note"], float(r["total"]),
                   f"To'langan {r['payout'].paid_at.astimezone(TZ):%d.%m.%Y}" if r["payout"] else "To'lanmagan"])
    t = report["totals"]
    ws.append(["", "JAMI", "", "", t["sales"], float(t["revenue"]), "", float(t["kpi"]), float(t["base"]), float(t["adjustment"]), "", float(t["total"]), ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, min_col=6, max_col=12):
        for c in row:
            if c.column != 11 and c.column != 7:
                c.number_format = "# ##0"
    for col, w in zip("ABCDEFGHIJKLM", (5, 26, 16, 22, 10, 16, 8, 14, 15, 13, 28, 16, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C4"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
