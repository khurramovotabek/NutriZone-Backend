"""
The transfer waybill as a real .xlsx (накладная на внутреннее перемещение) -- one row per product batch, sent vs received,
priced at COST (the transfer's unit_cost: latest goods-receipt cost, never the selling price), totals and a signature
block for hand signatures. Before dispatch it is the draft (planned quantities, "Принято" empty); after receiving it
is the final document with the differences.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from apps.stores.infrastructure import transfer_read
from apps.stores.infrastructure.models import StockTransfer

HEADERS = ("№", "Наименование товара", "Артикул / SKU", "Баркод", "Партия / Batch", "Срок годности", "Отправлено", "Принято", "Разница",
           "Цена за ед. (себестоимость), UZS", "Сумма, UZS", "Статус")
WIDTHS = (5, 42, 16, 16, 18, 13, 12, 11, 10, 18, 17, 15)
MONEY = '#,##0" UZS"'
STATUS_RU = {"draft": "Черновик", "ready": "Готов к отправке", "in_transit": "В пути", "receiving": "Приёмка", "partial": "Частично принят",
             "received": "Принят", "completed": "Завершён", "cancelled": "Отменён"}
THIN = Side(style="thin", color="94A3B8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1E293B")
SHORT_FILL = PatternFill("solid", fgColor="FEE2E2")
OVER_FILL = PatternFill("solid", fgColor="FEF3C7")
OK_FILL = PatternFill("solid", fgColor="DCFCE7")


def _row_status(r: dict) -> tuple[str, PatternFill | None]:
    if r["received"] is None:
        return "Отправлено" if r["sent"] else "План", None
    d = r["difference"]
    if d == 0:
        return "Совпадает", OK_FILL
    return ("Недостача", SHORT_FILL) if d < 0 else ("Излишек", OVER_FILL)


def _expiry(value: str) -> date | str:
    try:
        return date.fromisoformat(value) if value else "—"
    except ValueError:
        return value


def filename(t: StockTransfer) -> str:
    return f"{t.number}.xlsx"


def _meta(ws, t: StockTransfer, brand: str) -> int:
    """Title + document facts; returns the row of the table header."""
    last_col = get_column_letter(len(HEADERS))
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"{brand.upper()} — НАКЛАДНАЯ НА ВНУТРЕННЕЕ ПЕРЕМЕЩЕНИЕ ТОВАРОВ"
    ws["A1"].font = Font(size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    meta = [
        ("Номер документа", t.number, "Статус документа", STATUS_RU.get(t.status, t.status) + (" · расхождение на согласовании" if t.discrepancy_open else "")),
        ("Дата составления", transfer_read.local(t.created_at) or "—", "Составил", t.created_by or "—"),
        ("Отправитель", t.source.name + (f", {t.source.address}" if t.source.address else ""), "Дата отправки", transfer_read.local(t.dispatched_at) or "—"),
        ("Получатель", t.destination.name + (f", {t.destination.address}" if t.destination.address else ""), "Дата приёмки", transfer_read.local(t.received_at) or "—"),
    ]
    if t.notes:
        meta.append(("Примечание", t.notes, "", ""))
    for row, (a, b, c, d) in enumerate(meta, 3):
        for (start, end), value, bold_label in (((1, 2), a, True), ((3, 6 if c else 12), b, False), ((7, 9), c, True), ((10, 12), d, False)):
            if not value and start >= 7:
                continue
            ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
            ws.cell(row, start, value).font = Font(bold=True, color="475569") if bold_label else Font(bold=True)
    return 3 + len(meta) + 1


def _table_header(ws, head: int) -> None:
    for c, (title, width) in enumerate(zip(HEADERS, WIDTHS, strict=True), 1):
        cell = ws.cell(head, c, title)
        cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[head].height = 34


def _style(cell, c: int, v, fill) -> None:
    cell.border = BOX
    cell.alignment = Alignment(vertical="top", wrap_text=c in (2, 5, 12), horizontal="right" if c in (1, 7, 8, 9, 10, 11) else "left")
    if c == 6 and isinstance(v, date):
        cell.number_format = "DD.MM.YYYY"
    elif c in (10, 11):
        cell.number_format = MONEY
    elif c == 9 and isinstance(v, int) and v:
        cell.number_format = "+0;-0;0"
        cell.font = Font(bold=True, color="B91C1C" if v < 0 else "B45309")
    if fill is not None and c in (8, 9, 12):
        cell.fill = fill


def _lines(ws, rows: list, head: int) -> int:
    line = head
    for n, r in enumerate(rows, 1):
        line += 1
        status, fill = _row_status(r)
        values = (n, r["name"], r["sku"], r["barcode"], r["batch"] or "без партии", _expiry(r["expiry"]), r["sent"],
                  r["received"], r["difference"], float(r["unit_cost"]), float(r["value"]), status)
        for c, v in enumerate(values, 1):
            _style(ws.cell(line, c, v), c, v, fill)
    return line


def _totals(ws, rows: list, line: int, received_known: bool) -> int:
    sent = sum(r["sent"] for r in rows)
    got = sum(r["received"] or 0 for r in rows)
    value = sum(float(r["value"]) for r in rows)
    line += 1
    ws.cell(line, 2, "ИТОГО").font = Font(bold=True)
    for c, v in ((7, sent), (8, got if received_known else None), (9, (got - sent) if received_known else None), (11, value)):
        cell = ws.cell(line, c, v)
        cell.font, cell.border = Font(bold=True), BOX
        if c == 11:
            cell.number_format = MONEY
    for c in range(1, len(HEADERS) + 1):
        ws.cell(line, c).fill = PatternFill("solid", fgColor="F1F5F9")
    line += 2
    for label, v in (("Всего наименований (товар × партия)", len(rows)), ("Видов товара", len({r["item"].product_variant_id for r in rows})),
                     ("Всего отправлено, шт.", sent), ("Всего принято, шт.", got if received_known else "—"),
                     ("Общая сумма (по себестоимости), UZS", value), ("Строк с расхождением", sum(1 for r in rows if r["difference"]) if received_known else "—")):
        ws.cell(line, 2, label).font = Font(color="475569")
        cell = ws.cell(line, 3, v)
        cell.font = Font(bold=True)
        if isinstance(v, float):
            cell.number_format = MONEY
        line += 1
    return line


def _signatures(ws, t: StockTransfer, line: int) -> None:
    cols = ((1, 2), (3, 6), (7, 9), (10, 12))
    rows = [("Роль", "ФИО", "Подпись", "Дата"),
            ("Отпустил", t.dispatched_by, "", transfer_read.local(t.dispatched_at, "%d.%m.%Y")), ("Получил", t.received_by, "", transfer_read.local(t.received_at, "%d.%m.%Y")),
            ("Экспедитор", "", "", ""), ("Проверяющий", "", "", "")]
    for k, values in enumerate(rows):
        row = line + 1 + k
        if k:
            ws.row_dimensions[row].height = 30  # room for a hand signature
        for (a, b), v in zip(cols, values, strict=True):
            ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
            cell = ws.cell(row, a, v or "")
            cell.alignment = Alignment(vertical="center", horizontal="center" if not k else "left")
            if not k:
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
            for c in range(a, b + 1):
                ws.cell(row, c).border = BOX


def workbook(t: StockTransfer, brand: str = "NUTRIZONE") -> bytes:
    rows = transfer_read.batch_rows(t)
    wb = Workbook()
    ws = wb.active
    ws.title = t.number
    head = _meta(ws, t, brand)
    _table_header(ws, head)
    line = _lines(ws, rows, head)
    line = _totals(ws, rows, line, any(r["received"] is not None for r in rows))
    _signatures(ws, t, line)
    ws.freeze_panes = ws.cell(head + 1, 1)
    ws.print_title_rows = f"{head}:{head}"  # the table header repeats on every printed page
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.oddFooter.center.text = f"{t.number} · стр. &P из &N"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
