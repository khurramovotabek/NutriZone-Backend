"""
Dashboard → Loyalty & Wallet (NutriZone Coins).

- Mijozlar va coinlar (Wallet): every customer's tier and coin balance. Balances are never edited: a superadmin
  pins / releases a tier (services.set_tier) or credits / debits coins (services.adjust) -- both need a reason and
  land in the tier history / coin ledger with the staff member's name.
- Coin harakatlari (WalletTransaction): the append-only ledger, view-only, with an Excel export of the filtered rows.
- Loyalty darajalari (LoyaltyTier): the tier ladder, bulk-edited in one form (LadderFormSet checks the whole ladder).
- Dastur sozlamalari (LoyaltySettings): the singleton of programme knobs.
- Referal kodlari / Referallar: invite codes with their results, and every invitation's state.

Lists use the shared premium layout (platform.premium_admin.PremiumListMixin); add / change pages the premium form.
"""
from __future__ import annotations

from datetime import timedelta
from decimal import Decimal

from django import forms
from django.conf import settings
from django.contrib import admin, messages
from django.db.models import Count, Max, Q, Sum
from django.http import HttpResponse
from django.shortcuts import redirect
from django.urls import path, reverse
from django.utils import timezone
from django.utils.html import format_html, format_html_join

from apps.loyalty.infrastructure import admin_ui as lx
from apps.loyalty.infrastructure import services
from apps.loyalty.infrastructure.models import (
    LoyaltySettings,
    LoyaltyTier,
    Referral,
    ReferralCode,
    ReferralStatus,
    TierChange,
    Wallet,
    WalletTransaction,
    WalletTransactionKind,
)
from apps.platform.infrastructure import excel_export as xl
from apps.platform.infrastructure import premium_admin as ui
from apps.platform.infrastructure.excel_export import Col, ExcelExportMixin
from apps.platform.infrastructure.premium_admin import PremiumFormMixin, PremiumListMixin, hero_stat

CSS = ("admin/css/nutrizone-loyalty.css",)
_som = lx.som


def _titled(base, title):
    """A stock field filter with a human (Uzbek) title instead of the raw field name."""
    return type(f"{base.__name__}_{title}", (base,), {"__init__": lambda self, *a, **k: (base.__init__(self, *a, **k), setattr(self, "title", title))[0]})


def _month_start():
    return timezone.localtime(timezone.now(), ui.TZ).replace(day=1, hour=0, minute=0, second=0, microsecond=0)


def _panel(items_html, empty: str, extra="") -> str:
    body = format_html('<ul class="nzi-items">{}</ul>', items_html) if items_html else format_html('<p class="nzi-none">{}</p>', empty)
    return format_html('<div class="nzi-panel-body">{}{}</div>', body, extra)


# ================================================================ tiers


class LadderRowForm(forms.ModelForm):
    def _post_clean(self):
        # Runs model.full_clean(); the ladder-wide overlap check happens in LadderFormSet.clean instead.
        self.instance._skip_overlap_check = True
        super()._post_clean()


class LadderFormSet(forms.BaseModelFormSet):
    """
    Bulk-edit validation for the tier changelist: the submitted rows are
    merged with every tier NOT on this page, and the whole active ladder is
    checked for overlapping ranges at once. Per-row model validation skips
    its own DB-based overlap check (it would compare against neighbours'
    old, unsaved values and reject a perfectly valid multi-row shift).
    """

    def clean(self):
        super().clean()
        if any(self.errors):
            return
        from apps.loyalty.domain.tiers import TierRule, find_overlaps

        edited = {}
        for form in self.forms:
            if not form.has_changed() and form.instance.pk is None:
                continue
            inst = form.instance
            data = {**{f: getattr(inst, f) for f in ("min_spend", "max_spend", "percent", "is_active", "name", "slug")}, **form.cleaned_data}
            edited[inst.pk] = data
        others = LoyaltyTier.objects.exclude(pk__in=[pk for pk in edited if pk])
        rules = [TierRule(t.slug, t.name, t.min_spend, t.max_spend, t.percent) for t in others if t.is_active]
        rules += [TierRule(d["slug"], d["name"], d["min_spend"], d["max_spend"], d["percent"]) for d in edited.values() if d["is_active"]]
        for d in edited.values():
            if d["max_spend"] is not None and d["max_spend"] <= d["min_spend"]:
                raise forms.ValidationError(f"«{d['name']}»: maksimal summa minimaldan katta bo'lishi kerak.")
        overlaps = find_overlaps(rules)
        if overlaps:
            a, b = overlaps[0]
            raise forms.ValidationError(f"«{a.name}» va «{b.name}» oraliqlari ustma-ust tushyapti. Har bir daraja minimal summasi oldingi darajaning maksimalidan kichik bo'lmasin.")


class LoyaltyTierForm(forms.ModelForm):
    class Meta:
        model = LoyaltyTier
        fields = "__all__"
        labels = {"name": "Nomi", "slug": "Slug", "icon": "Belgi", "color": "Rang", "description": "Qisqa tavsif", "is_active": "Faol",
                  "min_spend": "Minimal xarid (so'm)", "max_spend": "Maksimal xarid (so'm)", "percent": "Keshbek (%)"}
        help_texts = {"min_spend": "Mijozning hisoblanadigan xaridlari shu summadan boshlab bu darajaga kiradi.",
                      "max_spend": "Bo'sh qoldirilsa — chegarasiz (eng yuqori daraja).",
                      "percent": "Har bir yangi buyurtma summasining shuncha foizi coin bo'lib qaytadi (1 so'm = 1 coin).",
                      "description": "Mijoz sahifasida ko'rinadi (faqat o'zbekcha).", "slug": "Dasturda ishlatiladi — o'zgartirmang."}


@admin.register(LoyaltyTier)
class LoyaltyTierAdmin(PremiumFormMixin, admin.ModelAdmin):
    """Loyalty darajalari: the whole ladder on one page -- visual steps, member counts, a calculator and the bulk editor."""

    form = LoyaltyTierForm
    list_display = ("name", "min_spend", "max_spend", "percent", "is_active")
    list_editable = ("min_spend", "max_spend", "percent", "is_active")
    list_display_links = ("name",)
    search_fields = ("name", "slug")
    prepopulated_fields = {"slug": ("name",)}
    ordering = ("min_spend",)
    list_per_page = 50
    change_list_template = "admin/loyalty/loyaltytier/change_list.html"
    fieldsets = (
        ("Daraja", {"fields": ("name", "slug", "icon", "color", "description", "is_active")}),
        ("Qoida", {"fields": ("min_spend", "max_spend", "percent"),
                   "description": "Mijozning hisoblanadigan xaridlari (yetkazilgan buyurtmalar va kassa xaridlari) [minimal, maksimal) oralig'ida bo'lsa, har bir yangi xaridda shu foiz coin oladi."}),
    )
    nz_css = CSS
    hero_icon, hero_tone, hero_noun = "fa-ranking-star", "#d4a017", "daraja"

    def get_changelist_form(self, request, **kwargs):
        kwargs.setdefault("form", LadderRowForm)
        return super().get_changelist_form(request, **kwargs)

    def get_changelist_formset(self, request, **kwargs):
        kwargs["formset"] = LadderFormSet
        return super().get_changelist_formset(request, **kwargs)

    @staticmethod
    def _members() -> dict:
        """Tier slug -> customers in it now (each wallet's current tier; skipped for very large programmes)."""
        if Wallet.objects.count() > 1500:
            return {}
        out: dict = {}
        for user_id in Wallet.objects.values_list("user_id", flat=True):
            tier = services.summary(user_id).current_tier
            if tier is not None:
                out[tier.slug] = out.get(tier.slug, 0) + 1
        return out

    def changelist_view(self, request, extra_context=None):
        tiers = list(LoyaltyTier.objects.order_by("min_spend"))
        active = [t for t in tiers if t.is_active]
        warnings = [f"«{a.name}» ({_som(a.max_spend)} so'mda tugaydi) va «{b.name}» ({_som(b.min_spend)} so'mdan boshlanadi) orasida bo'shliq bor: oradagi mijozlar pastki darajada qoladi."
                    for a, b in zip(active, active[1:], strict=False) if a.max_spend is not None and a.max_spend < b.min_spend]
        if active and active[0].min_spend > 0:
            warnings.append(f"Hech bir daraja 0 dan boshlanmaydi — yangi mijozlar «{active[0].name}» darajasiga yetguncha coin olmaydi.")
        members = self._members()
        ladder = [{"t": t, "glyph": lx.TIER_GLYPH.get(t.icon, "★"), "members": members.get(t.slug, 0), "percent": f"{t.percent.normalize():f}",
                   "min": _som(t.min_spend), "max": _som(t.max_spend) if t.max_spend is not None else "", "url": reverse("admin:loyalty_loyaltytier_change", args=[t.pk])}
                  for t in active]
        extra = {
            "title": "Loyalty darajalari", "nz_ladder": ladder, "nz_warnings": warnings, "nz_members_known": bool(members) or not Wallet.objects.exists(),
            "nz_stats": [
                ui.stat_card("Faol darajalar", len(active), "#d4a017", "fa-ranking-star", f"{len(tiers) - len(active)} ta o'chirilgan" if len(tiers) != len(active) else "barchasi faol"),
                ui.stat_card("Keshbek oralig'i", f"{active[0].percent.normalize():f}–{active[-1].percent.normalize():f}%" if active else "—", "#16a34a", "fa-percent", "eng past → eng yuqori"),
                ui.stat_card("Darajadagi mijozlar", sum(members.values()) if members else "—", "#2563eb", "fa-users", "hamyoni bor mijozlar"),
                ui.stat_card("Qo'lda belgilangan", Wallet.objects.filter(tier_override__isnull=False).count(), "#9333ea", "fa-thumbtack", "admin daraja bergan",
                             reverse("admin:loyalty_wallet_changelist") + "?manual=1"),
            ],
            "nz_window": LoyaltySettings.load().tier_window_days,
        }
        return super().changelist_view(request, {**(extra_context or {}), **extra})

    def form_hero(self, request, obj):
        if obj is None:
            return {"subtitle": "Nom, oraliq va keshbek foizini kiriting — oraliqlar boshqa darajalar bilan ustma-ust tushmasin"}
        from apps.loyalty.domain.tiers import reward_coins

        members = self._members().get(obj.slug)
        return {
            "title": f"{lx.TIER_GLYPH.get(obj.icon, '★')} {obj.name}", "subtitle": obj.description or f"slug: {obj.slug}", "tone": obj.color,
            "chips": [ui.chip("Faol", "#16a34a", "fa-circle-check") if obj.is_active else ui.chip("O'chirilgan", "#64748b", "fa-power-off")],
            "stats": [
                hero_stat("Keshbek", f"{obj.percent.normalize():f}%", obj.color, "har bir xariddan", "fa-percent"),
                hero_stat("Oraliq", f"{_som(obj.min_spend)} – {_som(obj.max_spend) if obj.max_spend is not None else '∞'}", "#2563eb", "so'm hisoblanadigan xarid", "fa-arrows-left-right"),
                hero_stat("1 000 000 so'mlik xaridga", f"+{_som(reward_coins(Decimal('1000000'), obj.percent))}", "#d4a017", "coin", "fa-coins"),
                hero_stat("Mijozlar", members if members is not None else "—", "#9333ea", "hozir shu darajada", "fa-users"),
            ],
            "links": [{"label": "Barcha darajalar", "url": reverse("admin:loyalty_loyaltytier_changelist"), "icon": "fa-ranking-star"}],
        }


# ================================================================ wallets


class WalletAdminForm(forms.ModelForm):
    """Superadmin controls: set / remove a manual tier and credit / debit coins -- each needs a reason and lands in the audit history."""

    set_tier = forms.ModelChoiceField(
        label="Daraja", queryset=LoyaltyTier.objects.filter(is_active=True).order_by("min_spend"), required=False,
        empty_label="Avtomatik — xaridlarga qarab", help_text="Mijozni xariddan qat'i nazar bitta darajaga qotirish yoki «Avtomatik» bilan qaytarish.",
    )
    tier_reason = forms.CharField(label="Daraja o'zgarishi sababi", required=False, max_length=255)
    adjust_amount = forms.IntegerField(
        label="Coin miqdori", required=False,
        help_text="Musbat — qo'shish, manfiy — ayirish. Ismingiz bilan coin tarixiga yoziladi; balans hech qachon to'g'ridan-to'g'ri o'zgartirilmaydi.",
    )
    adjust_note = forms.CharField(label="Sabab", required=False, max_length=255)

    class Meta:
        model = Wallet
        fields = ("user",)
        labels = {"user": "Mijoz"}

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        if self.instance and self.instance.pk:
            self.fields["set_tier"].initial = self.instance.tier_override_id

    def clean(self):
        data = super().clean()
        amount = data.get("adjust_amount")
        if amount:
            if not (data.get("adjust_note") or "").strip():
                self.add_error("adjust_note", "Qo'lda tuzatish uchun sabab yozilishi shart.")
            balance = self.instance.balance if self.instance and self.instance.pk else 0
            if balance + amount < 0:
                self.add_error("adjust_amount", f"Balans manfiy bo'lib qoladi (hozirgi balans: {_som(balance)}).")
        current = self.instance.tier_override_id if self.instance and self.instance.pk else None
        chosen = data.get("set_tier")
        if (chosen.pk if chosen else None) != current and not (data.get("tier_reason") or "").strip():
            self.add_error("tier_reason", "Darajani o'zgartirish uchun sabab yozilishi shart.")
        return data


class ManualTierFilter(admin.SimpleListFilter):
    title = "Daraja manbai"
    parameter_name = "manual"

    def lookups(self, request, model_admin):
        return [("1", "Admin belgilagan"), ("0", "Avtomatik")]

    def queryset(self, request, queryset):
        if self.value() in ("0", "1"):
            return queryset.filter(tier_override__isnull=self.value() == "0")
        return queryset


class CoinsFilter(admin.SimpleListFilter):
    title = "Coinlar"
    parameter_name = "coins"

    def lookups(self, request, model_admin):
        return [("has", "Coini bor"), ("none", "Coinsiz"), ("100k", "≥ 100 000"), ("expiring", "30 kunda yonadi")]

    def queryset(self, request, queryset):
        if self.value() == "expiring":
            now = timezone.now()
            return queryset.filter(transactions__remaining__gt=0, transactions__expires_at__gt=now, transactions__expires_at__lte=now + timedelta(days=30)).distinct()
        return {"has": queryset.filter(balance__gt=0), "none": queryset.filter(balance=0), "100k": queryset.filter(balance__gte=100_000)}.get(self.value(), queryset)


@admin.register(Wallet)
class WalletAdmin(ExcelExportMixin, PremiumListMixin, PremiumFormMixin, admin.ModelAdmin):
    """
    Mijozlar va coinlar: every customer with their tier and coin balance. Only a superadmin can change a tier or a
    balance -- always with a reason, always recorded (tier history + coin ledger with the staff member's name).
    """
    excel_title = "Mijozlar va coinlar"
    excel_select_related = ("user", "tier_override")
    excel_columns = [
        Col("Mijoz", lambda w: w.user.full_name if w.user else "", 24), Col("Telefon", lambda w: w.user.phone_number if w.user else "", 15),
        Col("Balans (coin)", "balance", 13, number=True), Col("Jami olingan", "lifetime_earned", 13, number=True), Col("Jami sarflangan", "lifetime_spent", 14, number=True),
        Col("Qo'lda daraja", lambda w: w.tier_override.name if w.tier_override else "", 16), Col("Yaratilgan", "created_at", 17),
    ]


    form = WalletAdminForm
    list_display = ("user", "balance", "updated_at")
    list_filter = (ManualTierFilter, ("tier_override", _titled(admin.RelatedFieldListFilter, "Qo'lda berilgan daraja")))
    search_fields = ("user__full_name", "user__phone_number", "user__telegram_username")
    autocomplete_fields = ("user",)
    list_select_related = ("user", "tier_override")
    ordering = ("-balance", "-updated_at")
    list_per_page = 30
    readonly_fields = ("ledger_panel", "lots_panel", "tiers_panel")

    nz_icon, nz_tone, nz_title = "fa-wallet", "#d4a017", "Mijozlar va coinlar"
    nz_subtitle = "Har bir mijozning darajasi, coin balansi va harakati. Daraja yoki balansni faqat superadmin sabab ko'rsatib o'zgartiradi."
    nz_columns = (("Mijoz", ""), ("Daraja", ""), ("Oxirgi harakat", ""), ("Jami ishlagan", "r"), ("Sarflagan", "r"), ("30 kunda yonadi", "r"), ("Balans", "r"))
    nz_grid = "minmax(230px,2fr) minmax(150px,1.3fr) 130px 120px 120px 120px 120px"
    nz_search = "Ism, telefon yoki Telegram…"
    nz_empty = ("Hali hamyon yo'q", "Mijoz birinchi xaridni qilganda yoki saytda loyalty sahifasini ochganda hamyon yaratiladi.")
    nz_sorts = (("Balans", ""), ("Oxirgi o'zgargan", "-3"))
    nz_css = CSS
    hero_icon, hero_tone, hero_noun = "fa-wallet", "#d4a017", "hamyon"

    def get_queryset(self, request):
        from django.db.models import Prefetch

        from apps.identity.infrastructure.models import UserAvatarAsset

        return super().get_queryset(request).prefetch_related(Prefetch("user__avatar_assets", queryset=UserAvatarAsset.objects.order_by("-created_at")))

    def get_list_filter(self, request):
        return (*self.list_filter, CoinsFilter)

    def nz_stats(self, request):
        now, month = timezone.now(), _month_start()
        tx = WalletTransaction.objects.filter(created_at__gte=month)
        w = Wallet.objects.aggregate(n=Count("id"), coins=Sum("balance"), holders=Count("id", filter=Q(balance__gt=0)), manual=Count("id", filter=Q(tier_override__isnull=False)))
        expiring = WalletTransaction.objects.filter(remaining__gt=0, expires_at__gt=now, expires_at__lte=now + timedelta(days=30)).aggregate(v=Sum("remaining"))["v"] or 0
        tx_url = reverse("admin:loyalty_wallettransaction_changelist")
        return [
            ui.stat_card("Muomaladagi coinlar", _som(w["coins"]), "#d4a017", "fa-coins", f"{w['holders']} mijozda · {w['n']} hamyon", "?coins=has"),
            ui.stat_card("Shu oy berilgan", f"+{_som(tx.filter(amount__gt=0).aggregate(v=Sum('amount'))['v'])}", "#16a34a", "fa-arrow-trend-up", "keshbek, bonus, tuzatish", tx_url + "?dir=in"),
            ui.stat_card("Shu oy sarflangan", _som(-(tx.filter(kind=WalletTransactionKind.SPEND).aggregate(v=Sum("amount"))["v"] or 0)), "#ea580c", "fa-bag-shopping", "xaridda to'langan",
                         tx_url + "?kind__exact=spend"),
            ui.stat_card("30 kunda yonadi", _som(expiring), "#dc2626" if expiring else "#64748b", "fa-hourglass-half", "ishlatilmasa muddati tugaydi", "?coins=expiring"),
            ui.stat_card("Qo'lda daraja", w["manual"], "#9333ea", "fa-thumbtack", "admin qotirgan", "?manual=1"),
        ]

    def nz_tabs(self, request):
        w = Wallet.objects.aggregate(n=Count("id"), has=Count("id", filter=Q(balance__gt=0)), none=Count("id", filter=Q(balance=0)), big=Count("id", filter=Q(balance__gte=100_000)))
        return ui.tabs_for(request, "coins", [("", "Hammasi", w["n"]), ("has", "Coini bor", w["has"]), ("none", "Coinsiz", w["none"]), ("100k", "≥ 100 000", w["big"])],
                           {"has": "#d4a017", "none": "#64748b", "100k": "#16a34a"})

    def nz_hero_actions(self, request):
        return [{"label": "Darajalar", "url": reverse("admin:loyalty_loyaltytier_changelist"), "icon": "fa-ranking-star"},
                {"label": "Coin harakatlari", "url": reverse("admin:loyalty_wallettransaction_changelist"), "icon": "fa-arrow-right-arrow-left"},
                {"label": "Sozlamalar", "url": reverse("admin:loyalty_loyaltysettings_changelist"), "icon": "fa-sliders"}]

    def nz_prepare(self, request, objs):
        now = timezone.now()
        ids = [w.pk for w in objs]
        expiring = dict(WalletTransaction.objects.filter(wallet_id__in=ids, remaining__gt=0, expires_at__gt=now, expires_at__lte=now + timedelta(days=30))
                        .values("wallet_id").annotate(v=Sum("remaining")).values_list("wallet_id", "v"))
        last = dict(WalletTransaction.objects.filter(wallet_id__in=ids).values("wallet_id").annotate(m=Max("created_at")).values_list("wallet_id", "m"))
        for w in objs:
            w._sum, w._expiring, w._last = services.summary(w.user_id), expiring.get(w.pk, 0), last.get(w.pk)

    def nz_row(self, w):
        s = w._sum
        nxt = format_html("<small>{} so'm → {}</small>", _som(s.amount_to_next), s.next_tier.name) if s.next_tier and s.amount_to_next is not None and not s.manual else ""
        return [
            ui.user_cell(w.user), format_html("{}{}", lx.tier_chip(s.current_tier, s.manual, w.tier_override_reason), nxt),
            format_html('<span title="{}">{}</span>', ui.local(w._last), ui.ago(w._last)) if w._last else ui.muted("harakat yo'q"),
            _som(w.lifetime_earned), _som(w.lifetime_spent),
            format_html('<b style="color:#dc2626">{}</b>', _som(w._expiring)) if w._expiring else ui.muted(), lx.coins(w.balance),
        ]

    # ---------------------------------------------------------------- the wallet card

    def get_fieldsets(self, request, obj=None):
        if obj is None:
            return (("Mijoz", {"fields": ("user",), "description": "Hamyon odatda avtomatik yaratiladi. Qo'lda faqat bonus berish uchun kerak bo'lsa yarating."}),)
        panels = (
            ("Coin harakatlari", {"classes": ("nzi-panel",), "fields": ("ledger_panel",)}),
            ("Muddati tugaydigan coinlar", {"classes": ("nzi-panel",), "fields": ("lots_panel",)}),
            ("Daraja tarixi", {"classes": ("nzi-panel",), "fields": ("tiers_panel",)}),
        )
        if not request.user.is_superuser:
            return panels
        return (
            ("Coin qo'shish yoki ayirish", {"fields": ("adjust_amount", "adjust_note")}),
            ("Darajani qo'lda belgilash", {"fields": ("set_tier", "tier_reason")}),
            *panels,
        )

    def get_readonly_fields(self, request, obj=None):
        return (*self.readonly_fields, "user") if obj else self.readonly_fields

    def form_hero(self, request, obj):
        if obj is None:
            return {"subtitle": "Mijozni qidirib tanlang"}
        s = services.summary(obj.user_id)
        risk = services.tier_risk(obj.user_id)
        chips = [lx.tier_chip(s.current_tier, s.manual, obj.tier_override_reason)]
        if s.manual:
            chips.append(ui.chip(f"Qo'lda: {obj.tier_override_reason or 'sabab yozilmagan'}", "#9333ea", "fa-thumbtack"))
        if risk is not None:
            chips.append(ui.chip(f"Daraja xavf ostida: {ui.local(risk.drop_at, '%d.%m.%Y')} gacha yana {_som(risk.needed)} so'm xarid kerak", "#dc2626", "fa-triangle-exclamation"))
        nxt = f"{_som(s.amount_to_next)} so'm → {s.next_tier.name}" if s.next_tier and s.amount_to_next is not None else "eng yuqori daraja"
        return {
            "title": obj.user.full_name or obj.user.phone_number, "subtitle": " · ".join(x for x in (obj.user.phone_number, f"@{obj.user.telegram_username}" if obj.user.telegram_username else "") if x),
            "image": ui.avatar_url(obj.user) or "", "icon": "fa-wallet", "tone": s.current_tier.color if s.current_tier else "#d4a017", "chips": chips,
            "stats": [
                hero_stat("Balans", _som(obj.balance), "#d4a017", "coin = so'm", "fa-coins"),
                hero_stat("Jami ishlagan", _som(obj.lifetime_earned), "#16a34a", f"sarflagan {_som(obj.lifetime_spent)}", "fa-arrow-trend-up"),
                hero_stat("Hisoblanadigan xarid", f"{_som(s.total_spend)} so'm", "#2563eb", f"{s.order_count} ta xarid", "fa-receipt"),
                hero_stat("Keyingi daraja", f"{s.progress_percent}%", "#9333ea", nxt, "fa-stairs"),
            ],
            "links": [
                {"label": "Mijoz kartasi", "url": reverse("admin:identity_user_change", args=[obj.user_id]), "icon": "fa-user"},
                {"label": "Barcha harakatlar", "url": reverse("admin:loyalty_wallettransaction_changelist") + f"?wallet__id__exact={obj.pk}", "icon": "fa-arrow-right-arrow-left"},
            ],
        }

    @admin.display(description="")
    def ledger_panel(self, obj):
        rows = list(obj.transactions.order_by("-created_at")[:15]) if obj and obj.pk else []
        items = format_html_join("", '<li><span class="nzi-ico" style="--t:{}"><i class="fas {}" aria-hidden="true"></i></span><span><b>{}</b><small>{}{}{}</small></span><span class="nzi-when">{}<small>qoldi {}</small></span></li>', (
            (lx.KIND.get(t.kind, ("", "fa-circle", "#64748b"))[2], lx.KIND.get(t.kind, ("", "fa-circle", ""))[1], lx.KIND.get(t.kind, (t.kind,))[0],
             ui.local(t.created_at), f" · {t.order_number}" if t.order_number else "", f" · {t.note}" if t.note else "", lx.coins(t.amount, signed=True), _som(t.balance_after))
            for t in rows))
        total = obj.transactions.count() if obj and obj.pk else 0
        more = format_html('<a class="nzi-more" href="{}?wallet__id__exact={}">Barcha {} ta harakat →</a>', reverse("admin:loyalty_wallettransaction_changelist"), obj.pk, total) if total > len(rows) else ""
        return _panel(items, "Hali coin harakati yo'q.", more)

    @admin.display(description="")
    def lots_panel(self, obj):
        now = timezone.now()
        lots = list(obj.transactions.filter(remaining__gt=0).order_by("expires_at")[:10]) if obj and obj.pk else []
        items = format_html_join("", '<li><span class="nzi-ico" style="--t:{}"><i class="fas fa-hourglass-half" aria-hidden="true"></i></span><span><b>{} coin</b><small>{} · {} dan</small></span><span class="nzi-when"><b>{}</b><small>{}</small></span></li>', (
            ("#dc2626" if t.expires_at and t.expires_at <= now + timedelta(days=30) else "#16a34a", _som(t.remaining), lx.KIND.get(t.kind, (t.kind,))[0], ui.local(t.created_at, "%d.%m.%Y"),
             ui.local(t.expires_at, "%d.%m.%Y") if t.expires_at else "muddatsiz", ui.ago(t.expires_at) if t.expires_at else "")
            for t in lots))
        return _panel(items, "Ishlatilmagan coin partiyasi yo'q.", format_html('<p class="nzi-help">{}</p>', "Coinlar partiyalar bo'yicha saqlanadi: xaridda avval muddati yaqini sarflanadi (FIFO)."))

    @admin.display(description="")
    def tiers_panel(self, obj):
        rows = list(TierChange.objects.filter(wallet=obj).order_by("-created_at")[:12]) if obj and obj.pk else []
        items = format_html_join("", '<li><span class="nzi-ico" style="--t:{}"><i class="fas fa-ranking-star" aria-hidden="true"></i></span><span><b>{} → {}</b><small>{} · {}{}</small></span><span class="nzi-when"><small>{}</small></span></li>', (
            (lx.TIER_CHANGE.get(c.kind, ("", "#64748b"))[1], c.from_tier or "—", c.to_tier or "avtomatik", lx.TIER_CHANGE.get(c.kind, (c.kind,))[0], c.actor or "tizim",
             f" · {c.reason}" if c.reason else "", ui.local(c.created_at)) for c in rows))
        return _panel(items, "Daraja hali o'zgarmagan.")

    def save_model(self, request, obj, form, change):
        if not change:
            obj.save()
        if not request.user.is_superuser:
            return
        actor = request.user.get_full_name() or request.user.get_username()
        chosen = form.cleaned_data.get("set_tier")
        if change and (chosen.pk if chosen else None) != obj.tier_override_id:
            services.set_tier(obj.user_id, chosen, form.cleaned_data.get("tier_reason", "").strip(), actor)
            messages.success(request, f"Daraja yangilandi: {chosen.name if chosen else 'avtomatik'}.")
        amount = form.cleaned_data.get("adjust_amount")
        if amount:
            try:
                services.adjust(obj.user_id, amount, form.cleaned_data.get("adjust_note", "").strip(), actor=actor)
                messages.success(request, f"Balans {amount:+,} coinga o'zgartirildi.".replace(",", " "))
            except ValueError as exc:
                messages.error(request, str(exc))
        obj.refresh_from_db()

    def has_delete_permission(self, request, obj=None) -> bool:
        return False


# ================================================================ ledger


class DirectionFilter(admin.SimpleListFilter):
    title = "Yo'nalish"
    parameter_name = "dir"

    def lookups(self, request, model_admin):
        return [("in", "Kirim (+)"), ("out", "Chiqim (−)")]

    def queryset(self, request, queryset):
        return {"in": queryset.filter(amount__gt=0), "out": queryset.filter(amount__lt=0)}.get(self.value(), queryset)


class OpenLotFilter(admin.SimpleListFilter):
    title = "Partiya"
    parameter_name = "lot"

    def lookups(self, request, model_admin):
        return [("open", "Ishlatilmagan qoldig'i bor"), ("soon", "30 kunda yonadi")]

    def queryset(self, request, queryset):
        now = timezone.now()
        if self.value() == "open":
            return queryset.filter(remaining__gt=0)
        if self.value() == "soon":
            return queryset.filter(remaining__gt=0, expires_at__gt=now, expires_at__lte=now + timedelta(days=30))
        return queryset


@admin.register(WalletTransaction)
class WalletTransactionAdmin(ExcelExportMixin, PremiumListMixin, admin.ModelAdmin):
    """Coin harakatlari: the append-only ledger (view-only) with an Excel export of exactly what is filtered."""
    excel_title = "Coin harakatlari"
    excel_select_related = ("wallet__user",)
    excel_columns = [
        Col("Vaqt", "created_at", 17), Col("Mijoz", lambda t: t.wallet.user.full_name if t.wallet and t.wallet.user else "", 24),
        Col("Telefon", lambda t: t.wallet.user.phone_number if t.wallet and t.wallet.user else "", 15), Col("Turi", "kind", 16),
        Col("Coin", "amount", 10, number=True), Col("Keyingi balans", "balance_after", 13), Col("Buyurtma", "order_number", 16),
        Col("Buyurtma summasi", "order_amount", 15, money=True), Col("Daraja", "tier_name", 14), Col("Muddati", "expires_at", 17), Col("Izoh", "note", 28),
    ]


    list_display = ("created_at", "amount")
    list_filter = (DirectionFilter, OpenLotFilter, ("created_at", _titled(admin.DateFieldListFilter, "Sana")))
    search_fields = ("order_number", "wallet__user__full_name", "wallet__user__phone_number", "note", "actor")
    list_select_related = ("wallet__user",)
    ordering = ("-created_at",)
    list_per_page = 40

    nz_icon, nz_tone, nz_title = "fa-arrow-right-arrow-left", "#16a34a", "Coin harakatlari"
    nz_subtitle = "Har bir coin qayerdan kelgani va qayerga ketgani: keshbek, referal, xaridda ishlatilgan, qaytarilgan va yonib ketganlar. Yozuvlar o'zgartirilmaydi."
    nz_columns = (("Mijoz", ""), ("Turi", ""), ("Asos", ""), ("Partiya", ""), ("Vaqt", ""), ("Balans keyin", "r"), ("Miqdor", "r"))
    nz_grid = "minmax(220px,1.8fr) minmax(170px,1.4fr) minmax(160px,1.4fr) minmax(130px,1fr) 100px 110px 110px"
    nz_search = "Buyurtma raqami, mijoz, izoh…"
    nz_empty = ("Hali coin harakati yo'q", "Keshbek yetkazilgan buyurtmalardan keyin avtomatik yoziladi.")
    nz_sorts = (("Yangilari", ""), ("Eskilari", "1"), ("Katta miqdor", "-2"))
    nz_css = CSS

    def get_list_filter(self, request):
        return (("kind", _titled(admin.ChoicesFieldListFilter, "Turi")), *self.list_filter)

    def has_add_permission(self, request) -> bool:
        return False

    def has_change_permission(self, request, obj=None) -> bool:
        return False

    def has_delete_permission(self, request, obj=None) -> bool:
        return False

    def get_urls(self):
        return [path("export.xlsx", self.admin_site.admin_view(self.export_view), name="loyalty_wallettransaction_export"), *super().get_urls()]

    def nz_stats(self, request):
        tx = WalletTransaction.objects.filter(created_at__gte=_month_start())
        by = dict(tx.values("kind").annotate(v=Sum("amount")).values_list("kind", "v"))
        issued = sum(v for k, v in by.items() if (v or 0) > 0)
        return [
            ui.stat_card("Shu oy berilgan", f"+{_som(issued)}", "#16a34a", "fa-arrow-trend-up", f"keshbek {_som(by.get('earn'))} · referal {_som(by.get('referral'))}", "?dir=in"),
            ui.stat_card("Shu oy sarflangan", _som(-(by.get("spend") or 0)), "#ea580c", "fa-bag-shopping", "xaridda to'langan", "?kind__exact=spend"),
            ui.stat_card("Yonib ketgan", _som(-(by.get("expiry") or 0)), "#64748b", "fa-hourglass-end", "shu oy muddati tugagan", "?kind__exact=expiry"),
            ui.stat_card("Qaytarib olingan", _som(-(by.get("reversal") or 0)), "#dc2626", "fa-ban", "bekor qilingan buyurtmalar", "?kind__exact=reversal"),
        ]

    def nz_tabs(self, request):
        counts = dict(WalletTransaction.objects.values_list("kind").annotate(n=Count("id")))
        return ui.tabs_for(request, "kind__exact", [("", "Hammasi", sum(counts.values()))] + [(k, lx.KIND[k][0].split(" (")[0], counts[k]) for k in lx.KIND if counts.get(k)],
                           {k: v[2] for k, v in lx.KIND.items()})

    def nz_hero_actions(self, request):
        q = request.GET.copy()
        q.pop("p", None)
        return [{"label": "Excel'ga yuklash", "url": reverse("admin:loyalty_wallettransaction_export") + (f"?{q.urlencode()}" if q else ""), "icon": "fa-file-excel"},
                {"label": "Mijozlar va coinlar", "url": reverse("admin:loyalty_wallet_changelist"), "icon": "fa-wallet"}]

    def nz_href(self, obj):
        return reverse("admin:loyalty_wallet_change", args=[obj.wallet_id])

    def nz_row(self, t):
        if t.amount > 0:
            lot = format_html("{} qoldi<small>{}</small>", _som(t.remaining), f"muddati {ui.local(t.expires_at, '%d.%m.%Y')}" if t.expires_at else "muddatsiz") if t.remaining else ui.muted("to'liq ishlatilgan")
        else:
            lot = ui.muted()
        basis = " · ".join(x for x in (t.order_number, f"{t.tier_name} {t.tier_percent.normalize():f}%" if t.tier_name and t.tier_percent is not None else t.tier_name, t.note) if x)
        return [
            ui.user_cell(t.wallet.user), lx.kind_chip(t.kind), format_html("{}<small>{}</small>", basis or "—", f"kim: {t.actor}" if t.actor else ""), lot,
            format_html('<span title="{}">{}</span><small>{}</small>', ui.local(t.created_at), ui.local(t.created_at, "%d.%m.%Y"), ui.local(t.created_at, "%H:%M")),
            _som(t.balance_after), lx.coins(t.amount, signed=True),
        ]

    def export_view(self, request):
        """The filtered ledger as .xlsx (same filters / search as the page)."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        if not self.has_view_permission(request):
            from django.core.exceptions import PermissionDenied

            raise PermissionDenied
        cl = self.get_changelist_instance(request)
        wb = Workbook()
        ws = wb.active
        ws.title = "Coin harakatlari"
        ws.append(["Sana", "Mijoz", "Telefon", "Turi", "Miqdor", "Balans keyin", "Partiya qoldig'i", "Muddati", "Buyurtma", "Daraja", "Izoh", "Kim"])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="B45309")
        for t in cl.get_queryset(request).select_related("wallet__user")[:20000]:
            ws.append([ui.local(t.created_at), t.wallet.user.full_name, t.wallet.user.phone_number, lx.KIND.get(t.kind, (t.kind,))[0], t.amount, t.balance_after,
                       t.remaining if t.amount > 0 else None, ui.local(t.expires_at, "%d.%m.%Y") if t.expires_at else "", t.order_number, t.tier_name, t.note, t.actor])
        for col, width in zip("ABCDEFGHIJKL", (17, 26, 16, 30, 12, 13, 15, 12, 18, 12, 30, 18), strict=False):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        import io

        buf = io.BytesIO()
        wb.save(buf)
        response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="coin_harakatlari_{timezone.localdate():%Y%m%d}.xlsx"'
        return response


# ================================================================ settings


class LoyaltySettingsForm(forms.ModelForm):
    class Meta:
        model = LoyaltySettings
        fields = "__all__"
        labels = {
            "tier_window_days": "Daraja hisoblash davri (kun)", "coin_expiry_days": "Coin amal qilish muddati (kun)", "expiry_reminder_days": "Eslatma (muddatdan necha kun oldin)",
            "redeem_enabled": "Xaridda coin bilan to'lash", "max_redeem_percent": "Buyurtmaning coin bilan to'lanadigan ulushi (maks. %)", "min_redeem_coins": "Bir to'lovda minimal coin",
            "referral_enabled": "Referal dasturi", "referrer_reward": "Taklif qilganga mukofot (coin)", "referee_reward": "Yangi mijozga mukofot (coin)",
            "referral_min_order": "Referal: birinchi buyurtma minimal summasi (so'm)", "subscribe_gift_coins": "Telegram yangiliklariga obuna sovg'asi (coin)",
            "telegram_notifications": "Telegram xabarnomalari",
        }
        help_texts = {
            "tier_window_days": "Daraja oxirgi shuncha kundagi xaridlar bo'yicha hisoblanadi (siljuvchi davr). 0 — butun tarix.",
            "coin_expiry_days": "Har bir coin partiyasi berilganidan shuncha kun o'tib ishlatilmasa yonib ketadi. 0 — muddatsiz.",
            "expiry_reminder_days": "Mijozga coin yonishidan shuncha kun oldin Telegram va bildirishnoma orqali eslatiladi.",
            "max_redeem_percent": "1 coin = 1 so'm. Mijoz buyurtma summasining ko'pi bilan shuncha foizini coin bilan to'lay oladi.",
            "min_redeem_coins": "Bundan kam coin bilan to'lab bo'lmaydi.",
            "referral_min_order": "Mukofot taklif qilingan do'stning birinchi buyurtmasi yetkazilganda va shu summadan kam bo'lmasa beriladi.",
            "subscribe_gift_coins": "Botda yangiliklarga obuna bo'lgan mijoz birinchi xaridida bir marta oladi. 0 — sovg'a yo'q.",
            "telegram_notifications": "Coin tushishi, daraja oshishi, referal bonuslari va yonish eslatmalari.",
            "referrer_reward": "", "referee_reward": "",
        }


@admin.register(LoyaltySettings)
class LoyaltySettingsAdmin(PremiumFormMixin, admin.ModelAdmin):
    """Dastur sozlamalari: the singleton -- the changelist jumps straight to the one settings row."""

    form = LoyaltySettingsForm
    fieldsets = (
        ("Darajalar", {"fields": ("tier_window_days",)}),
        ("Coin muddati", {"fields": ("coin_expiry_days", "expiry_reminder_days")}),
        ("Xaridda coin bilan to'lash", {"fields": ("redeem_enabled", "max_redeem_percent", "min_redeem_coins")}),
        ("Referal dasturi", {"fields": ("referral_enabled", "referrer_reward", "referee_reward", "referral_min_order")}),
        ("Telegram", {"fields": ("subscribe_gift_coins", "telegram_notifications")}),
        ("Mijoz uchun qanday ko'rinadi", {"classes": ("nzi-panel",), "fields": ("example_panel",)}),
    )
    readonly_fields = ("example_panel",)
    nz_css = CSS
    hero_icon, hero_tone = "fa-sliders", "#d4a017"

    def changelist_view(self, request, extra_context=None):
        return redirect(reverse("admin:loyalty_loyaltysettings_change", args=[LoyaltySettings.load().pk]))

    def has_add_permission(self, request) -> bool:
        return False

    def has_delete_permission(self, request, obj=None) -> bool:
        return False

    def response_change(self, request, obj):
        messages.success(request, "Dastur sozlamalari saqlandi — o'zgarishlar darhol kuchga kiradi.")
        return redirect(request.path)

    def form_hero(self, request, obj):
        def onoff(flag, on, off):
            return ui.chip(on, "#16a34a", "fa-circle-check") if flag else ui.chip(off, "#64748b", "fa-power-off")

        return {
            "title": "Loyalty dasturi sozlamalari", "subtitle": f"oxirgi o'zgarish: {ui.local(obj.updated_at)}" if obj else "",
            "chips": [onoff(obj.redeem_enabled, "Coin bilan to'lash yoqilgan", "Coin bilan to'lash o'chiq"), onoff(obj.referral_enabled, "Referal yoqilgan", "Referal o'chiq"),
                      onoff(obj.telegram_notifications, "Telegram xabarnomalari", "Telegram xabarnomalari o'chiq")],
            "stats": [
                hero_stat("Daraja davri", f"{obj.tier_window_days} kun" if obj.tier_window_days else "butun tarix", "#2563eb", "xaridlar hisoblanadi", "fa-calendar"),
                hero_stat("Coin muddati", f"{obj.coin_expiry_days} kun" if obj.coin_expiry_days else "muddatsiz", "#dc2626", f"eslatma {obj.expiry_reminder_days} kun oldin", "fa-hourglass-half"),
                hero_stat("Coin bilan to'lash", f"≤ {obj.max_redeem_percent}%", "#ea580c", f"kamida {_som(obj.min_redeem_coins)} coin", "fa-bag-shopping"),
                hero_stat("Referal", f"{_som(obj.referrer_reward)} + {_som(obj.referee_reward)}", "#9333ea", "taklif qilgan + yangi mijoz", "fa-user-plus"),
            ],
            "links": [{"label": "Darajalar", "url": reverse("admin:loyalty_loyaltytier_changelist"), "icon": "fa-ranking-star"},
                      {"label": "Mijozlar va coinlar", "url": reverse("admin:loyalty_wallet_changelist"), "icon": "fa-wallet"}],
        }

    @admin.display(description="")
    def example_panel(self, obj):
        """What the saved settings mean for one customer -- a 1 000 000 so'm order at the entry and the top tier."""
        from apps.loyalty.domain.tiers import reward_coins

        tiers = list(LoyaltyTier.objects.filter(is_active=True).order_by("min_spend"))
        order = Decimal("1000000")
        lines = []
        if tiers:
            lo, hi = tiers[0], tiers[-1]
            lines.append(("fa-coins", "#d4a017", f"1 000 000 so'mlik buyurtma: {lo.name} darajasida +{_som(reward_coins(order, lo.percent))} coin, {hi.name} darajasida +{_som(reward_coins(order, hi.percent))} coin."))
        if obj.redeem_enabled:
            lines.append(("fa-bag-shopping", "#ea580c", f"Shu buyurtmaning {_som(order * obj.max_redeem_percent / 100)} so'migacha coin bilan to'lash mumkin (kamida {_som(obj.min_redeem_coins)} coin)."))
        else:
            lines.append(("fa-power-off", "#64748b", "Coin bilan to'lash hozir o'chiq — mijozlar faqat coin yig'adi."))
        if obj.coin_expiry_days:
            lines.append(("fa-hourglass-half", "#dc2626", f"Bugun tushgan coin {ui.local(timezone.now() + timedelta(days=obj.coin_expiry_days), '%d.%m.%Y')} gacha amal qiladi; {obj.expiry_reminder_days} kun oldin eslatiladi."))
        if obj.referral_enabled:
            minimum = f", buyurtma kamida {_som(obj.referral_min_order)} so'm bo'lsa" if obj.referral_min_order else ""
            lines.append(("fa-user-plus", "#9333ea", f"Do'stini taklif qilgan mijoz {_som(obj.referrer_reward)} coin, do'sti {_som(obj.referee_reward)} coin oladi — do'stning birinchi buyurtmasi yetkazilganda{minimum}."))
        items = format_html_join("", '<li><span class="nzi-ico" style="--t:{}"><i class="fas {}" aria-hidden="true"></i></span><span><b>{}</b></span></li>', ((c, i, t) for i, c, t in lines))
        return _panel(items, "", format_html('<p class="nzi-help">{}</p>', "Misol saqlangan sozlamalar bo'yicha hisoblanadi — o'zgartirgandan keyin saqlang."))


# ================================================================ referrals


def _share_link(code: str) -> str:
    base = (getattr(settings, "FRONTEND_URL", "") or "").rstrip("/")
    return f"{base}/?ref={code}" if base else f"/?ref={code}"


@admin.register(ReferralCode)
class ReferralCodeAdmin(PremiumListMixin, admin.ModelAdmin):
    """Referal kodlari: every customer's invite code with what it brought -- invitations, rewarded friends, coins earned."""

    list_display = ("code", "created_at")
    search_fields = ("code", "user__full_name", "user__phone_number")
    list_select_related = ("user",)
    ordering = ("-created_at",)
    list_per_page = 30

    nz_icon, nz_tone, nz_title = "fa-ticket", "#9333ea", "Referal kodlari"
    nz_subtitle = "Har bir mijozning taklif kodi va natijasi: nechta do'st keldi, nechtasi xarid qildi va qancha coin ishlandi. Kod mijoz loyalty sahifasini ochganda yaratiladi."
    nz_columns = (("Kod", ""), ("Egasi", ""), ("Taklif qilganlar", "r"), ("Xarid qilganlar", "r"), ("Ishlagan coin", "r"), ("Ulashish havolasi", ""), ("Yaratilgan", "r"))
    nz_grid = "150px minmax(220px,1.8fr) 120px 120px 120px minmax(170px,1.3fr) 110px"
    nz_search = "Kod yoki mijoz…"
    nz_empty = ("Hali referal kodi yo'q", "Mijoz saytdagi loyalty sahifasini ochganda kodi avtomatik yaratiladi.")
    nz_sorts = (("Yangilari", ""), ("Eskilari", "2"))
    nz_css = CSS

    def has_add_permission(self, request) -> bool:
        return False

    def has_change_permission(self, request, obj=None) -> bool:
        return False

    def has_delete_permission(self, request, obj=None) -> bool:
        return False

    def nz_stats(self, request):
        r = Referral.objects.aggregate(n=Count("id"), ok=Count("id", filter=Q(status=ReferralStatus.REWARDED)))
        paid = WalletTransaction.objects.filter(kind=WalletTransactionKind.REFERRAL).aggregate(v=Sum("amount"))["v"] or 0
        active = Referral.objects.values("referrer_id").distinct().count()
        cfg = LoyaltySettings.load()
        return [
            ui.stat_card("Kodlar", ReferralCode.objects.count(), "#9333ea", "fa-ticket", f"{active} tasi kimnidir taklif qilgan"),
            ui.stat_card("Takliflar", r["n"], "#2563eb", "fa-user-plus", "kod bilan kelgan mijozlar", reverse("admin:loyalty_referral_changelist")),
            ui.stat_card("Xarid qilganlar", r["ok"], "#16a34a", "fa-circle-check", f"{round(r['ok'] * 100 / r['n']) if r['n'] else 0}% konversiya", reverse("admin:loyalty_referral_changelist") + "?status__exact=rewarded"),
            ui.stat_card("To'langan bonus", _som(paid), "#d4a017", "fa-coins", f"hozir {_som(cfg.referrer_reward)} + {_som(cfg.referee_reward)} coin" if cfg.referral_enabled else "dastur o'chiq"),
        ]

    def nz_hero_actions(self, request):
        return [{"label": "Referallar", "url": reverse("admin:loyalty_referral_changelist"), "icon": "fa-people-arrows"},
                {"label": "Sozlamalar", "url": reverse("admin:loyalty_loyaltysettings_changelist"), "icon": "fa-sliders"}]

    def nz_prepare(self, request, objs):
        users = [c.user_id for c in objs]
        made = {r["referrer_id"]: r for r in Referral.objects.filter(referrer_id__in=users).values("referrer_id").annotate(n=Count("id"), ok=Count("id", filter=Q(status=ReferralStatus.REWARDED)))}
        earned = dict(WalletTransaction.objects.filter(kind=WalletTransactionKind.REFERRAL, wallet__user_id__in=users, referral__referrer_id__in=users)
                      .values("wallet__user_id").annotate(v=Sum("amount")).values_list("wallet__user_id", "v"))
        for c in objs:
            c._made, c._earned = made.get(c.user_id, {"n": 0, "ok": 0}), earned.get(c.user_id, 0)

    def nz_href(self, obj):
        return reverse("admin:loyalty_referral_changelist") + f"?q={obj.code}"

    def nz_row(self, c):
        link = _share_link(c.code)
        return [
            format_html('<span class="nzl-code"><code>{}</code><button type="button" class="nzl-copy" data-copy="{}" title="Nusxa olish" aria-label="Kodni nusxalash"><i class="fas fa-copy" aria-hidden="true"></i></button></span>', c.code, c.code),
            ui.user_cell(c.user), format_html("<b>{}</b>", c._made["n"]) if c._made["n"] else ui.muted("0"),
            format_html('<b style="color:#16a34a">{}</b>', c._made["ok"]) if c._made["ok"] else ui.muted("0"), lx.coins(c._earned) if c._earned else ui.muted(),
            format_html('<span class="nzl-link"><span>{}</span><button type="button" class="nzl-copy" data-copy="{}" title="Havolani nusxalash" aria-label="Havolani nusxalash"><i class="fas fa-link" aria-hidden="true"></i></button></span>', link.split("//")[-1], link),
            ui.local(c.created_at, "%d.%m.%Y"),
        ]


class ChannelFilter(admin.SimpleListFilter):
    title = "Kanal"
    parameter_name = "channel"

    def lookups(self, request, model_admin):
        return [(k, v[0]) for k, v in lx.CHANNEL.items()]

    def queryset(self, request, queryset):
        return queryset.filter(channel=self.value()) if self.value() else queryset


@admin.register(Referral)
class ReferralAdmin(ExcelExportMixin, PremiumListMixin, PremiumFormMixin, admin.ModelAdmin):
    """Referallar: who invited whom, through which channel, and whether the friend's first delivered order paid both sides."""
    excel_title = "Referallar"
    excel_select_related = ("referrer", "referee")
    excel_columns = [
        Col("Sana", "created_at", 17), Col("Taklif qilgan", lambda r: r.referrer.full_name if r.referrer else "", 22), Col("Telefon", lambda r: r.referrer.phone_number if r.referrer else "", 15),
        Col("Taklif qilingan", lambda r: r.referee.full_name if r.referee else "", 22), Col("Telefoni", lambda r: r.referee.phone_number if r.referee else "", 15),
        Col("Kod", "code", 12), Col("Kanal", "channel", 10), Col("Holat", "status", 12), Col("Mukofotlangan", "rewarded_at", 17),
    ]


    list_display = ("created_at", "rewarded_at")
    list_filter = (ChannelFilter, ("created_at", _titled(admin.DateFieldListFilter, "Sana")))
    search_fields = ("code", "referrer__full_name", "referrer__phone_number", "referee__full_name", "referee__phone_number")
    list_select_related = ("referrer", "referee")
    ordering = ("-created_at",)
    list_per_page = 30
    actions = ["void_selected"]
    fields = ("referrer", "referee", "code", "channel", "status", "qualifying_order_id", "created_at", "rewarded_at")
    readonly_fields = fields

    nz_icon, nz_tone, nz_title = "fa-people-arrows", "#9333ea", "Referallar"
    nz_subtitle = "Kim kimni taklif qildi va qaysi kanal orqali. Ikkala tomon ham do'stning birinchi buyurtmasi yetkazilganda coin oladi."
    nz_columns = (("Taklif qilgan", ""), ("Taklif qilingan", ""), ("Kod", ""), ("Kanal", ""), ("Holat", ""), ("Sana", "r"))
    nz_grid = "minmax(220px,1.6fr) minmax(220px,1.6fr) 120px 140px minmax(170px,1.3fr) 120px"
    nz_search = "Kod yoki ism / telefon…"
    nz_empty = ("Hali referal yo'q", "Mijoz o'z kodini do'stiga yuborsa va do'sti ro'yxatdan o'tsa, shu yerda paydo bo'ladi.")
    nz_css = CSS
    hero_icon, hero_tone = "fa-people-arrows", "#9333ea"

    def has_add_permission(self, request) -> bool:
        return False

    def has_change_permission(self, request, obj=None) -> bool:
        return obj is None and super().has_change_permission(request, obj)  # the list's void action; a referral itself is read-only

    def has_delete_permission(self, request, obj=None) -> bool:
        return False

    def nz_stats(self, request):
        k = Referral.objects.aggregate(n=Count("id"), pending=Count("id", filter=Q(status="pending")), ok=Count("id", filter=Q(status="rewarded")), void=Count("id", filter=Q(status="void")))
        paid = WalletTransaction.objects.filter(kind=WalletTransactionKind.REFERRAL).aggregate(v=Sum("amount"))["v"] or 0
        return [
            ui.stat_card("Jami takliflar", k["n"], "#9333ea", "fa-people-arrows", f"{Referral.objects.filter(created_at__gte=_month_start()).count()} tasi shu oy"),
            ui.stat_card("Kutilmoqda", k["pending"], "#d97706", "fa-hourglass-half", "birinchi buyurtma yetkazilmagan", "?status__exact=pending"),
            ui.stat_card("Mukofotlangan", k["ok"], "#16a34a", "fa-circle-check", f"{round(k['ok'] * 100 / k['n']) if k['n'] else 0}% konversiya", "?status__exact=rewarded"),
            ui.stat_card("To'langan bonus", _som(paid), "#d4a017", "fa-coins", "ikkala tomonga", reverse("admin:loyalty_wallettransaction_changelist") + "?kind__exact=referral"),
        ]

    def nz_tabs(self, request):
        counts = dict(Referral.objects.values_list("status").annotate(n=Count("id")))
        return ui.tabs_for(request, "status__exact", [("", "Hammasi", sum(counts.values()))] + [(k, v[0].split(" ")[0] if k == "pending" else v[0], counts.get(k, 0)) for k, v in lx.REFERRAL_STATUS.items()],
                           {k: v[1] for k, v in lx.REFERRAL_STATUS.items()})

    def nz_hero_actions(self, request):
        return [{"label": "Referal kodlari", "url": reverse("admin:loyalty_referralcode_changelist"), "icon": "fa-ticket"},
                {"label": "Sozlamalar", "url": reverse("admin:loyalty_loyaltysettings_changelist"), "icon": "fa-sliders"}]

    def nz_row(self, r):
        label, color = lx.REFERRAL_STATUS.get(r.status, (r.status, "#64748b"))
        ch_label, ch_icon = lx.CHANNEL.get(r.channel, (r.channel, "fa-circle"))
        return [
            ui.user_cell(r.referrer), format_html('<span class="nzl-arrow"><i class="fas fa-arrow-right" aria-hidden="true"></i>{}</span>', ui.user_cell(r.referee)),
            format_html("<code>{}</code>", r.code), ui.chip(ch_label, "#0891b2", ch_icon),
            format_html('<span class="nzi-pill" style="--t:{}">{}</span>{}', color, label, format_html("<small>{}</small>", ui.local(r.rewarded_at, "%d.%m.%Y")) if r.rewarded_at else ""),
            format_html('<span title="{}">{}</span>', ui.local(r.created_at), ui.local(r.created_at, "%d.%m.%Y")),
        ]

    def form_hero(self, request, obj):
        label, color = lx.REFERRAL_STATUS.get(obj.status, (obj.status, "#64748b"))
        ch_label, ch_icon = lx.CHANNEL.get(obj.channel, (obj.channel, "fa-circle"))
        return {
            "title": f"{obj.referrer.full_name or obj.referrer.phone_number} → {obj.referee.full_name or obj.referee.phone_number}", "subtitle": f"kod {obj.code} · {ui.local(obj.created_at)}",
            "tone": color, "chips": [format_html('<span class="nzi-pill" style="--t:{}">{}</span>', color, label), ui.chip(ch_label, "#0891b2", ch_icon)],
            "links": [{"label": "Taklif qilgan", "url": reverse("admin:identity_user_change", args=[obj.referrer_id]), "icon": "fa-user"},
                      {"label": "Taklif qilingan", "url": reverse("admin:identity_user_change", args=[obj.referee_id]), "icon": "fa-user-plus"}],
        }

    @admin.action(description="Bekor qilish (faqat kutilayotganlar)", permissions=["change"])
    def void_selected(self, request, queryset):
        n = queryset.filter(status=ReferralStatus.PENDING).update(status=ReferralStatus.VOID)  # nothing was paid yet, so no ledger rows to reverse
        messages.success(request, f"{n} ta referal bekor qilindi." if n else "Faqat «kutilmoqda» holatidagi referallar bekor qilinadi.")


__all__ = ["LoyaltyTierAdmin", "WalletAdmin", "WalletTransactionAdmin", "WalletTransactionKind"]
